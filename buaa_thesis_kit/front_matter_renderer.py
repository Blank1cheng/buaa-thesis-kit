from __future__ import annotations

import copy
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from buaa_thesis_kit.models import ContentBlock, Metadata, ThesisModel
from buaa_thesis_kit.template_fill import _add_toc_field


SAMPLE_BODY_START = "论文封面书脊"
ASSETS_DIR = Path(__file__).resolve().parents[1] / "assets"
SEAL_ASSET = ASSETS_DIR / "buaa_seal.png"
WORDMARK_ASSET = ASSETS_DIR / "buaa_wordmark.png"
FRONT_MATTER_HEADER = "北京航空航天大学毕业设计(论文)"
SPINE_XML_MARKER = "BUAA_VERTICAL_SPINE"


def render_front_matter(
    document,
    model: ThesisModel,
    body_sections: list[ContentBlock],
) -> None:
    """Render fixed BUAA front matter before editable body content."""
    _ensure_fixed_assets()
    _replace_cover_top_block(document, model.metadata)
    _replace_cover_fields(document, model)
    _compact_cover_spacing(document)
    _trim_template_after_cover(document)
    _append_spine_page(document, model.metadata)
    _append_task_book_page(document, model.metadata)
    _append_declaration_page(document, model.metadata)
    _append_chinese_abstract_page(document, model)
    _append_english_abstract_page(document, model)
    _append_toc_page(document, body_sections)
    _start_body_section(document)


def _ensure_fixed_assets() -> None:
    missing = [str(path) for path in (SEAL_ASSET, WORDMARK_ASSET) if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing fixed BUAA front-matter asset(s): " + ", ".join(missing))


def _replace_cover_top_block(document, metadata: Metadata) -> None:
    _remove_cover_top_field_paragraphs(document)
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    row = table.rows[0]
    row.height = Cm(2.15)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    left_cell, right_cell = row.cells
    left_cell.width = Cm(8.4)
    right_cell.width = Cm(8.4)
    left_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _clear_cell(left_cell)
    _clear_cell(right_cell)

    seal_paragraph = left_cell.add_paragraph()
    seal_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    seal_paragraph.paragraph_format.space_after = Pt(0)
    seal_paragraph.paragraph_format.line_spacing = 1
    seal_paragraph.add_run().add_picture(str(SEAL_ASSET), width=Cm(1.95), height=Cm(1.95))

    for label, value in (
        ("单位代码", metadata.unit_code or "10006"),
        ("学    号", metadata.student_id),
        ("分类号", metadata.classification),
    ):
        paragraph = right_cell.add_paragraph(f"{label}       {_metadata_value(value)}")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.25
        _format_runs(paragraph, size=12)

    body = document._body._element
    body.remove(table._tbl)
    body.insert(0, table._tbl)


def _remove_cover_top_field_paragraphs(document) -> None:
    for paragraph in list(_cover_paragraphs(document)):
        text = paragraph.text
        compact = re.sub(r"\s+", "", text)
        if "单位代码" in compact or "分类号" in compact or ("学" in compact and "号" in compact and len(compact) < 32):
            paragraph._element.getparent().remove(paragraph._element)


def _replace_cover_fields(document, model: ThesisModel) -> None:
    metadata = model.metadata
    title_lines = _title_cn_lines(model)
    for paragraph in _cover_paragraphs(document):
        text = paragraph.text
        stripped = text.strip()
        if "单位代码" in text:
            _replace_paragraph_text_preserving_style(
                paragraph,
                f"单位代码       {_metadata_value(metadata.unit_code or '10006')}",
                alignment=WD_ALIGN_PARAGRAPH.RIGHT,
            )
        elif "学" in text and "号" in text and "学号" not in text:
            _replace_paragraph_text_preserving_style(
                paragraph,
                f"学    号       {_metadata_value(metadata.student_id)}",
                alignment=WD_ALIGN_PARAGRAPH.RIGHT,
            )
        elif "分类号" in text:
            _replace_paragraph_text_preserving_style(
                paragraph,
                f"分类号       {_metadata_value(metadata.classification)}",
                alignment=WD_ALIGN_PARAGRAPH.RIGHT,
            )
        elif stripped in {"（题目）", "(题目)"}:
            _replace_paragraph_lines_preserving_style(paragraph, title_lines)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _format_runs(paragraph, bold=True, size=16)
        elif re.fullmatch(r"\d{4}年\d{1,2}月", stripped):
            _replace_paragraph_text_preserving_style(
                paragraph,
                _date_year_month(metadata.date),
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
            )

    for table in document.tables:
        _replace_cover_table(table, metadata)


def _cover_paragraphs(document):
    for paragraph in document.paragraphs:
        if SAMPLE_BODY_START in paragraph.text:
            break
        yield paragraph


def _replace_cover_table(table, metadata: Metadata) -> None:
    values = {
        "学院名称": metadata.college,
        "院（系）名称": metadata.college,
        "专业名称": metadata.major,
        "学生姓名": metadata.student_name,
        "指导教师": metadata.advisor,
    }
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        if len(row.cells) < 2:
            continue
        label = re.sub(r"\s+", "", row.cells[0].text)
        for key, value in values.items():
            if key in label:
                _replace_cell_text_preserving_style(
                    row.cells[1],
                    _metadata_value(value),
                    font_size=_cover_table_value_font_size(value),
                    alignment=WD_ALIGN_PARAGRAPH.CENTER,
                )
                break


def _trim_template_after_cover(document) -> None:
    body = document._body._element
    remove = False
    for child in list(body):
        if child.tag == qn("w:p") and SAMPLE_BODY_START in _xml_text(child):
            _remove_trailing_blank_paragraphs(body, child)
            remove = True
        if remove and child.tag != qn("w:sectPr"):
            body.remove(child)


def _compact_cover_spacing(document) -> None:
    body = document._body._element
    consecutive_blanks = 0
    for child in list(body):
        if child.tag == qn("w:p") and SAMPLE_BODY_START in _xml_text(child):
            break
        if child.tag == qn("w:p") and _is_removable_blank_paragraph(child):
            consecutive_blanks += 1
            if consecutive_blanks > 2:
                body.remove(child)
            continue
        consecutive_blanks = 0


def _remove_trailing_blank_paragraphs(body, marker) -> None:
    current = marker.getprevious()
    while current is not None and current.tag == qn("w:p") and _is_removable_blank_paragraph(current):
        previous = current.getprevious()
        body.remove(current)
        current = previous


def _is_removable_blank_paragraph(paragraph_element) -> bool:
    if _xml_text(paragraph_element).strip():
        return False
    protected_tags = {qn("w:drawing"), qn("w:pict")}
    return not any(node.tag in protected_tags for node in paragraph_element.iter())


def _append_spine_page(document, metadata: Metadata) -> None:
    section = _new_section(document)
    _clear_header_footer(section)
    _set_a4_page(section, top=2.0, bottom=2.0, left=2.0, right=2.0)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run()._r.append(_vertical_spine_pict(metadata))


def _vertical_spine_pict(metadata: Metadata):
    blocks = [
        _metadata_value(metadata.title_cn or metadata.title_en),
        _metadata_value(metadata.student_name),
        "北京航空航天大学",
    ]
    paragraphs = "\n".join(_spine_block_paragraph_xml(block, index) for index, block in enumerate(blocks) if block)
    xml = f"""
    <w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            xmlns:v="urn:schemas-microsoft-com:vml">
      <v:shape id="{SPINE_XML_MARKER}" type="#_x0000_t202"
        style="position:absolute;margin-left:274pt;margin-top:14pt;width:50pt;height:715pt;z-index:1;
               mso-position-horizontal-relative:page;mso-position-vertical-relative:page"
        filled="f" stroked="t" strokeweight="0.75pt">
        <v:textbox inset="0,96pt,0,0">
          <w:txbxContent>{paragraphs}</w:txbxContent>
        </v:textbox>
      </v:shape>
    </w:pict>
    """
    return parse_xml(xml)


def _spine_block_paragraph_xml(text: str, index: int) -> str:
    spacing_after = "420" if index < 2 else "0"
    return (
        '<w:p>'
        '<w:pPr>'
        '<w:textDirection w:val="tbRl"/>'
        '<w:jc w:val="center"/>'
        f'<w:spacing w:after="{spacing_after}"/>'
        '</w:pPr>'
        '<w:r><w:rPr><w:rFonts w:ascii="SimSun" w:eastAsia="SimSun"/>'
        '<w:sz w:val="24"/><w:szCs w:val="24"/></w:rPr>'
        f'{_spine_vertical_text_xml(text)}</w:r>'
        '</w:p>'
    )


def _spine_vertical_text_xml(text: str) -> str:
    chars = [char for char in _metadata_value(text) if not char.isspace()]
    if not chars:
        return "<w:t></w:t>"
    parts: list[str] = []
    for index, char in enumerate(chars):
        parts.append(f"<w:t>{escape(char)}</w:t>")
        if index < len(chars) - 1:
            parts.append("<w:br/>")
    return "".join(parts)


def _append_task_book_page(document, metadata: Metadata) -> None:
    section = _new_section(document)
    _clear_header_footer(section)
    _set_a4_page(section, top=2.4, bottom=2.4, left=2.8, right=2.4)
    _center_heading(document, "本科毕业设计（论文）任务书", size=16)
    _metadata_table(
        document,
        [
            ("毕业设计（论文）题目", metadata.title_cn or metadata.title_en),
            ("学生姓名", metadata.student_name),
            ("学号", metadata.student_id),
            ("学院", metadata.college),
            ("专业", metadata.major),
            ("指导教师", metadata.advisor),
        ],
    )
    _labeled_blank_block(document, "毕业设计（论文）使用的原始资料（数据）及设计技术要求：", lines=4)
    _labeled_blank_block(document, "毕业设计（论文）工作内容：", lines=5)
    _labeled_blank_block(document, "主要参考资料：", lines=4)


def _append_declaration_page(document, metadata: Metadata) -> None:
    section = _new_section(document)
    _clear_header_footer(section)
    _set_a4_page(section, top=2.8, bottom=2.4, left=3.0, right=2.6)
    _center_heading(document, "本人声明", size=16)
    _body_paragraph(
        document,
        "本人郑重声明：所呈交的毕业设计（论文）是在指导教师指导下独立完成的。"
        "除文中已经注明引用的内容外，本论文不包含任何其他个人或集体已经发表或撰写过的研究成果。",
        first_line_indent=True,
    )
    _body_paragraph(
        document,
        "对本文研究做出重要贡献的个人和集体，均已在文中以明确方式标明。本人完全意识到本声明的法律结果由本人承担。",
        first_line_indent=True,
    )
    _signature_table(document, metadata)


def _append_chinese_abstract_page(document, model: ThesisModel) -> None:
    section = _new_section(document)
    _set_front_matter_header_footer(section, start=1)
    metadata = model.metadata
    _abstract_title(document, metadata.title_cn or metadata.title_en)
    _abstract_author_row(document, f"学生：{metadata.student_name}", f"指导老师：{metadata.advisor}")
    _center_heading(document, "摘    要", size=15)
    _add_text_paragraphs(document, _front_matter_value(model, "chinese_abstract", "abstract_cn", "cn_abstract"))
    keywords = _front_matter_value(model, "keywords_cn", "chinese_keywords", "cn_keywords", "keywords")
    _body_paragraph(document, f"关键词：{keywords}")


def _append_english_abstract_page(document, model: ThesisModel) -> None:
    section = _new_section(document)
    _set_front_matter_header_footer(section)
    title = _front_matter_value(model, "title_en", "english_title") or model.metadata.title_en
    _abstract_title(document, title or model.metadata.title_cn)
    author = _front_matter_value(model, "author_en") or model.metadata.student_name
    tutor = _front_matter_value(model, "tutor_en") or model.metadata.advisor
    _abstract_author_row(document, f"Author: {author}", f"Tutor: {tutor}")
    _center_heading(document, "Abstract", size=15)
    _add_text_paragraphs(document, _front_matter_value(model, "english_abstract", "abstract_en", "en_abstract"))
    keywords = _front_matter_value(model, "keywords_en", "english_keywords", "en_keywords")
    _body_paragraph(document, f"Key Words: {keywords}")


def _append_toc_page(document, body_sections: list[ContentBlock]) -> None:
    if not body_sections:
        return
    section = _new_section(document)
    _set_front_matter_header_footer(section)
    _center_heading(document, "目录", size=16)
    _add_toc_field(document.add_paragraph())


def _start_body_section(document) -> None:
    section = _new_section(document)
    _set_body_header_footer(section)


def _new_section(document):
    return document.add_section(WD_SECTION.NEW_PAGE)


def _set_a4_page(section, *, top: float, bottom: float, left: float, right: float) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)


def _clear_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    _clear_part(section.header)
    _clear_part(section.footer)


def _set_front_matter_header_footer(section, *, start: int | None = None) -> None:
    _set_a4_page(section, top=2.4, bottom=2.2, left=3.0, right=2.6)
    _set_page_numbering(section, fmt="upperRoman", start=start)
    _set_header_footer(section, FRONT_MATTER_HEADER, roman=True)


def _set_body_header_footer(section) -> None:
    _set_a4_page(section, top=2.4, bottom=2.2, left=3.0, right=2.6)
    _set_page_numbering(section, fmt="decimal", start=1)
    _set_header_footer(section, FRONT_MATTER_HEADER, roman=False)


def _set_header_footer(section, header_text: str, *, roman: bool) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    _clear_part(section.header)
    _clear_part(section.footer)
    header = section.header.add_paragraph(header_text)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_runs(header, size=10.5)
    footer = section.footer.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    _add_page_field(footer)
    footer.add_run(" 页")
    _format_runs(footer, size=10.5 if roman else 10.5)


def _clear_part(part) -> None:
    for child in list(part._element):
        part._element.remove(child)


def _clear_cell(cell) -> None:
    for child in list(cell._tc):
        if child.tag != qn("w:tcPr"):
            cell._tc.remove(child)


def _set_page_numbering(section, *, fmt: str, start: int | None) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:fmt"), fmt)
    if start is not None:
        pg_num.set(qn("w:start"), str(start))
    sect_pr.append(pg_num)


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, end):
        run._r.append(node)


def _metadata_table(document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, (label, value) in zip(table.rows, rows):
        row.cells[0].width = Cm(5)
        row.cells[1].width = Cm(9)
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _replace_cell_text(row.cells[0], f"{label}：", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        _replace_cell_text(row.cells[1], _metadata_value(value), alignment=WD_ALIGN_PARAGRAPH.CENTER)


def _labeled_blank_block(document, label: str, *, lines: int) -> None:
    _body_paragraph(document, label)
    for _ in range(lines):
        paragraph = document.add_paragraph(" " * 2)
        _format_runs(paragraph, size=12)
        _add_bottom_border(paragraph)


def _signature_table(document, metadata: Metadata) -> None:
    table = document.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    values = [
        ("学生签名", metadata.student_name),
        ("日期", _date_year_month(metadata.date)),
        ("指导教师签名", metadata.advisor),
    ]
    for row, (label, value) in zip(table.rows, values):
        _replace_cell_text(row.cells[0], f"{label}：", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        _replace_cell_text(row.cells[1], _metadata_value(value), alignment=WD_ALIGN_PARAGRAPH.CENTER)


def _abstract_title(document, title: str) -> None:
    for line in _title_lines_from_text(title):
        paragraph = document.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _format_runs(paragraph, bold=True, size=14)


def _abstract_author_row(document, left: str, right: str) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    _replace_cell_text(table.cell(0, 0), left, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _replace_cell_text(table.cell(0, 1), right, alignment=WD_ALIGN_PARAGRAPH.RIGHT)


def _center_heading(document, text: str, *, size: float) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_runs(paragraph, bold=True, size=size)


def _add_text_paragraphs(document, text: str) -> None:
    for line in str(text or "").splitlines():
        value = line.strip()
        if value:
            _body_paragraph(document, value, first_line_indent=True)


def _body_paragraph(document, text: str, *, first_line_indent: bool = False) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_line_indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing = 1.5
    _format_runs(paragraph, size=12)


def _add_bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)


def _replace_cell_text(cell, text: str, *, alignment=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _replace_paragraph_text_preserving_style(paragraph, text, alignment=alignment)
    _format_runs(paragraph, size=12)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def _replace_cell_text_preserving_style(
    cell,
    text: str,
    *,
    font_size: float | None = None,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _replace_paragraph_text_preserving_style(paragraph, text, alignment=alignment)
    if font_size is not None:
        _format_runs(paragraph, size=font_size)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def _replace_paragraph_text_preserving_style(
    paragraph,
    text: str,
    *,
    alignment=None,
) -> None:
    first_run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_run_properties = copy.deepcopy(paragraph.runs[0]._r.rPr)
    paragraph.clear()
    if alignment is not None:
        paragraph.alignment = alignment
    if not text:
        return
    run = paragraph.add_run(text)
    if first_run_properties is not None:
        run._r.insert(0, first_run_properties)


def _replace_paragraph_lines_preserving_style(paragraph, lines: list[str]) -> None:
    first_run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_run_properties = copy.deepcopy(paragraph.runs[0]._r.rPr)
    paragraph.clear()
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        if first_run_properties is not None:
            run._r.insert(0, copy.deepcopy(first_run_properties))


def _format_runs(paragraph, *, bold: bool = False, size: float = 12) -> None:
    for run in paragraph.runs:
        run.bold = bold
        run.font.size = Pt(size)


def _title_cn_lines(model: ThesisModel) -> list[str]:
    explicit = [
        _metadata_value(model.front_matter.get("title_cn_line1", "")),
        _metadata_value(model.front_matter.get("title_cn_line2", "")),
    ]
    if any(explicit):
        return [line for line in explicit if line]
    return _title_cn_lines_from_text(model.metadata.title_cn or model.metadata.title_en or "Untitled Thesis")


def _title_cn_lines_from_text(title: str) -> list[str]:
    value = _metadata_value(title)
    if len(value) <= 18:
        return [value]
    for marker in ("关键技术", "方法研究", "系统研究", "性能评估"):
        index = value.find(marker)
        if 8 <= index <= len(value) - 4:
            return [value[:index], value[index:]]
    split_at = min(18, max(10, len(value) // 2 + 2))
    if len(value) - split_at <= 2:
        split_at = max(1, len(value) - 6)
    return [value[:split_at], value[split_at:]]


def _title_lines_from_text(title: str) -> list[str]:
    value = _metadata_value(title)
    if not value:
        return []
    if re.search(r"[A-Za-z]", value) and " " in value and not re.search(r"[\u4e00-\u9fff]", value):
        return _wrap_english_title(value, max_chars=58)
    return _title_cn_lines_from_text(value)


def _wrap_english_title(title: str, *, max_chars: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for word in title.split():
        candidate = f"{current} {word}".strip()
        if current and len(candidate) > max_chars:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def _front_matter_value(model: ThesisModel, *keys: str) -> str:
    for key in keys:
        value = model.front_matter.get(key)
        if isinstance(value, list):
            value = "\n".join(str(item) for item in value if str(item).strip())
        value = str(value or "").strip()
        if value:
            return value
    return ""


def _date_year_month(value: str) -> str:
    text = _metadata_value(value)
    match = re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月", text)
    if match:
        return f"{match.group(1)}年{int(match.group(2))}月"
    match = re.search(r"(\d{4})[-/.](\d{1,2})", text)
    if match:
        return f"{match.group(1)}年{int(match.group(2))}月"
    return text


def _cover_table_value_font_size(value: str) -> float | None:
    length = len(_metadata_value(value))
    if length >= 16:
        return 10.5
    if length >= 10:
        return 12
    return None


def _metadata_value(value: str) -> str:
    return str(value or "").strip()


def _xml_text(element) -> str:
    return "".join(node.text or "" for node in element.iter() if node.tag == qn("w:t"))


__all__ = [
    "SEAL_ASSET",
    "SPINE_XML_MARKER",
    "WORDMARK_ASSET",
    "render_front_matter",
]
