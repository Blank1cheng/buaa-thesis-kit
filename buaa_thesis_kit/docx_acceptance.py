from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from docx import Document

from buaa_thesis_kit.models import Metadata


SPINE_MARKERS = ("Book Spine", "书脊")
PDF_PLACEHOLDER_HEADING = "PDF Extracted Text"
EMU_PER_INCH = 914400
PAGE_SCREENSHOT_MIN_WIDTH_IN = 5.0
PAGE_SCREENSHOT_MIN_HEIGHT_IN = 7.4
MAX_BODY_SNIPPETS = 8
MIN_BODY_SNIPPET_CHARS = 20
MAX_BODY_SNIPPET_CHARS = 120
FORBIDDEN_VISIBLE_TEXT_MARKERS = (
    "[Figure inserted]",
    "[Figure requires review]",
    "[Equation preview inserted]",
    "本页由规范化流水线",
    "需人工复核",
    "References 作为中文参考文献标题",
    "MERGEFORMAT",
    "公式章",
    "下一章",
    ".worktrees",
    "output_work_",
    "D:\\",
    "image1.png",
    ".wmf",
    ".emf",
    ".png",
)
COVER_FIELD_LEAK_MARKERS = (
    "院（系）名称",
    "专业名称",
    "学生姓名",
    "指导教师",
)
FORBIDDEN_VISIBLE_TEXT_PATTERNS = (
    re.compile(r"image\d+\.(?:wmf|emf)", flags=re.IGNORECASE),
)


@dataclass
class DocxOutputInspection:
    blocking_items: list[str] = field(default_factory=list)
    manual_review: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    editability: dict[str, int] = field(default_factory=dict)


def inspect_docx_output(
    docx_path: Path,
    metadata: Metadata,
    *,
    source_kind: str,
    require_spine: bool,
    required_body_snippets: Iterable[str] | None = None,
    expected_omml_equation_count: int = 0,
) -> DocxOutputInspection:
    """Check hard acceptance gates for an authoritative, editable Word output."""
    result = DocxOutputInspection()
    path = Path(docx_path)
    if not path.exists() or not path.is_file():
        result.blocking_items.append(f"word_output_missing: {path}")
        return result

    try:
        package_info = _inspect_docx_package(path)
        document = Document(str(path))
    except Exception as exc:
        result.blocking_items.append(f"word_output_unreadable: {exc}")
        return result

    visible_text = _merge_visible_text(_document_visible_text(document), str(package_info["document_text"]))
    compact_text = _compact_text(visible_text)
    editable_chars = len(compact_text)
    body_snippets = _normalized_body_snippets(required_body_snippets or [])
    body_snippet_hits = _body_snippet_hits(compact_text, body_snippets)
    result.editability = {
        "editable_characters": editable_chars,
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "drawing_count": package_info["drawing_count"],
        "page_screenshot_drawing_count": package_info["page_screenshot_drawing_count"],
        "omml_equation_count": package_info["omml_equation_count"],
        "expected_omml_equation_count": max(0, expected_omml_equation_count),
        "media_count": package_info["media_count"],
        "body_snippet_count": len(body_snippets),
        "body_snippet_hits": body_snippet_hits,
    }

    if require_spine and not (
        _contains_any(visible_text, SPINE_MARKERS) or int(package_info["vertical_spine_count"]) > 0
    ):
        result.blocking_items.append("spine_missing: authoritative Word output has no book spine marker.")

    _inspect_required_editable_text(
        result,
        compact_text,
        metadata,
        package_info,
        source_kind=source_kind,
    )
    _inspect_required_editable_body_text(
        result,
        body_snippets,
        body_snippet_hits,
        source_kind=source_kind,
    )

    if source_kind.lower() == "pdf" and package_info["page_screenshot_drawing_count"]:
        result.blocking_items.append(
            "word_page_screenshot: PDF-derived Word output contains "
            f"{package_info['page_screenshot_drawing_count']} page-sized drawing(s); "
            "keep page renders only as OCR evidence and rebuild Word content as editable text/tables/equations."
        )

    if PDF_PLACEHOLDER_HEADING in visible_text:
        result.blocking_items.append(
            "pdf_placeholder_heading_visible: remove internal PDF Extracted Text heading from Word output."
        )

    if "[Equation preview inserted]" in visible_text:
        result.blocking_items.append(
            "editable_equation_missing: equation preview image was inserted instead of an editable Word/OMML/OLE equation object."
        )

    expected_omml = max(0, expected_omml_equation_count)
    actual_omml = package_info["omml_equation_count"]
    if expected_omml and actual_omml < expected_omml:
        result.blocking_items.append(
            "editable_omml_equation_missing: expected "
            f"{expected_omml} editable OMML equation object(s), found {actual_omml}."
        )

    if "__BUAA_EDITABLE_EQUATION_OBJECT__" in visible_text:
        result.blocking_items.append(
            "editable_equation_object_token_visible: internal editable equation placeholder was not converted to a Word/OLE object."
        )

    _inspect_forbidden_visible_text(result, visible_text, document)
    _inspect_formula_token_dump(result, document)

    unresolved = re.findall(r"\{\{\s*[A-Z_]+\s*\}\}", visible_text)
    if unresolved:
        result.blocking_items.append(
            f"unresolved_word_placeholders: {', '.join(sorted(set(unresolved)))}"
        )

    if not result.blocking_items:
        result.notes.append(
            f"editable Word validation passed: {editable_chars} editable characters, "
            f"{package_info['drawing_count']} drawings, "
            f"{package_info['page_screenshot_drawing_count']} page-sized screenshots."
        )
    return result


def _inspect_required_editable_text(
    result: DocxOutputInspection,
    compact_text: str,
    metadata: Metadata,
    package_info: dict[str, int],
    *,
    source_kind: str,
) -> None:
    source_label = source_kind.upper() if source_kind else "SOURCE"
    required_values = [
        value
        for value in [
            metadata.title_cn or metadata.title_en,
            metadata.student_id,
        ]
        if _compact_text(value)
    ]
    missing_values = [
        value for value in required_values if _compact_text(value) not in compact_text
    ]
    if missing_values:
        result.blocking_items.append(
            f"editable_text_missing: {source_label}-derived Word output does not expose required "
            f"metadata as editable text: {', '.join(missing_values)}"
        )

    if package_info["drawing_count"] > 0 and len(compact_text) < 30:
        result.blocking_items.append(
            f"editable_text_missing: {source_label}-derived Word output appears image-only."
        )


def _inspect_required_editable_body_text(
    result: DocxOutputInspection,
    body_snippets: list[str],
    body_snippet_hits: int,
    *,
    source_kind: str,
) -> None:
    if not body_snippets or body_snippet_hits:
        return
    source_label = source_kind.upper() if source_kind else "SOURCE"
    result.blocking_items.append(
        f"editable_body_text_missing: {source_label}-derived Word output does not expose "
        f"sampled source body text as editable text: {len(body_snippets)} sample(s) checked."
    )


def _inspect_docx_package(path: Path) -> dict[str, int | str]:
    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        document_xml = package.read("word/document.xml").decode("utf-8", errors="replace")
    drawing_sizes = _drawing_sizes(document_xml)
    return {
        "document_text": _xml_visible_text(document_xml),
        "media_count": sum(1 for name in names if name.startswith("word/media/")),
        "drawing_count": document_xml.count("<w:drawing"),
        "vertical_spine_count": _vertical_spine_count(document_xml),
        "page_screenshot_drawing_count": sum(
            1 for width, height in drawing_sizes if _looks_like_page_screenshot(width, height)
        ),
        "omml_equation_count": _omml_equation_count(document_xml),
    }


def _omml_equation_count(document_xml: str) -> int:
    omath_para_count = len(re.findall(r"<m:oMathPara\b", document_xml))
    if omath_para_count:
        return omath_para_count
    return len(re.findall(r"<m:oMath\b", document_xml))


def _drawing_sizes(document_xml: str) -> list[tuple[float, float]]:
    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    }
    try:
        root = ET.fromstring(document_xml)
    except ET.ParseError:
        return []

    sizes: list[tuple[float, float]] = []
    for drawing in root.findall(".//w:drawing", namespaces):
        extent = drawing.find(".//wp:extent", namespaces)
        if extent is None:
            continue
        try:
            width = int(extent.attrib.get("cx", "0")) / EMU_PER_INCH
            height = int(extent.attrib.get("cy", "0")) / EMU_PER_INCH
        except ValueError:
            continue
        sizes.append((width, height))
    return sizes


def _looks_like_page_screenshot(width_in: float, height_in: float) -> bool:
    long_edge = max(width_in, height_in)
    short_edge = min(width_in, height_in)
    return long_edge >= PAGE_SCREENSHOT_MIN_HEIGHT_IN and short_edge >= PAGE_SCREENSHOT_MIN_WIDTH_IN


def _xml_visible_text(document_xml: str) -> str:
    try:
        root = ET.fromstring(document_xml)
    except ET.ParseError:
        return ""
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    return "\n".join(node.text or "" for node in root.findall(".//w:t", namespaces))


def _vertical_spine_count(document_xml: str) -> int:
    count = len(re.findall(r"<w:textDirection\b[^>]*w:val=\"tbRl\"", document_xml))
    if "BUAA_VERTICAL_SPINE" in document_xml and count == 0:
        return 1
    return count


def _merge_visible_text(*parts: str) -> str:
    merged: list[str] = []
    seen: set[str] = set()
    for part in parts:
        for line in str(part or "").splitlines():
            if not line or line in seen:
                continue
            merged.append(line)
            seen.add(line)
    return "\n".join(merged)


def _document_visible_text(document) -> str:
    parts: list[str] = []
    parts.extend(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(part for part in parts if part is not None)


def _inspect_forbidden_visible_text(
    result: DocxOutputInspection,
    visible_text: str,
    document,
) -> None:
    found = [marker for marker in FORBIDDEN_VISIBLE_TEXT_MARKERS if marker in visible_text]
    for pattern in FORBIDDEN_VISIBLE_TEXT_PATTERNS:
        found.extend(sorted(set(pattern.findall(visible_text))))
    found.extend(_cover_field_leaks(document))
    if found:
        result.blocking_items.append(
            "unsafe_word_body_text: final Word output exposes process/debug text or local asset paths: "
            + ", ".join(found)
        )
    for paragraph_text in _document_visible_paragraphs(document):
        if paragraph_text.strip() == "References":
            result.blocking_items.append(
                "english_references_heading_visible: Chinese BUAA thesis output must use 参考文献."
            )
            break


def _inspect_formula_token_dump(result: DocxOutputInspection, document) -> None:
    consecutive = 0
    for paragraph_text in _document_visible_paragraphs(document):
        if _looks_like_formula_token_dump_line(paragraph_text):
            consecutive += 1
            if consecutive >= 5:
                result.blocking_items.append(
                    "formula_token_dump_visible: five or more consecutive equation-like token lines were rendered as normal body text."
                )
                return
        else:
            consecutive = 0


def _cover_field_leaks(document) -> list[str]:
    leaks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        compact = _compact_text(text)
        exact_hits = [
            marker
            for marker in COVER_FIELD_LEAK_MARKERS
            if compact == _compact_text(marker)
        ]
        cluster_hits = [
            marker
            for marker in COVER_FIELD_LEAK_MARKERS
            if _compact_text(marker) in compact
        ]
        if exact_hits:
            leaks.extend(exact_hits)
        elif len(cluster_hits) >= 2:
            leaks.extend(cluster_hits)
    return sorted(set(leaks))


def _document_visible_paragraphs(document) -> list[str]:
    parts: list[str] = []
    parts.extend(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return [part for part in parts if part is not None]


def _looks_like_formula_token_dump_line(text: str) -> bool:
    value = str(text or "").strip()
    if not value or len(value) > 80:
        return False
    if value.startswith("[公式缺失："):
        return False
    if re.search(r"[\u4e00-\u9fff]", value):
        return False
    if not re.search(r"[A-Za-z]", value):
        return False
    if _looks_like_reference_fragment(value):
        return False
    if "=" in value:
        return True
    if re.search(r"\\(?:frac|sum|int|sqrt|left|right|theta|alpha|beta|gamma|omega)\b", value):
        return True
    if re.search(r"\b(?:OTF|MTF|PSF)\s*\(", value, flags=re.IGNORECASE):
        return True
    if re.match(r"^[A-Za-z]\s*\([A-Za-z0-9_,\s]+\)(?:\s*[+\-*/^].*)?$", value):
        return True
    compact = re.sub(r"\s+", "", value)
    operator_count = len(re.findall(r"[+*/^_\\]", compact))
    word_count = len(re.findall(r"[A-Za-z]{2,}", value))
    if len(compact) <= 40 and operator_count >= 1 and word_count <= 2:
        return True
    return False


def _looks_like_reference_fragment(value: str) -> bool:
    if re.match(r"^\[\d+\]", value):
        return True
    if re.search(r"\[(?:J|M|C|D|R|P|S|EB/OL|OL)\]", value, flags=re.IGNORECASE):
        return True
    if re.search(r"\b(?:19|20)\d{2}\b", value):
        return True
    if re.search(
        r"\b(?:Journal|Engineering|Science|Technology|Proceedings|Transactions|"
        r"Prediction|Reliability|Mechanical|Nuclear)\b",
        value,
        flags=re.IGNORECASE,
    ):
        return True
    return False


def _contains_any(text: str, needles: Iterable[str]) -> bool:
    return any(needle in text for needle in needles)


def _normalized_body_snippets(snippets: Iterable[str]) -> list[str]:
    normalized: list[str] = []
    seen: set[str] = set()
    for snippet in snippets:
        value = _compact_text(snippet)
        if len(value) < MIN_BODY_SNIPPET_CHARS:
            continue
        value = value[:MAX_BODY_SNIPPET_CHARS]
        if value in seen:
            continue
        normalized.append(value)
        seen.add(value)
        if len(normalized) >= MAX_BODY_SNIPPETS:
            break
    return normalized


def _body_snippet_hits(compact_text: str, snippets: Iterable[str]) -> int:
    return sum(1 for snippet in snippets if snippet in compact_text)


def _compact_text(text: str) -> str:
    return re.sub(r"\s+", "", str(text or ""))


__all__ = ["DocxOutputInspection", "inspect_docx_output"]
