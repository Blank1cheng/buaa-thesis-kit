from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Mm, Pt
from lxml import etree

from buaa_thesis_kit.models import AssetItem, ContentBlock, EquationItem, ThesisModel


PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Z_]+)\s*\}\}")
BLOCK_PLACEHOLDERS = {"BODY", "REFERENCES", "TABLES", "FIGURES", "EQUATIONS", "APPENDICES"}
SUPPORTED_IMAGE_SUFFIXES = {".bmp", ".gif", ".jpg", ".jpeg", ".png", ".tif", ".tiff"}
OCR_EVIDENCE_FIGURE_TYPES = {"pdf-page-image"}
EQUATION_OBJECT_TOKEN_PREFIX = "__BUAA_EDITABLE_EQUATION_OBJECT__"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
IMAGE_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
OLE_OBJECT_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject"
OLE_OBJECT_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.oleObject"
IMAGE_CONTENT_TYPES = {
    ".bmp": "image/bmp",
    ".emf": "image/x-emf",
    ".gif": "image/gif",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
    ".wmf": "image/x-wmf",
}


@dataclass(frozen=True)
class RenderedParagraph:
    text: str
    style: str | None = None
    alignment: WD_ALIGN_PARAGRAPH | None = None
    bold: bool = False


@dataclass(frozen=True)
class RenderedTable:
    block: ContentBlock


@dataclass(frozen=True)
class RenderedFigure:
    figure: AssetItem


@dataclass(frozen=True)
class RenderedEquation:
    equation: EquationItem


RenderedItem = RenderedParagraph | RenderedTable | RenderedFigure | RenderedEquation


def fill_word_template(template_path: Path, model: ThesisModel, output_path: Path) -> None:
    """Fill a DOCX template from a ThesisModel and save the authoritative Word file."""
    template_path = Path(template_path)
    output_path = Path(output_path)
    document = Document(str(template_path))

    has_placeholders = _document_has_placeholders(document)
    _apply_conservative_formatting(document, apply_page_setup=False)
    if has_placeholders:
        represented_placeholders = _represented_placeholders(document)
        _fill_placeholder_document(document, model)
        _remove_remaining_placeholders(document)
        _append_missing_content(document, model, represented_placeholders)
    else:
        _clear_body_preserving_sections(document)
        _apply_conservative_formatting(document, apply_page_setup=True)
        _build_fallback_document(document, model)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    finalize_equation_objects(output_path, model.equations)


def _document_has_placeholders(document) -> bool:
    return any(PLACEHOLDER_RE.search(paragraph.text) for paragraph in _iter_all_paragraphs(document))


def _represented_placeholders(document) -> set[str]:
    represented: set[str] = set()
    for paragraph in _iter_all_paragraphs(document):
        for match in PLACEHOLDER_RE.finditer(paragraph.text):
            represented.add(match.group(1))
    return represented


def _fill_placeholder_document(document, model: ThesisModel) -> None:
    for paragraph in list(_iter_all_paragraphs(document)):
        if not PLACEHOLDER_RE.search(paragraph.text):
            continue
        if _has_block_placeholder(paragraph.text):
            rendered = _render_placeholder_items(paragraph.text, model)
            _replace_paragraph_with_items(paragraph, rendered)
        else:
            if not _replace_inline_placeholders_in_runs(paragraph, model):
                replacement = _replace_inline_placeholders(paragraph.text, model)
                _replace_paragraph_text(paragraph, _remove_unresolved_placeholders(replacement))


def _has_block_placeholder(text: str) -> bool:
    return any(match.group(1) in BLOCK_PLACEHOLDERS for match in PLACEHOLDER_RE.finditer(text))


def _render_placeholder_items(text: str, model: ThesisModel) -> list[RenderedItem]:
    rendered: list[RenderedItem] = []
    cursor = 0
    for match in PLACEHOLDER_RE.finditer(text):
        leading = text[cursor : match.start()]
        if leading:
            inline = _remove_unresolved_placeholders(_replace_inline_placeholders(leading, model))
            if inline:
                rendered.append(RenderedParagraph(inline))

        token = match.group(1)
        if token in BLOCK_PLACEHOLDERS:
            rendered.extend(_render_block_placeholder(token, model))
        else:
            inline = _remove_unresolved_placeholders(_scalar_placeholders(model).get(token, ""))
            if inline:
                rendered.append(RenderedParagraph(inline))
        cursor = match.end()

    trailing = text[cursor:]
    if trailing:
        inline = _remove_unresolved_placeholders(_replace_inline_placeholders(trailing, model))
        if inline:
            rendered.append(RenderedParagraph(inline))
    return rendered


def _render_block_placeholder(token: str, model: ThesisModel) -> list[RenderedItem]:
    if token == "BODY":
        return _render_sections(model.sections)
    if token == "REFERENCES":
        return _render_references(model.references)
    if token == "TABLES":
        return [RenderedTable(table) for table in model.tables]
    if token == "FIGURES":
        return [RenderedFigure(figure) for figure in model.figures]
    if token == "EQUATIONS":
        return _render_equations(model.equations)
    if token == "APPENDICES":
        return _render_appendices(model.appendices)
    return []


def _replace_inline_placeholders(text: str, model: ThesisModel) -> str:
    values = _scalar_placeholders(model)

    def replace(match: re.Match[str]) -> str:
        return values.get(match.group(1), match.group(0))

    return PLACEHOLDER_RE.sub(replace, text)


def _replace_inline_placeholders_in_runs(paragraph: Paragraph, model: ThesisModel) -> bool:
    original_text = paragraph.text
    changed = False
    for run in paragraph.runs:
        if PLACEHOLDER_RE.search(run.text):
            run.text = _remove_unresolved_placeholders_preserving_spacing(
                _replace_inline_placeholders(run.text, model)
            )
            changed = True

    if not changed:
        return False
    if PLACEHOLDER_RE.search(paragraph.text):
        replacement = _replace_inline_placeholders(original_text, model)
        _replace_paragraph_text(paragraph, _remove_unresolved_placeholders(replacement))
        return False
    return True


def _scalar_placeholders(model: ThesisModel) -> dict[str, str]:
    metadata = model.metadata
    return {
        "TITLE_CN": metadata.title_cn,
        "TITLE_EN": metadata.title_en,
        "STUDENT_NAME": metadata.student_name,
        "STUDENT_ID": metadata.student_id,
        "COLLEGE": metadata.college,
        "MAJOR": metadata.major,
        "ADVISOR": metadata.advisor,
        "DATE": metadata.date,
        "CLASSIFICATION": metadata.classification,
        "UNIT_CODE": metadata.unit_code,
        "ABSTRACT_CN": _front_matter_value(
            model,
            "chinese_abstract",
            "abstract_cn",
            "cn_abstract",
        ),
        "ABSTRACT_EN": _front_matter_value(
            model,
            "english_abstract",
            "abstract_en",
            "en_abstract",
        ),
    }


def _front_matter_value(model: ThesisModel, *keys: str) -> str:
    for key in keys:
        value = model.front_matter.get(key)
        if isinstance(value, str):
            return value
        if isinstance(value, list):
            return "\n".join(str(item) for item in value if item)
    return ""


def _remove_remaining_placeholders(document) -> None:
    for paragraph in _iter_all_paragraphs(document):
        if PLACEHOLDER_RE.search(paragraph.text):
            _replace_paragraph_text(paragraph, _remove_unresolved_placeholders(paragraph.text))


def _remove_unresolved_placeholders(text: str) -> str:
    text = PLACEHOLDER_RE.sub("", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def _remove_unresolved_placeholders_preserving_spacing(text: str) -> str:
    return PLACEHOLDER_RE.sub("", text)


def _replace_paragraph_with_items(
    paragraph: Paragraph,
    rendered: list[RenderedItem],
) -> None:
    if not rendered:
        _replace_paragraph_text(paragraph, "")
        return

    first = rendered[0]
    if isinstance(first, RenderedParagraph):
        _apply_rendered_paragraph(paragraph, first)
        anchor = paragraph._p
    else:
        _replace_paragraph_text(paragraph, "")
        anchor = _insert_item_after_anchor(paragraph, paragraph._p, first)

    for item in rendered[1:]:
        anchor = _insert_item_after_anchor(paragraph, anchor, item)


def _insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    return Paragraph(new_element, paragraph._parent)


def _insert_paragraph_after_anchor(paragraph: Paragraph, anchor) -> Paragraph:
    new_element = OxmlElement("w:p")
    anchor.addnext(new_element)
    return Paragraph(new_element, paragraph._parent)


def _insert_item_after_anchor(
    paragraph: Paragraph,
    anchor,
    item: RenderedItem,
):
    if isinstance(item, RenderedParagraph):
        inserted = _insert_paragraph_after_anchor(paragraph, anchor)
        _apply_rendered_paragraph(inserted, item)
        return inserted._p
    if isinstance(item, RenderedTable):
        return _insert_table_after_anchor(paragraph, anchor, item.block)
    if isinstance(item, RenderedFigure):
        return _insert_figure_after_anchor(paragraph, anchor, item.figure)
    return _insert_equation_after_anchor(paragraph, anchor, item.equation)


def _insert_table_after_anchor(paragraph: Paragraph, anchor, table_block: ContentBlock):
    if table_block.title:
        title = _insert_paragraph_after_anchor(paragraph, anchor)
        _replace_paragraph_text(title, table_block.title)
        anchor = title._p

    rows = _parse_table_rows(table_block.text)
    if not rows:
        inserted = _insert_paragraph_after_anchor(paragraph, anchor)
        _replace_paragraph_text(inserted, table_block.text)
        return inserted._p

    column_count = max(len(row) for row in rows)
    try:
        table = paragraph._parent.add_table(rows=len(rows), cols=column_count, width=Inches(6))
    except TypeError:
        table = paragraph._parent.add_table(rows=len(rows), cols=column_count)
    _safe_set_table_style(table, "Table Grid")
    for row_index, row_values in enumerate(rows):
        for column_index in range(column_count):
            value = row_values[column_index] if column_index < len(row_values) else ""
            table.cell(row_index, column_index).text = value
    anchor.addnext(table._tbl)
    return table._tbl


def _insert_figure_after_anchor(paragraph: Paragraph, anchor, figure: AssetItem):
    if _is_resolved_ocr_evidence_figure(figure):
        return anchor
    image_path = Path(figure.path) if figure.path else None
    caption = _figure_caption(figure)
    if image_path and _is_supported_existing_image(image_path) and not _is_ocr_evidence_figure(figure):
        picture = _insert_paragraph_after_anchor(paragraph, anchor)
        try:
            picture.add_run().add_picture(str(image_path), width=Inches(5.5))
            label = _insert_paragraph_after_anchor(paragraph, anchor)
            _replace_paragraph_text(label, f"[Figure inserted] {caption}".strip())
            return picture._p
        except Exception:
            _replace_paragraph_text(picture, _figure_review_text(figure))
            return picture._p

    review = _insert_paragraph_after_anchor(paragraph, anchor)
    _replace_paragraph_text(review, _figure_review_text(figure))
    return review._p


def _insert_equation_after_anchor(paragraph: Paragraph, anchor, equation: EquationItem):
    inserted = _insert_paragraph_after_anchor(paragraph, anchor)
    _apply_equation(inserted, equation)
    return inserted._p


def _apply_rendered_paragraph(paragraph: Paragraph, rendered: RenderedParagraph) -> None:
    _replace_paragraph_text(paragraph, rendered.text)
    if rendered.style:
        _safe_set_style(paragraph, rendered.style)
    if rendered.alignment is not None:
        paragraph.alignment = rendered.alignment
    if rendered.bold:
        for run in paragraph.runs:
            run.bold = True


def _replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    paragraph.clear()
    if text:
        paragraph.add_run(text)


def _build_fallback_document(document, model: ThesisModel) -> None:
    _add_title_block(document, model)
    _add_abstracts(document, model)
    _add_sections(document, model.sections)
    _add_tables(document, model.tables)
    _add_figures(document, model.figures)
    _add_equations(document, model.equations)
    _add_references(document, model.references)
    _add_appendices(document, model.appendices)


def _append_missing_content(document, model: ThesisModel, represented_placeholders: set[str]) -> None:
    _append_missing_abstracts(document, model, represented_placeholders)
    for token in ("BODY", "TABLES", "FIGURES", "EQUATIONS", "REFERENCES", "APPENDICES"):
        if token in represented_placeholders:
            continue
        _append_block(document, token, model)


def _append_missing_abstracts(document, model: ThesisModel, represented_placeholders: set[str]) -> None:
    abstract_cn = _front_matter_value(model, "chinese_abstract", "abstract_cn", "cn_abstract")
    abstract_en = _front_matter_value(model, "english_abstract", "abstract_en", "en_abstract")
    if abstract_cn and "ABSTRACT_CN" not in represented_placeholders:
        _add_front_matter_heading(document, "摘    要")
        _add_text_paragraphs(document, abstract_cn)
        _add_keywords(
            document,
            "关键词：",
            _front_matter_value(model, "keywords_cn", "chinese_keywords", "cn_keywords", "keywords"),
        )
    if abstract_en and "ABSTRACT_EN" not in represented_placeholders:
        if abstract_cn:
            document.add_page_break()
        _add_front_matter_heading(document, "Abstract")
        _add_text_paragraphs(document, abstract_en)
        _add_keywords(
            document,
            "Key Words: ",
            _front_matter_value(model, "keywords_en", "english_keywords", "en_keywords", "keywords"),
        )


def _append_block(document, token: str, model: ThesisModel) -> None:
    if token == "BODY":
        _add_sections(document, model.sections)
    elif token == "TABLES":
        _add_tables(document, model.tables)
    elif token == "FIGURES":
        _add_figures(document, model.figures)
    elif token == "EQUATIONS":
        _add_equations(document, model.equations)
    elif token == "REFERENCES":
        _add_references(document, model.references)
    elif token == "APPENDICES":
        _add_appendices(document, model.appendices)


def _add_title_block(document, model: ThesisModel) -> None:
    metadata = model.metadata
    title = metadata.title_cn or metadata.title_en or "Untitled Thesis"
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(16)

    if metadata.title_en and metadata.title_en != title:
        english_title = document.add_paragraph()
        english_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = english_title.add_run(metadata.title_en)
        run.bold = True
        run.font.size = Pt(14)

    metadata_lines = [
        ("Student Name", metadata.student_name),
        ("Student ID", metadata.student_id),
        ("College", metadata.college),
        ("Major", metadata.major),
        ("Advisor", metadata.advisor),
        ("Date", metadata.date),
        ("Classification", metadata.classification),
        ("Unit Code", metadata.unit_code),
    ]
    for label, value in metadata_lines:
        if value:
            document.add_paragraph(f"{label}: {value}")


def _add_abstracts(document, model: ThesisModel) -> None:
    abstract_cn = _front_matter_value(model, "chinese_abstract", "abstract_cn", "cn_abstract")
    abstract_en = _front_matter_value(model, "english_abstract", "abstract_en", "en_abstract")
    if abstract_cn:
        _add_front_matter_heading(document, "摘    要")
        _add_text_paragraphs(document, abstract_cn)
        _add_keywords(
            document,
            "关键词：",
            _front_matter_value(model, "keywords_cn", "chinese_keywords", "cn_keywords", "keywords"),
        )
    if abstract_en:
        if abstract_cn:
            document.add_page_break()
        _add_front_matter_heading(document, "Abstract")
        _add_text_paragraphs(document, abstract_en)
        _add_keywords(
            document,
            "Key Words: ",
            _front_matter_value(model, "keywords_en", "english_keywords", "en_keywords", "keywords"),
        )


def _add_sections(document, sections: Iterable[ContentBlock]) -> None:
    for item in _render_sections(sections):
        paragraph = document.add_paragraph()
        _apply_rendered_paragraph(paragraph, item)


def _add_tables(document, tables: Iterable[ContentBlock]) -> None:
    for table_block in tables:
        if table_block.title:
            document.add_paragraph(table_block.title)
        rows = _parse_table_rows(table_block.text)
        if not rows:
            if table_block.text:
                _add_text_paragraphs(document, table_block.text)
            continue
        column_count = max(len(row) for row in rows)
        table = document.add_table(rows=len(rows), cols=column_count)
        _safe_set_table_style(table, "Table Grid")
        for row_index, row_values in enumerate(rows):
            for column_index in range(column_count):
                value = row_values[column_index] if column_index < len(row_values) else ""
                table.cell(row_index, column_index).text = value


def _add_figures(document, figures: Iterable[AssetItem]) -> None:
    for figure in figures:
        if _is_resolved_ocr_evidence_figure(figure):
            continue
        image_path = Path(figure.path) if figure.path else None
        caption = _figure_caption(figure)
        if image_path and _is_supported_existing_image(image_path) and not _is_ocr_evidence_figure(figure):
            picture = document.add_paragraph()
            try:
                picture.add_run().add_picture(str(image_path), width=Inches(5.5))
                document.add_paragraph(f"[Figure inserted] {caption}".strip())
            except Exception:
                _replace_paragraph_text(picture, _figure_review_text(figure))
        else:
            document.add_paragraph(_figure_review_text(figure))


def _add_equations(document, equations: Iterable[EquationItem]) -> None:
    for item in _render_equations(equations):
        paragraph = document.add_paragraph()
        if isinstance(item, RenderedEquation):
            _apply_equation(paragraph, item.equation)
        else:
            _apply_rendered_paragraph(paragraph, item)


def _add_references(document, references: Iterable[ContentBlock]) -> None:
    references = list(references)
    if not references:
        return
    _add_heading(document, "References", level=1)
    for item in _render_references(references):
        document.add_paragraph(item.text)


def _add_appendices(document, appendices: Iterable[ContentBlock]) -> None:
    for item in _render_appendices(appendices):
        paragraph = document.add_paragraph()
        _apply_rendered_paragraph(paragraph, item)


def _render_sections(sections: Iterable[ContentBlock]) -> list[RenderedParagraph]:
    rendered: list[RenderedParagraph] = []
    for section in sections:
        if section.title:
            rendered.append(RenderedParagraph(section.title, style=_heading_style(section.level)))
        rendered.extend(RenderedParagraph(text) for text in _split_paragraph_text(section.text))
    return rendered


def _render_references(references: Iterable[ContentBlock]) -> list[RenderedParagraph]:
    rendered: list[RenderedParagraph] = []
    for reference in references:
        text = reference.text or reference.title
        rendered.extend(RenderedParagraph(item) for item in _split_paragraph_text(text))
    return rendered


def _render_tables_as_paragraphs(tables: Iterable[ContentBlock]) -> list[RenderedParagraph]:
    rendered: list[RenderedParagraph] = []
    for table in tables:
        if table.title:
            rendered.append(RenderedParagraph(table.title))
        rows = _parse_table_rows(table.text)
        if rows:
            rendered.extend(RenderedParagraph(" | ".join(row)) for row in rows)
        else:
            rendered.extend(RenderedParagraph(text) for text in _split_paragraph_text(table.text))
    return rendered


def _render_figures_as_review_paragraphs(figures: Iterable[AssetItem]) -> list[RenderedParagraph]:
    return [RenderedParagraph(_figure_review_text(figure)) for figure in figures]


def _render_equations(equations: Iterable[EquationItem]) -> list[RenderedEquation | RenderedParagraph]:
    rendered: list[RenderedEquation | RenderedParagraph] = []
    for equation in equations:
        if equation.omml.strip() or _has_preserved_ole_equation(equation) or equation.preview_path:
            rendered.append(RenderedEquation(equation))
        else:
            rendered.append(RenderedParagraph(_equation_review_text(equation)))
    return rendered


def _apply_equation(paragraph: Paragraph, equation: EquationItem) -> None:
    omml = equation.omml.strip()
    if omml:
        try:
            paragraph.clear()
            paragraph._p.append(parse_xml(omml))
            if equation.number:
                paragraph.add_run(f" {equation.number}")
            return
        except Exception:
            pass
    if _has_preserved_ole_equation(equation):
        _replace_paragraph_text(paragraph, _equation_object_token(equation))
        return
    preview_path = Path(equation.preview_path) if equation.preview_path else None
    if preview_path is not None and _is_supported_existing_image(preview_path):
        try:
            paragraph.clear()
            paragraph.add_run().add_picture(str(preview_path), width=Inches(4.8))
            label = _insert_paragraph_after(paragraph)
            _replace_paragraph_text(label, f"[Equation preview inserted] {_equation_caption(equation)}".strip())
            return
        except Exception:
            pass
    _replace_paragraph_text(paragraph, _equation_review_text(equation))


def finalize_equation_objects(docx_path: Path, equations: Iterable[EquationItem]) -> None:
    """Patch preserved Word/MathType OLE equations into a saved DOCX package."""
    equations_by_token = {
        _equation_object_token(equation): equation
        for equation in equations
        if _has_preserved_ole_equation(equation)
    }
    if not equations_by_token:
        return

    docx_path = Path(docx_path)
    if not docx_path.exists():
        return

    patched_path = docx_path.with_name(f"{docx_path.stem}.equation-objects{docx_path.suffix}")
    additional_parts: dict[str, bytes] = {}
    try:
        with zipfile.ZipFile(docx_path, "r") as package:
            package_names = set(package.namelist())
            document_root = etree.fromstring(package.read("word/document.xml"))
            rels_xml = (
                package.read("word/_rels/document.xml.rels")
                if "word/_rels/document.xml.rels" in package_names
                else b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
            )
            rels_root = etree.fromstring(rels_xml)
            content_types_root = etree.fromstring(package.read("[Content_Types].xml"))

            next_rel_index = _next_relationship_index(rels_root)
            changed = False
            for token, equation in equations_by_token.items():
                paragraph = _find_token_paragraph(document_root, token)
                if paragraph is None:
                    continue
                try:
                    next_rel_index = _patch_equation_object(
                        paragraph,
                        equation,
                        rels_root,
                        content_types_root,
                        package_names | set(additional_parts),
                        additional_parts,
                        next_rel_index,
                    )
                except Exception:
                    _replace_package_paragraph_text(paragraph, _equation_review_text(equation))
                changed = True

            if not changed:
                return

            document_xml = _serialize_xml(document_root)
            rels_xml = _serialize_xml(rels_root)
            content_types_xml = _serialize_xml(content_types_root)

            with zipfile.ZipFile(patched_path, "w") as target:
                for info in package.infolist():
                    if info.filename == "word/document.xml":
                        target.writestr(info, document_xml)
                    elif info.filename == "word/_rels/document.xml.rels":
                        target.writestr(info, rels_xml)
                    elif info.filename == "[Content_Types].xml":
                        target.writestr(info, content_types_xml)
                    elif info.filename not in additional_parts:
                        target.writestr(info, package.read(info.filename))
                if "word/_rels/document.xml.rels" not in package_names:
                    target.writestr("word/_rels/document.xml.rels", rels_xml)
                for part_name, payload in additional_parts.items():
                    target.writestr(part_name, payload)
        patched_path.replace(docx_path)
    finally:
        if patched_path.exists():
            patched_path.unlink()


def _patch_equation_object(
    paragraph,
    equation: EquationItem,
    rels_root,
    content_types_root,
    package_names: set[str],
    additional_parts: dict[str, bytes],
    next_rel_index: int,
) -> int:
    object_path = Path(equation.object_path)
    object_xml = equation.object_xml.strip()
    if not object_path.exists() or not object_xml:
        _replace_package_paragraph_text(paragraph, _equation_review_text(equation))
        return next_rel_index

    object_node = etree.fromstring(object_xml.encode("utf-8"))
    if etree.QName(object_node).localname == "r":
        run_node = object_node
    else:
        run_node = etree.Element(f"{{{W_NS}}}r")
        run_node.append(object_node)

    object_rel_id = f"rIdBuaaEquationObject{next_rel_index}"
    next_rel_index += 1
    stem = _safe_equation_package_stem(equation)
    object_suffix = object_path.suffix.lower() or ".bin"
    object_part_name = _unique_package_part_name(
        package_names | set(additional_parts),
        f"word/embeddings/{stem}{object_suffix}",
    )
    additional_parts[object_part_name] = object_path.read_bytes()
    _add_relationship(
        rels_root,
        object_rel_id,
        OLE_OBJECT_REL_TYPE,
        object_part_name.removeprefix("word/"),
    )
    _ensure_default_content_type(content_types_root, object_suffix, OLE_OBJECT_CONTENT_TYPE)
    for ole_node in run_node.xpath(".//*[local-name()='OLEObject']"):
        ole_node.set(f"{{{R_NS}}}id", object_rel_id)

    preview_path = Path(equation.preview_path) if equation.preview_path else None
    if preview_path is not None and preview_path.exists():
        image_rel_id = f"rIdBuaaEquationPreview{next_rel_index}"
        next_rel_index += 1
        image_suffix = preview_path.suffix.lower() or ".png"
        image_part_name = _unique_package_part_name(
            package_names | set(additional_parts),
            f"word/media/{stem}{image_suffix}",
        )
        additional_parts[image_part_name] = preview_path.read_bytes()
        _add_relationship(
            rels_root,
            image_rel_id,
            IMAGE_REL_TYPE,
            image_part_name.removeprefix("word/"),
        )
        _ensure_default_content_type(
            content_types_root,
            image_suffix,
            IMAGE_CONTENT_TYPES.get(image_suffix, "application/octet-stream"),
        )
        for image_node in run_node.xpath(".//*[local-name()='imagedata']"):
            image_node.set(f"{{{R_NS}}}id", image_rel_id)

    _replace_package_paragraph_with_run(paragraph, run_node)
    return next_rel_index


def _has_preserved_ole_equation(equation: EquationItem) -> bool:
    return bool(equation.object_xml.strip() and equation.object_path and Path(equation.object_path).exists())


def _equation_object_token(equation: EquationItem) -> str:
    return f"{EQUATION_OBJECT_TOKEN_PREFIX}{_safe_equation_package_stem(equation)}__"


def _safe_equation_package_stem(equation: EquationItem) -> str:
    value = equation.id or equation.text or "equation"
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", value).strip("._")
    return f"buaa-equation-{safe or 'equation'}"


def _find_token_paragraph(root, token: str):
    for text_node in root.xpath(".//*[local-name()='t']"):
        if text_node.text != token:
            continue
        paragraph = text_node
        while paragraph is not None and etree.QName(paragraph).localname != "p":
            paragraph = paragraph.getparent()
        return paragraph
    return None


def _replace_package_paragraph_with_run(paragraph, run_node) -> None:
    _clear_package_paragraph(paragraph)
    paragraph.append(run_node)


def _replace_package_paragraph_text(paragraph, text: str) -> None:
    _clear_package_paragraph(paragraph)
    run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
    text_node = etree.SubElement(run, f"{{{W_NS}}}t")
    text_node.text = text


def _clear_package_paragraph(paragraph) -> None:
    for child in list(paragraph):
        if child.tag != f"{{{W_NS}}}pPr":
            paragraph.remove(child)


def _add_relationship(rels_root, rel_id: str, relationship_type: str, target: str) -> None:
    relationship = etree.SubElement(rels_root, f"{{{REL_NS}}}Relationship")
    relationship.set("Id", rel_id)
    relationship.set("Type", relationship_type)
    relationship.set("Target", target)


def _next_relationship_index(rels_root) -> int:
    highest = 0
    for relationship in rels_root.xpath(".//*[local-name()='Relationship']"):
        rel_id = relationship.get("Id") or ""
        match = re.search(r"(\d+)$", rel_id)
        if match:
            highest = max(highest, int(match.group(1)))
    return highest + 1


def _ensure_default_content_type(content_types_root, suffix: str, content_type: str) -> None:
    extension = suffix.lower().lstrip(".")
    for node in content_types_root.xpath(".//*[local-name()='Default']"):
        if (node.get("Extension") or "").lower() == extension:
            return
    default = etree.SubElement(content_types_root, f"{{{CONTENT_TYPES_NS}}}Default")
    default.set("Extension", extension)
    default.set("ContentType", content_type)


def _unique_package_part_name(existing_names: set[str], desired_name: str) -> str:
    path = Path(desired_name)
    parent = path.parent.as_posix()
    suffix = path.suffix
    stem = path.stem
    candidate = desired_name
    counter = 2
    while candidate in existing_names:
        candidate = f"{parent}/{stem}-{counter}{suffix}"
        counter += 1
    return candidate


def _serialize_xml(root) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _equation_review_text(equation: EquationItem) -> str:
    content = equation.latex or equation.text or "manual conversion required"
    if equation.preview_path:
        content = f"{content} preview={Path(equation.preview_path).name}"
    number = f" {equation.number}" if equation.number else ""
    prefix = "[Equation requires review]" if equation.requires_review else "[Equation]"
    return f"{prefix}{number} {content}".strip()


def _equation_caption(equation: EquationItem) -> str:
    return equation.text or equation.number or equation.id


def _render_appendices(appendices: Iterable[ContentBlock]) -> list[RenderedParagraph]:
    rendered: list[RenderedParagraph] = []
    for appendix in appendices:
        if appendix.title:
            rendered.append(RenderedParagraph(appendix.title, style="Heading 1"))
        rendered.extend(RenderedParagraph(text) for text in _split_paragraph_text(appendix.text))
    return rendered


def _split_paragraph_text(text: str) -> list[str]:
    return [line.strip() for line in str(text or "").splitlines() if line.strip()]


def _parse_table_rows(text: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in str(text or "").splitlines():
        if not line.strip():
            continue
        rows.append([cell.strip() for cell in line.split("\t")])
    return rows


def _figure_caption(figure: AssetItem) -> str:
    parts = [figure.caption, figure.path]
    return " ".join(part for part in parts if part)


def _figure_review_text(figure: AssetItem) -> str:
    if _is_ocr_evidence_figure(figure):
        return f"[OCR evidence requires transcription] {_figure_caption(figure)}".strip()
    return f"[Figure requires review] {_figure_caption(figure)}".strip()


def _is_supported_existing_image(path: Path) -> bool:
    return path.exists() and path.suffix.lower() in SUPPORTED_IMAGE_SUFFIXES


def _is_ocr_evidence_figure(figure: AssetItem) -> bool:
    return str(figure.type or "").strip().lower() in OCR_EVIDENCE_FIGURE_TYPES


def _is_resolved_ocr_evidence_figure(figure: AssetItem) -> bool:
    return _is_ocr_evidence_figure(figure) and not figure.requires_review


def _heading_style(level: int) -> str:
    return "Heading 1" if level <= 1 else "Heading 2"


def _add_heading(document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.bold = True


def _add_front_matter_heading(document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True


def _add_keywords(document, label: str, keywords: str) -> None:
    value = str(keywords or "").strip()
    if value:
        document.add_paragraph(f"{label}{value}")


def _add_text_paragraphs(document, text: str) -> None:
    for paragraph_text in _split_paragraph_text(text):
        document.add_paragraph(paragraph_text)


def _safe_set_style(paragraph: Paragraph, style_name: str) -> None:
    try:
        paragraph.style = style_name
    except KeyError:
        return


def _safe_set_table_style(table, style_name: str) -> None:
    try:
        table.style = style_name
    except KeyError:
        return


def _iter_all_paragraphs(document) -> Iterable[Paragraph]:
    # Task 6 intentionally limits placeholder scanning to the document body.
    # Accessing missing header/footer objects through python-docx creates parts on save.
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        yield from _iter_table_paragraphs(table)


def _iter_table_paragraphs(table) -> Iterable[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested_table in cell.tables:
                yield from _iter_table_paragraphs(nested_table)


def _clear_body_preserving_sections(document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _apply_conservative_formatting(document, *, apply_page_setup: bool) -> None:
    if apply_page_setup and len(document.sections) == 1:
        section = document.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(30)
        section.bottom_margin = Mm(25)
        section.left_margin = Mm(30)
        section.right_margin = Mm(20)

    try:
        normal = document.styles["Normal"]
    except KeyError:
        return

    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    rpr = normal._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "SimSun")
