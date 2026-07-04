from __future__ import annotations

import copy
import re
from dataclasses import replace
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from buaa_thesis_kit.models import AssetItem, ContentBlock, Metadata, ThesisModel
from buaa_thesis_kit.template_fill import (
    _add_abstracts,
    _add_appendices,
    _add_equations,
    _add_references,
    _add_sections,
    _add_table_of_contents,
    _add_tables,
    _apply_conservative_formatting,
    finalize_equation_objects,
)


SPINE_MARKER = "书脊"
SAMPLE_BODY_START = "论文封面书脊"
PDF_EXTRACTED_TEXT_TITLE = "PDF Extracted Text"
SUPPORTED_INLINE_IMAGE_SUFFIXES = {".bmp", ".gif", ".jpg", ".jpeg", ".png", ".tif", ".tiff"}
FIGURE_MARKER_RE = re.compile(
    "^(?:(?:\u56fe)|fig(?:ure)?\\.?)\\s*(\\d+(?:[.\\-]\\d+)*)",
    flags=re.IGNORECASE,
)
FRONT_MATTER_SECTION_TITLES = {
    "本科毕业设计（论文）任务书",
    "本科毕业设计(论文)任务书",
    "毕业设计（论文）任务书",
    "毕业设计(论文)任务书",
    "任务书",
    "本人声明",
    "摘要",
    "Abstract",
    "目录",
}
COVER_METADATA_LABELS = (
    "学院",
    "院（系）名称",
    "专业",
    "专业名称",
    "学生姓名",
    "指导教师",
    "学号",
    "题目",
)


def render_editable_buaa_docx(template_path: Path, model: ThesisModel, output_path: Path) -> None:
    """Render an editable BUAA thesis DOCX by reusing the official Word template."""
    document = Document(str(template_path))
    _apply_conservative_formatting(document, apply_page_setup=False)
    _replace_cover_fields(document, model.metadata)
    _compact_cover_spacing(document)
    _trim_template_after_cover(document)
    _append_spine_page(document, model.metadata)
    _append_model_content(document, model)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    finalize_equation_objects(output, model.equations)


def _replace_cover_fields(document, metadata: Metadata) -> None:
    title = _metadata_value(metadata.title_cn or metadata.title_en or "Untitled Thesis")
    for paragraph in _cover_paragraphs(document):
        text = paragraph.text
        stripped = text.strip()
        if "单位代码" in text:
            _replace_paragraph_text_preserving_style(
                paragraph,
                f"                                                    单位代码       {_metadata_value(metadata.unit_code or '10006')}        ",
            )
        elif "学" in text and "号" in text and "学号" not in text:
            _replace_paragraph_text_preserving_style(
                paragraph,
                f"学    号     {_metadata_value(metadata.student_id)}       ",
            )
        elif "分类号" in text:
            _replace_paragraph_text_preserving_style(
                paragraph,
                f"分类号    {_metadata_value(metadata.classification)}         ",
            )
        elif stripped in {"（题目）", "(题目)"}:
            _replace_paragraph_text_preserving_style(paragraph, title)
        elif re.fullmatch(r"\d{4}年\d{1,2}月", stripped):
            _replace_paragraph_text_preserving_style(
                paragraph,
                _metadata_value(metadata.date),
            )

    if document.tables:
        _replace_cover_table(document.tables[0], metadata)


def _cover_paragraphs(document):
    for paragraph in document.paragraphs:
        if SAMPLE_BODY_START in paragraph.text:
            break
        yield paragraph


def _replace_cover_table(table, metadata: Metadata) -> None:
    values = {
        "学院名称": metadata.college,
        "专业名称": metadata.major,
        "学生姓名": metadata.student_name,
        "指导教师": metadata.advisor,
    }
    for row in table.rows:
        if len(row.cells) < 2:
            continue
        label = re.sub(r"\s+", "", row.cells[0].text)
        for key, value in values.items():
            if key in label:
                _replace_cell_text_preserving_style(
                    row.cells[1],
                    _metadata_value(value),
                    font_size=_cover_table_value_font_size(value),
                )
                break


def _trim_template_after_cover(document) -> None:
    body = document._body._element
    remove = False
    for child in list(body):
        if child.tag == qn("w:p") and SAMPLE_BODY_START in _xml_text(child):
            _remove_trailing_blank_paragraphs(body, child)
            remove = True
        if remove and child.tag != qn("w:sectPr"):
            body.remove(child)


def _compact_cover_spacing(document) -> None:
    body = document._body._element
    consecutive_blanks = 0
    for child in list(body):
        if child.tag == qn("w:p") and SAMPLE_BODY_START in _xml_text(child):
            break
        if child.tag == qn("w:p") and _is_removable_blank_paragraph(child):
            consecutive_blanks += 1
            if consecutive_blanks > 2:
                body.remove(child)
            continue
        consecutive_blanks = 0


def _remove_trailing_blank_paragraphs(body, marker) -> None:
    current = marker.getprevious()
    while current is not None and current.tag == qn("w:p") and _is_removable_blank_paragraph(current):
        previous = current.getprevious()
        body.remove(current)
        current = previous


def _is_removable_blank_paragraph(paragraph_element) -> bool:
    if _xml_text(paragraph_element).strip():
        return False
    protected_tags = {qn("w:drawing"), qn("w:pict")}
    return not any(node.tag in protected_tags for node in paragraph_element.iter())


def _append_spine_page(document, metadata: Metadata) -> None:
    document.add_page_break()
    marker = document.add_paragraph(SPINE_MARKER)
    marker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_runs(marker, bold=True, size=14)

    for value in _spine_values(metadata):
        paragraph = document.add_paragraph(value)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _format_runs(paragraph, size=12)


def _append_model_content(document, model: ThesisModel) -> None:
    body_sections = _body_sections(model.sections)
    document.add_page_break()
    _append_task_book_page(document, model.metadata)
    document.add_page_break()
    _append_declaration_page(document, model.metadata)
    document.add_page_break()
    _add_abstracts(document, model)
    _add_table_of_contents(document, body_sections)
    if body_sections:
        document.add_page_break()
    _add_sections_with_page_breaks(
        document,
        body_sections,
        model.figures,
    )
    _add_tables(document, model.tables)
    _add_equations(document, _trusted_editable_equations(model))
    _add_references(document, model.references)
    _add_appendices(document, model.appendices)


def _append_task_book_page(document, metadata: Metadata) -> None:
    heading = document.add_paragraph("本科毕业设计（论文）任务书")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_runs(heading, bold=True, size=14)
    for label, value in [
        ("题目", metadata.title_cn or metadata.title_en),
        ("学生姓名", metadata.student_name),
        ("学号", metadata.student_id),
        ("学院", metadata.college),
        ("专业", metadata.major),
        ("指导教师", metadata.advisor),
        ("日期", metadata.date),
    ]:
        if _metadata_value(value):
            document.add_paragraph(f"{label}：{_metadata_value(value)}")
    document.add_paragraph("任务内容、进度安排和指导记录请以学校原始任务书为准；本页由规范化流水线按模板生成，需人工复核。")


def _append_declaration_page(document, metadata: Metadata) -> None:
    heading = document.add_paragraph("本人声明")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_runs(heading, bold=True, size=14)
    document.add_paragraph(
        "本人郑重声明：所提交的毕业设计（论文）是在指导教师的指导下独立完成的。"
        "除文中已经注明引用的内容外，本论文不包含他人已经发表或撰写过的研究成果。"
    )
    if _metadata_value(metadata.student_name):
        document.add_paragraph(f"学生签名：{_metadata_value(metadata.student_name)}")


def _trusted_editable_equations(model: ThesisModel):
    return [
        equation
        for equation in model.equations
        if equation.omml.strip()
    ]


def _body_sections(sections: list[ContentBlock]) -> list[ContentBlock]:
    return [
        section
        for section in _sections_without_pdf_placeholder_heading(sections)
        if not _is_front_matter_section(section)
    ]


def _is_front_matter_section(section: ContentBlock) -> bool:
    title = _metadata_value(section.title)
    compact_title = re.sub(r"\s+", "", title)
    if compact_title in {re.sub(r"\s+", "", item) for item in FRONT_MATTER_SECTION_TITLES}:
        return True
    lines = [line.strip() for line in str(section.text or "").splitlines() if line.strip()]
    if not title and 1 <= len(lines) <= 10:
        hit_count = sum(1 for line in lines if any(label in line for label in COVER_METADATA_LABELS))
        return hit_count >= 2
    return False


def _add_sections_with_page_breaks(
    document,
    sections: list[ContentBlock],
    figures: list[AssetItem] | None = None,
) -> set[str]:
    rendered_figure_ids: set[str] = set()
    first = True
    for section in sections:
        if not first and _section_starts_new_page(section):
            document.add_page_break()
        _add_section_with_inline_figures(document, section, figures or [], rendered_figure_ids)
        first = False
    return rendered_figure_ids


def _add_section_with_inline_figures(
    document,
    section: ContentBlock,
    figures: list[AssetItem],
    rendered_figure_ids: set[str],
) -> None:
    if section.title:
        _add_sections(document, [replace(section, text="")])
    for text in _split_section_text(section.text):
        figure = _matching_captioned_figure(text, figures, rendered_figure_ids)
        if figure is not None and _add_inline_figure(document, figure):
            rendered_figure_ids.add(figure.id)
        document.add_paragraph(text)


def _matching_captioned_figure(
    paragraph_text: str,
    figures: list[AssetItem],
    rendered_figure_ids: set[str],
) -> AssetItem | None:
    paragraph_key = _caption_key(paragraph_text)
    if not paragraph_key:
        return None
    paragraph_marker = _figure_marker_key(paragraph_text)
    for figure in figures:
        if figure.id in rendered_figure_ids:
            continue
        if _caption_key(figure.caption) == paragraph_key:
            return figure
        if paragraph_marker and _figure_marker_key(figure.caption) == paragraph_marker:
            return figure
    return None


def _add_inline_figure(document, figure: AssetItem) -> bool:
    image_path = Path(figure.path) if figure.path else None
    if image_path is None or not _is_supported_inline_image(image_path):
        return False
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        paragraph.add_run().add_picture(str(image_path), width=Inches(5.5))
    except Exception:
        parent = paragraph._element.getparent()
        if parent is not None:
            parent.remove(paragraph._element)
        return False
    return True


def _is_supported_inline_image(path: Path) -> bool:
    return path.exists() and path.is_file() and path.suffix.lower() in SUPPORTED_INLINE_IMAGE_SUFFIXES


def _split_section_text(text: str) -> list[str]:
    return [line.strip() for line in str(text or "").splitlines() if line.strip()]


def _caption_key(text: str) -> str:
    return re.sub(r"\s+", "", str(text or "")).casefold()


def _figure_marker_key(text: str) -> str:
    match = FIGURE_MARKER_RE.match(str(text or "").strip())
    if not match:
        return ""
    return match.group(1).replace("-", ".").casefold()


def _section_starts_new_page(section: ContentBlock) -> bool:
    title = _metadata_value(section.title)
    if title in {"本人声明", "摘要", "Abstract"}:
        return True
    return section.level <= 1 and bool(
        re.match(r"^\d+\s+\S+", title)
        or re.match(r"^第[一二三四五六七八九十百]+章\s+\S+$", title)
    )


def _sections_without_pdf_placeholder_heading(sections: list[ContentBlock]) -> list[ContentBlock]:
    rendered: list[ContentBlock] = []
    for section in sections:
        if section.title == PDF_EXTRACTED_TEXT_TITLE:
            rendered.append(replace(section, title=""))
        else:
            rendered.append(section)
    return rendered


def _spine_values(metadata: Metadata) -> list[str]:
    return [
        value
        for value in [
            metadata.title_cn or metadata.title_en,
            metadata.student_name,
            metadata.college,
            metadata.major,
            metadata.date,
        ]
        if _metadata_value(value)
    ]


def _cover_table_value_font_size(value: str) -> float | None:
    length = len(_metadata_value(value))
    if length >= 16:
        return 10.5
    if length >= 10:
        return 12
    return None


def _replace_cell_text_preserving_style(
    cell,
    text: str,
    *,
    font_size: float | None = None,
) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _replace_paragraph_text_preserving_style(paragraph, text)
    if font_size is not None:
        _format_runs(paragraph, size=font_size)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def _replace_paragraph_text_preserving_style(paragraph, text: str) -> None:
    first_run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        first_run_properties = copy.deepcopy(paragraph.runs[0]._r.rPr)
    paragraph.clear()
    if not text:
        return
    run = paragraph.add_run(text)
    if first_run_properties is not None:
        run._r.insert(0, first_run_properties)


def _format_runs(paragraph, *, bold: bool = False, size: float = 12) -> None:
    for run in paragraph.runs:
        run.bold = bold
        run.font.size = Pt(size)


def _metadata_value(value: str) -> str:
    return str(value or "").strip()


def _xml_text(element) -> str:
    return "".join(node.text or "" for node in element.iter() if node.tag == qn("w:t"))


__all__ = ["render_editable_buaa_docx"]
