from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document

from buaa_thesis_kit.graph import GraphState, NodeResult
from buaa_thesis_kit.models import Metadata
from buaa_thesis_kit.template_fill import fill_word_template


SPINE_MARKERS = ("Book Spine", "\u4e66\u810a")
SPINE_REQUIRED_FIELDS = ("title_cn", "student_name", "college", "major", "date")


def inspect_source_features(state: GraphState) -> NodeResult:
    state.source_features["has_spine"] = _source_has_spine(state.extraction_source)
    return NodeResult(next_node="diagnose_compliance")


def diagnose_compliance(state: GraphState) -> NodeResult:
    if state.model.status == "failed":
        details = "; ".join(state.model.extraction_warnings) if state.model.extraction_warnings else "no details"
        state.blocking_items.append(f"Extraction failed: {details}")
        return NodeResult(next_node="finalize_output")

    if not state.source_features.get("has_spine", False):
        _append_once(state.findings, "missing_spine")
    return NodeResult(next_node="plan_minimal_fixes")


def plan_minimal_fixes(state: GraphState) -> NodeResult:
    if "missing_spine" in state.findings:
        missing_fields = _missing_spine_fields(state.model.metadata)
        if missing_fields:
            _append_once(
                state.manual_review,
                f"spine_metadata_missing: {', '.join(missing_fields)}",
            )
        _append_once(state.repair_actions, "insert_spine")
        _append_once(state.notes, "Finding missing_spine planned as repair action insert_spine.")
    return NodeResult(next_node="apply_word_fixes")


def apply_word_fixes(state: GraphState) -> NodeResult:
    state.work_dir.mkdir(parents=True, exist_ok=True)
    repaired_docx = state.work_dir / "repaired.docx"

    if _can_repair_docx_source(state):
        shutil.copy2(Path(state.extraction_source), repaired_docx)
        document = Document(str(repaired_docx))
    else:
        if state.template_path is None:
            state.blocking_items.append("Word repair failed: missing template for non-DOCX source.")
            return NodeResult(next_node="export_pdf")
        fill_word_template(state.template_path, state.model, repaired_docx)
        document = Document(str(repaired_docx))

    if "insert_spine" in state.repair_actions and not _document_has_spine(document):
        _append_spine_page(document, state.model.metadata)
        _append_once(state.applied_repairs, "insert_spine")
        _append_once(state.notes, "Applied repair insert_spine for missing_spine.")

    document.save(str(repaired_docx))
    state.authoritative_docx = repaired_docx
    state.outputs["word"] = "pass" if repaired_docx.exists() and repaired_docx.stat().st_size > 0 else "failed"
    return NodeResult(next_node="export_pdf")


def visual_compare(state: GraphState) -> NodeResult:
    if "missing_spine" in state.findings and "insert_spine" not in state.applied_repairs:
        _append_once(state.blocking_items, "spine_visual_mismatch: insert_spine repair was not applied.")
    return NodeResult(next_node="decide")


def decide(state: GraphState) -> NodeResult:
    if state.blocking_items:
        return NodeResult(next_node="finalize_output")
    unresolved = [
        finding
        for finding in state.findings
        if finding == "missing_spine" and "insert_spine" not in state.applied_repairs
    ]
    if unresolved:
        return NodeResult(next_node="revise_plan")
    return NodeResult(next_node="finalize_output")


def _source_has_spine(source: Path | None) -> bool:
    if source is None or source.suffix.lower() != ".docx" or not source.exists():
        return False
    try:
        document = Document(str(source))
    except Exception:
        return False
    return _document_has_spine(document)


def _document_has_spine(document) -> bool:
    return any(_text_has_spine_marker(paragraph.text) for paragraph in document.paragraphs)


def _text_has_spine_marker(text: str) -> bool:
    return any(marker in str(text) for marker in SPINE_MARKERS)


def _missing_spine_fields(metadata: Metadata) -> list[str]:
    missing: list[str] = []
    if not (metadata.title_cn or metadata.title_en):
        missing.append("title")
    for field in SPINE_REQUIRED_FIELDS:
        if field == "title_cn":
            continue
        if not getattr(metadata, field):
            missing.append(field)
    return missing


def _can_repair_docx_source(state: GraphState) -> bool:
    source = state.extraction_source
    return state.source_kind == "docx" and source is not None and source.suffix.lower() == ".docx" and source.exists()


def _append_spine_page(document, metadata: Metadata) -> None:
    document.add_page_break()
    document.add_paragraph("Book Spine")
    document.add_paragraph(_metadata_value(metadata.title_cn or metadata.title_en))
    document.add_paragraph(_metadata_value(metadata.student_name))
    document.add_paragraph(_metadata_value(metadata.college))
    document.add_paragraph(_metadata_value(metadata.major))
    document.add_paragraph(_metadata_value(metadata.date))


def _metadata_value(value: str) -> str:
    return str(value or "").strip()


def _append_once(items: list[str], item: str) -> None:
    if item not in items:
        items.append(item)


__all__ = [
    "apply_word_fixes",
    "decide",
    "diagnose_compliance",
    "inspect_source_features",
    "plan_minimal_fixes",
    "visual_compare",
]
