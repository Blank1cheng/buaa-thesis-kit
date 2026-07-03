from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


PAGE_IMAGE_INSET_INCHES = 0.35


def render_pdf_pages_as_docx(pdf_path: Path, output_path: Path, work_dir: Path, scale: float = 2.0) -> int:
    """Render every source PDF page as a full-page image inside an authoritative DOCX."""
    source = Path(pdf_path)
    destination = Path(output_path)
    work = Path(work_dir)
    image_dir = work / "pdf-page-render"
    image_dir.mkdir(parents=True, exist_ok=True)
    destination.parent.mkdir(parents=True, exist_ok=True)

    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError(f"PyMuPDF unavailable for PDF visual rendering: {exc}") from exc

    document = fitz.open(str(source))
    try:
        if document.page_count == 0:
            raise RuntimeError(f"PDF visual rendering failed: source has no pages: {source}")

        first_page = document.load_page(0)
        page_width_inches = first_page.rect.width / 72
        page_height_inches = first_page.rect.height / 72

        docx = Document()
        section = docx.sections[0]
        section.page_width = Inches(page_width_inches)
        section.page_height = Inches(page_height_inches)
        section.top_margin = Inches(0)
        section.bottom_margin = Inches(0)
        section.left_margin = Inches(0)
        section.right_margin = Inches(0)

        for page_index in range(document.page_count):
            page = document.load_page(page_index)
            image_path = image_dir / f"pdf-page-{page_index + 1:03d}.png"
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            pixmap.save(str(image_path))

            if page_index > 0:
                docx.add_page_break()
            paragraph = docx.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            image_width = max(0.5, page_width_inches - PAGE_IMAGE_INSET_INCHES)
            paragraph.add_run().add_picture(str(image_path), width=Inches(image_width))

        docx.save(str(destination))
        return document.page_count
    finally:
        document.close()


__all__ = ["render_pdf_pages_as_docx"]
