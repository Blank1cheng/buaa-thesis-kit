from pathlib import Path

import base64
import json
import zipfile
import fitz
from docx import Document
from pypdf import PdfWriter

from buaa_thesis_kit.models import ContentBlock, EquationItem, Metadata, ThesisModel


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "buaa_undergraduate_thesis_template.docx"
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)
OMML_FRAGMENT = (
    '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
    "<m:oMath><m:r><m:t>x+y</m:t></m:r></m:oMath>"
    "</m:oMathPara>"
)


def _document_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as docx_zip:
        return docx_zip.read("word/document.xml").decode("utf-8")


def _write_minimal_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Title: Repair First Thesis")
    document.add_paragraph("Student Name: Zhang San")
    document.add_paragraph("Student ID: 20370001")
    document.add_paragraph("1 Introduction")
    document.add_paragraph("Body text that must be preserved.")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _write_valid_pdf(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as handle:
        writer.write(handle)


def _write_marker_pdf(path: Path, pages: list[list[str]]) -> None:
    document = fitz.open()
    for lines in pages:
        page = document.new_page(width=595, height=842)
        page.insert_text((72, 72), "\n".join(lines), fontsize=12)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    document.close()


def _write_image_only_docx(path: Path) -> None:
    image = path.with_suffix(".png")
    image.write_bytes(TINY_PNG)
    document = Document()
    document.add_picture(str(image))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def test_graph_runner_revises_after_visual_compare_failure_before_finalizing(tmp_path):
    from buaa_thesis_kit.graph import GraphRunner, GraphState, NodeResult

    source = tmp_path / "source.docx"
    output = tmp_path / "output"
    work = tmp_path / "output_work"
    _write_minimal_docx(source)

    def ingest(state: GraphState) -> NodeResult:
        state.notes.append("ingested")
        return NodeResult(next_node="visual_compare")

    def visual_compare(state: GraphState) -> NodeResult:
        if state.retry_count == 0:
            state.findings.append("visual_mismatch")
        return NodeResult(next_node="decide")

    def decide(state: GraphState) -> NodeResult:
        if state.findings and state.retry_count == 0:
            return NodeResult(next_node="revise_plan")
        return NodeResult(next_node="finalize_output")

    def revise_plan(state: GraphState) -> NodeResult:
        state.findings.clear()
        return NodeResult(next_node="visual_compare")

    def finalize_output(state: GraphState) -> NodeResult:
        state.outputs["report.md"] = "pass"
        return NodeResult(next_node=None)

    state = GraphState(source_path=source, output_root=output, work_dir=work, max_retries=2)
    result = GraphRunner(
        {
            "ingest": ingest,
            "visual_compare": visual_compare,
            "decide": decide,
            "revise_plan": revise_plan,
            "finalize_output": finalize_output,
        },
        start_node="ingest",
    ).run(state)

    assert result.status == "pass"
    assert result.retry_count == 1
    assert result.history == [
        "ingest",
        "visual_compare",
        "decide",
        "revise_plan",
        "visual_compare",
        "decide",
        "finalize_output",
    ]


def test_graph_runner_fails_when_revise_loop_exceeds_max_retries(tmp_path):
    from buaa_thesis_kit.graph import GraphRunner, GraphState, NodeResult

    source = tmp_path / "source.docx"
    output = tmp_path / "output"
    work = tmp_path / "output_work"
    _write_minimal_docx(source)

    def always_retry(_state: GraphState) -> NodeResult:
        return NodeResult(next_node="revise_plan")

    def revise(_state: GraphState) -> NodeResult:
        return NodeResult(next_node="decide")

    state = GraphState(source_path=source, output_root=output, work_dir=work, max_retries=1)
    result = GraphRunner(
        {"decide": always_retry, "revise_plan": revise},
        start_node="decide",
    ).run(state)

    assert result.status == "failed"
    assert result.retry_count == 1
    assert result.history == ["decide", "revise_plan", "decide"]
    assert any("maximum graph retry count exceeded" in item for item in result.blocking_items)


def test_graph_runner_fails_when_node_cycle_exceeds_max_steps(tmp_path):
    from buaa_thesis_kit.graph import GraphRunner, GraphState, NodeResult

    source = tmp_path / "source.docx"
    output = tmp_path / "output"
    work = tmp_path / "output_work"
    _write_minimal_docx(source)

    def loop(_state: GraphState) -> NodeResult:
        return NodeResult(next_node="loop")

    state = GraphState(source_path=source, output_root=output, work_dir=work, max_steps=3)
    result = GraphRunner({"loop": loop}, start_node="loop").run(state)

    assert result.status == "failed"
    assert result.history == ["loop", "loop", "loop"]
    assert any("maximum graph step count exceeded" in item for item in result.blocking_items)


def test_plan_minimal_fixes_promotes_missing_spine_to_repair_action(tmp_path):
    from buaa_thesis_kit.graph import GraphState
    from buaa_thesis_kit.graph_nodes import plan_minimal_fixes

    source = tmp_path / "source.docx"
    work = tmp_path / "output_work"
    output = tmp_path / "output"
    _write_minimal_docx(source)
    state = GraphState(source_path=source, output_root=output, work_dir=work)
    state.findings.append("missing_spine")
    state.model = ThesisModel(
        metadata=Metadata(
            title_cn="Repair First Thesis",
            student_name="Zhang San",
            college="Automation College",
            major="Automation",
            date="2026",
        )
    )

    result = plan_minimal_fixes(state)

    assert result.next_node == "apply_word_fixes"
    assert state.repair_actions == ["insert_spine"]
    assert not state.blocking_items


def test_diagnose_compliance_blocks_reference_citation_without_entry(tmp_path):
    from buaa_thesis_kit.graph import GraphState
    from buaa_thesis_kit.graph_nodes import diagnose_compliance

    source = tmp_path / "source.docx"
    work = tmp_path / "output_work"
    output = tmp_path / "output"
    _write_minimal_docx(source)
    state = GraphState(source_path=source, output_root=output, work_dir=work)
    state.source_features["has_spine"] = True
    state.model = ThesisModel(
        sections=[
            ContentBlock(
                id="sec-1",
                type="chapter",
                title="1 Introduction",
                text="This paragraph cites [2].",
                level=1,
            )
        ],
        references=[
            ContentBlock(id="ref-1", type="reference", text="[1] Wang. Flight control. 2026.")
        ],
        status="needs_review",
    )

    result = diagnose_compliance(state)

    assert result.next_node == "plan_minimal_fixes"
    assert "citation_without_reference: [2]" in state.blocking_items


def test_diagnose_compliance_truncated_sample_demotes_reference_gaps_to_review(tmp_path):
    from buaa_thesis_kit.graph import GraphState
    from buaa_thesis_kit.graph_nodes import diagnose_compliance

    source = tmp_path / "source.docx"
    work = tmp_path / "output_work"
    output = tmp_path / "output"
    _write_minimal_docx(source)
    state = GraphState(source_path=source, output_root=output, work_dir=work, sample_mode="truncated")
    state.source_features["has_spine"] = True
    state.model = ThesisModel(
        sections=[
            ContentBlock(
                id="sec-1",
                type="chapter",
                title="1 Introduction",
                text="This truncated sample cites [2] but intentionally omits the later reference list.",
                level=1,
            )
        ],
        references=[
            ContentBlock(id="ref-1", type="reference", text="[1] Wang. Flight control. 2026.")
        ],
        status="needs_review",
    )

    result = diagnose_compliance(state)

    assert result.next_node == "plan_minimal_fixes"
    assert not any("citation_without_reference" in item for item in state.blocking_items)
    assert any("truncated_sample_reference_check_skipped" in item for item in state.manual_review)


def test_apply_word_fixes_renders_docx_source_through_buaa_template(tmp_path):
    from buaa_thesis_kit.graph import GraphState
    from buaa_thesis_kit.graph_nodes import apply_word_fixes

    source = tmp_path / "source.docx"
    work = tmp_path / "output_work"
    output = tmp_path / "output"
    document = Document()
    document.add_paragraph("UNNORMALIZED SOURCE COVER ARTIFACT")
    document.add_paragraph("Body text that must be preserved.")
    document.save(source)
    state = GraphState(source_path=source, output_root=output, work_dir=work, template_path=TEMPLATE)
    state.extraction_source = source
    state.source_kind = "docx"
    state.repair_actions.append("insert_spine")
    state.model = ThesisModel(
        metadata=Metadata(
            title_cn="Repair First Thesis",
            student_name="Zhang San",
            college="Automation College",
            major="Automation",
            date="2026",
        )
    )
    state.model.sections.append(
        ContentBlock(
            id="sec-1",
            type="chapter",
            title="1 Introduction",
            text="Body text that must be preserved.",
            level=1,
        )
    )

    result = apply_word_fixes(state)

    assert result.next_node == "export_pdf"
    assert state.authoritative_docx is not None
    output_doc = Document(str(state.authoritative_docx))
    paragraphs = [paragraph.text for paragraph in output_doc.paragraphs]
    assert "Body text that must be preserved." in paragraphs
    assert "UNNORMALIZED SOURCE COVER ARTIFACT" not in paragraphs
    assert any("毕业设计(论文)" in text for text in paragraphs)
    document_xml = _document_xml(state.authoritative_docx)
    assert "BUAA_VERTICAL_SPINE" in document_xml
    assert 'w:textDirection w:val="tbRl"' in document_xml
    assert any("Repair First Thesis" in text for text in paragraphs)


def test_apply_word_fixes_does_not_treat_template_spine_instruction_as_existing_spine(tmp_path):
    from buaa_thesis_kit.graph import GraphState
    from buaa_thesis_kit.graph_nodes import apply_word_fixes

    source = tmp_path / "source.docx"
    work = tmp_path / "output_work"
    output = tmp_path / "output"
    document = Document()
    document.add_paragraph("论文封面书脊")
    document.add_paragraph("Body text that must be preserved.")
    document.save(source)

    state = GraphState(source_path=source, output_root=output, work_dir=work, template_path=TEMPLATE)
    state.extraction_source = source
    state.source_kind = "docx"
    state.repair_actions.append("insert_spine")
    state.model = ThesisModel(
        metadata=Metadata(
            title_cn="Repair First Thesis",
            student_name="Zhang San",
            college="Automation College",
            major="Automation",
            date="2026",
        )
    )

    apply_word_fixes(state)

    output_doc = Document(str(state.authoritative_docx))
    paragraphs = [paragraph.text for paragraph in output_doc.paragraphs]
    assert "论文封面书脊" not in paragraphs
    assert not any(text.strip() == "书脊" for text in paragraphs)
    assert "BUAA_VERTICAL_SPINE" in _document_xml(state.authoritative_docx)
    assert state.applied_repairs == ["insert_spine"]


def test_apply_word_fixes_fallback_spine_has_no_debug_title(tmp_path, monkeypatch):
    import buaa_thesis_kit.graph_nodes as graph_nodes
    from buaa_thesis_kit.graph import GraphState
    from buaa_thesis_kit.graph_nodes import apply_word_fixes

    source = tmp_path / "source.docx"
    work = tmp_path / "output_work"
    output = tmp_path / "output"
    _write_minimal_docx(source)

    def fake_render_editable_buaa_docx(_template, model, destination):
        document = Document()
        document.add_paragraph(model.metadata.title_cn)
        document.add_paragraph(model.metadata.student_id)
        document.add_paragraph("Body text that must be preserved.")
        document.save(destination)

    monkeypatch.setattr(graph_nodes, "render_editable_buaa_docx", fake_render_editable_buaa_docx)
    state = GraphState(source_path=source, output_root=output, work_dir=work, template_path=TEMPLATE)
    state.extraction_source = source
    state.source_kind = "docx"
    state.repair_actions.append("insert_spine")
    state.model = ThesisModel(
        metadata=Metadata(
            title_cn="Fallback Spine Thesis",
            student_id="20370001",
            student_name="Zhang San",
            college="Automation College",
            major="Automation",
            date="2026",
        ),
        sections=[
            ContentBlock(
                id="sec-1",
                type="chapter",
                title="1 Introduction",
                text="Body text that must be preserved.",
                level=1,
            )
        ],
    )

    apply_word_fixes(state)

    assert state.authoritative_docx is not None
    visible_text = "\n".join(paragraph.text for paragraph in Document(str(state.authoritative_docx)).paragraphs)
    assert "Book Spine" not in visible_text
    assert "书脊" not in visible_text
    document_xml = _document_xml(state.authoritative_docx)
    assert "BUAA_VERTICAL_SPINE" in document_xml
    assert 'w:textDirection w:val="tbRl"' in document_xml


def test_visual_compare_blocks_pdf_word_without_editable_text(tmp_path):
    from buaa_thesis_kit.graph import GraphState
    from buaa_thesis_kit.graph_nodes import visual_compare

    output = tmp_path / "output"
    work = tmp_path / "work"
    source = tmp_path / "source.pdf"
    authoritative = tmp_path / "image-only.docx"
    _write_valid_pdf(source)
    _write_image_only_docx(authoritative)

    state = GraphState(source_path=source, output_root=output, work_dir=work)
    state.source_kind = "pdf"
    state.authoritative_docx = authoritative
    state.model = ThesisModel(
        metadata=Metadata(title_cn="PDF Pipeline Thesis", student_id="20370001")
    )

    result = visual_compare(state)

    assert result.next_node == "decide"
    assert any("editable_text_missing" in item for item in state.blocking_items)


def test_visual_compare_blocks_when_expected_omml_equation_is_not_in_word(tmp_path):
    from buaa_thesis_kit.graph import GraphState
    from buaa_thesis_kit.graph_nodes import visual_compare

    output = tmp_path / "output"
    work = tmp_path / "work"
    source = tmp_path / "source.pdf"
    authoritative = tmp_path / "missing-omml.docx"
    _write_valid_pdf(source)

    document = Document()
    document.add_paragraph("Equation Thesis")
    document.add_paragraph("20370001")
    document.add_paragraph("Book Spine")
    document.add_paragraph("Editable body text that must be preserved.")
    document.add_paragraph("x + y")
    document.save(authoritative)

    state = GraphState(source_path=source, output_root=output, work_dir=work)
    state.source_kind = "pdf"
    state.authoritative_docx = authoritative
    state.model = ThesisModel(
        metadata=Metadata(title_cn="Equation Thesis", student_id="20370001"),
        sections=[
            ContentBlock(
                id="sec-1",
                type="chapter",
                text="Editable body text that must be preserved.",
                level=1,
            )
        ],
        equations=[
            EquationItem(
                id="eq-1",
                kind="pdf-text-equation",
                text="x+y",
                omml=OMML_FRAGMENT,
                requires_review=False,
            )
        ],
    )

    result = visual_compare(state)

    assert result.next_node == "decide"
    assert any("editable_omml_equation_missing" in item for item in state.blocking_items)


def test_visual_compare_blocks_collapsed_exported_pdf_front_matter(tmp_path):
    from buaa_thesis_kit.graph import GraphState
    from buaa_thesis_kit.graph_nodes import visual_compare

    output = tmp_path / "output"
    work = tmp_path / "work"
    source = tmp_path / "source.pdf"
    thesis_pdf = output / "thesis.pdf"
    _write_valid_pdf(source)
    _write_marker_pdf(
        thesis_pdf,
        [[
            "Graduation Thesis",
            "Book Spine",
            "Declaration",
            "Chinese Abstract",
            "Abstract",
            "1 Introduction",
        ]],
    )

    state = GraphState(source_path=source, output_root=output, work_dir=work)
    state.source_kind = "pdf"
    state.outputs["pdf"] = "pass"

    result = visual_compare(state)

    assert result.next_node == "decide"
    assert any("pdf_layout_collapsed" in item for item in state.blocking_items)


def test_pipeline_report_records_graph_history_and_spine_repair(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.docx"
    output = tmp_path / "output"
    _write_minimal_docx(source)

    def stub_pdf_export(_source: Path, target: Path) -> tuple[bool, str]:
        _write_valid_pdf(target)
        return True, "stubbed Word PDF export"

    def pass_validate_output_text_file(candidate, output_report=None):
        report = {
            "status": "pass",
            "candidate": str(candidate),
            "failures": [],
            "counts": {},
        }
        if output_report is not None:
            Path(output_report).parent.mkdir(parents=True, exist_ok=True)
            output_report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report

    monkeypatch.setattr(pipeline, "export_pdf_from_docx", stub_pdf_export)
    monkeypatch.setattr(pipeline, "validate_output_text_file", pass_validate_output_text_file)

    report = pipeline.run_pipeline(source, output)

    assert report["outputs"]["word"] == "pass"
    assert report["outputs"]["pdf"] == "pass"
    assert report["outputs"]["tex"] in {"pass", "needs_review"}
    assert report["outputs"]["image"] == "pass"
    assert any("Graph nodes executed:" in note for note in report["notes"])
    assert any("missing_spine" in note for note in report["notes"])
    assert (output / "thesis.docx").is_file()
    assert (output / "thesis.pdf").is_file()
    assert (output / "thesis.tex").is_file()
    assert (output / "report.md").is_file()
    assert (output / "image").is_dir()
