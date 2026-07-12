from pathlib import Path

from docx import Document
from pypdf import PdfWriter


def _write_docx(path: Path, paragraphs: list[str]) -> None:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def test_output_text_reports_specific_template_leak_reasons(tmp_path):
    from buaa_thesis_kit.harness.validators import validate_output_text_file

    candidate = tmp_path / "thesis7-like.docx"
    _write_docx(
        candidate,
        [
            "分类号 T P 2 7 3",
            "论文封面书脊",
            "四号黑体字",
            "小四号黑体字",
            "注：任务书应该附在已完成的毕业设计（论文）的首页。",
            "摘    要",
            "Research on Key Technologies",
            "Author: CUI Runhao",
            "目    录",
            "本人声明 ........................................ 1",
            "Mao Xia, Ding Yukun. Template reference.",
            "刘国钧. 图书馆目录[M]. 北京: 高等教育出版社.",
            "I 级叶/盘转子错频方案的对比分析 ................ 15",
            "1 绪论",
            "第 48 页",
        ],
    )

    report = validate_output_text_file(candidate)

    reason_ids = {item["id"] for item in report["failures"]}
    assert report["status"] == "failed"
    assert "template_instructions_or_sample_leak" in reason_ids
    assert "spine_template_instruction_leak" in reason_ids
    assert "taskbook_template_note_leak" in reason_ids
    assert "toc_contains_declaration" in reason_ids
    assert "toc_contains_template_sample_reference" in reason_ids
    assert "body_contains_template_page_number_48" in reason_ids
    assert "cover_classification_split" in reason_ids
    assert "cn_abstract_contains_english_title" in reason_ids


def test_pipeline_reruns_output_text_after_pdf_export_mutates_docx(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    def dirtying_pdf_export(source: Path, target: Path) -> tuple[bool, str]:
        document = Document(source)
        document.add_paragraph("论文封面书脊")
        document.save(source)
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with target.open("wb") as handle:
            writer.write(handle)
        return True, "stubbed dirtying Word finalize"

    output = tmp_path / "pipeline_gate"
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", dirtying_pdf_export)

    report = pipeline.run_pipeline(
        Path("tests/fixtures/truncated_input.docx"),
        output,
        sample_mode="truncated",
    )

    post_report = output / "harness" / "output_text_post_finalize_report.json"
    assert report["status"] == "failed"
    assert post_report.is_file()
    assert any("harness_output_text_post_finalize_failed" in item for item in report["blocking_items"])
