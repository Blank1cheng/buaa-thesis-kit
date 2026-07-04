import base64
from pathlib import Path

from docx import Document
from docx.shared import Inches

from buaa_thesis_kit.editable_template_render import render_editable_buaa_docx
from buaa_thesis_kit.models import ContentBlock, EquationItem, Metadata, ThesisModel
from buaa_thesis_kit.docx_acceptance import inspect_docx_output


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "buaa_undergraduate_thesis_template.docx"
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)
OMML_FRAGMENT = (
    '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
    "<m:oMath><m:r><m:t>x+y</m:t></m:r></m:oMath>"
    "</m:oMathPara>"
)


def _write_image_only_docx(path: Path) -> None:
    image = path.with_suffix(".png")
    image.write_bytes(TINY_PNG)
    document = Document()
    document.add_picture(str(image))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _write_editable_docx_with_full_page_screenshot(path: Path) -> None:
    image = path.with_suffix(".png")
    image.write_bytes(TINY_PNG)
    document = Document()
    document.add_paragraph("PDF Pipeline Thesis")
    document.add_paragraph("20370001")
    document.add_paragraph("Editable body text extracted from the source PDF.")
    document.add_picture(str(image), width=Inches(6.0), height=Inches(8.5))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def test_inspect_docx_output_blocks_pdf_word_without_editable_text(tmp_path):
    output = tmp_path / "image-only.docx"
    _write_image_only_docx(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="PDF Pipeline Thesis", student_id="20370001"),
        source_kind="pdf",
        require_spine=False,
    )

    assert any("editable_text_missing" in item for item in result.blocking_items)


def test_inspect_docx_output_blocks_pdf_full_page_screenshot_even_with_editable_text(tmp_path):
    output = tmp_path / "editable-plus-screenshot.docx"
    _write_editable_docx_with_full_page_screenshot(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="PDF Pipeline Thesis", student_id="20370001"),
        source_kind="pdf",
        require_spine=False,
    )

    assert any("word_page_screenshot" in item for item in result.blocking_items)


def test_inspect_docx_output_blocks_pdf_word_when_source_body_text_is_not_editable(tmp_path):
    image = tmp_path / "body-screenshot.png"
    image.write_bytes(TINY_PNG)
    output = tmp_path / "metadata-plus-body-image.docx"
    document = Document()
    document.add_paragraph("PDF Pipeline Thesis")
    document.add_paragraph("20370001")
    document.add_paragraph("Body content is represented by the raster image below.")
    document.add_picture(str(image), width=Inches(5.5), height=Inches(6.0))
    document.save(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="PDF Pipeline Thesis", student_id="20370001"),
        source_kind="pdf",
        require_spine=False,
        required_body_snippets=["This source body paragraph must be editable text."],
    )

    assert any("editable_body_text_missing" in item for item in result.blocking_items)


def test_inspect_docx_output_blocks_docx_word_without_editable_text(tmp_path):
    output = tmp_path / "image-only.docx"
    _write_image_only_docx(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="Editable Thesis", student_id="20370001"),
        source_kind="docx",
        require_spine=False,
    )

    assert any("editable_text_missing" in item for item in result.blocking_items)


def test_inspect_docx_output_blocks_equation_preview_screenshot_substitute(tmp_path):
    output = tmp_path / "equation-preview.docx"
    document = Document()
    document.add_paragraph("Editable Thesis")
    document.add_paragraph("20370001")
    document.add_paragraph("[Equation preview inserted] equation.bin")
    document.save(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="Editable Thesis", student_id="20370001"),
        source_kind="docx",
        require_spine=False,
    )

    assert any("editable_equation_missing" in item for item in result.blocking_items)


def test_inspect_docx_output_blocks_figure_debug_and_local_paths(tmp_path):
    output = tmp_path / "debug-text.docx"
    document = Document()
    document.add_paragraph("Editable Thesis")
    document.add_paragraph("20370001")
    document.add_paragraph("[Figure inserted] D:\\Work\\.worktrees\\image1.png")
    document.save(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="Editable Thesis", student_id="20370001"),
        source_kind="docx",
        require_spine=False,
    )

    assert any("unsafe_word_body_text" in item for item in result.blocking_items)


def test_inspect_docx_output_blocks_english_references_heading_for_chinese_output(tmp_path):
    output = tmp_path / "english-references.docx"
    document = Document()
    document.add_paragraph("Editable Thesis")
    document.add_paragraph("20370001")
    document.add_paragraph("References")
    document.add_paragraph("[1] Wang. Test. 2026.")
    document.save(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="Editable Thesis", student_id="20370001"),
        source_kind="docx",
        require_spine=False,
    )

    assert any("english_references_heading_visible" in item for item in result.blocking_items)


def test_inspect_docx_output_blocks_consecutive_formula_token_dump(tmp_path):
    output = tmp_path / "formula-dump.docx"
    document = Document()
    document.add_paragraph("Editable Thesis")
    document.add_paragraph("20370001")
    for token in ("h(x,y)", "OTF(u,v)", "H(u,v)=F(u,v)G(u,v)", "MTF(f)", "g(x,y)=h(x,y)*f(x,y)"):
        document.add_paragraph(token)
    document.save(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="Editable Thesis", student_id="20370001"),
        source_kind="docx",
        require_spine=False,
    )

    assert any("formula_token_dump_visible" in item for item in result.blocking_items)


def test_inspect_docx_output_allows_wrapped_english_reference_fragments(tmp_path):
    output = tmp_path / "references.docx"
    document = Document()
    document.add_paragraph("Editable Thesis")
    document.add_paragraph("20370001")
    for fragment in (
        "on of roller in a hot strip mill based on multi-scale LSTM with multi-head a",
        "ttention[J].Reliability Engineering and System Safety,2024,248110161-.",
        "[22]Zhang T ,Wang H .Quantile regression network-based cross-domain prediction",
        "hod for Multi-Component System Considering Maintenance: Subsea Christmas",
        "Tree System as A Case Study[J].China Ocean Engineering,2024,38(2):198-209.",
        "[30]Jang I ,Kim H C .Prediction of Remaining Useful Life (RUL) of Electronic",
        "Components in the POSAFE-Q PLC Platform under NPP Dynamic Stress Con",
        "ditions[J].Nuclear Engineering and Technology,2024,56(5):1863-1873.",
    ):
        document.add_paragraph(fragment)
    document.save(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="Editable Thesis", student_id="20370001"),
        source_kind="docx",
        require_spine=False,
    )

    assert not any("formula_token_dump_visible" in item for item in result.blocking_items)


def test_inspect_docx_output_blocks_internal_equation_object_token(tmp_path):
    output = tmp_path / "equation-token.docx"
    document = Document()
    document.add_paragraph("Editable Thesis")
    document.add_paragraph("20370001")
    document.add_paragraph("__BUAA_EDITABLE_EQUATION_OBJECT__buaa-equation-eq-1__")
    document.save(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="Editable Thesis", student_id="20370001"),
        source_kind="docx",
        require_spine=False,
    )

    assert any("editable_equation_object_token_visible" in item for item in result.blocking_items)


def test_inspect_docx_output_blocks_when_expected_omml_equation_is_missing(tmp_path):
    output = tmp_path / "missing-omml-equation.docx"
    document = Document()
    document.add_paragraph("Equation Thesis")
    document.add_paragraph("20370001")
    document.add_paragraph("Book Spine")
    document.add_paragraph("x + y")
    document.save(output)

    result = inspect_docx_output(
        output,
        Metadata(title_cn="Equation Thesis", student_id="20370001"),
        source_kind="docx",
        require_spine=True,
        expected_omml_equation_count=1,
    )

    assert any("editable_omml_equation_missing" in item for item in result.blocking_items)


def test_inspect_docx_output_accepts_vertical_spine_textbox_without_debug_marker(tmp_path):
    output = tmp_path / "vertical-spine.docx"
    model = ThesisModel(
        metadata=Metadata(
            title_cn="Vertical Spine Thesis",
            student_id="20370001",
            student_name="Zhang San",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026年7月",
            classification="TN953",
        ),
        sections=[
            ContentBlock(
                id="section-1",
                type="section",
                title="1 Introduction",
                text="This PDF contains extractable thesis text.",
            )
        ],
    )
    render_editable_buaa_docx(TEMPLATE, model, output)
    visible_text = "\n".join(paragraph.text for paragraph in Document(str(output)).paragraphs)

    result = inspect_docx_output(
        output,
        model.metadata,
        source_kind="pdf",
        require_spine=True,
    )

    assert "Book Spine" not in visible_text
    assert "书脊" not in visible_text
    assert not any("spine_missing" in item for item in result.blocking_items)


def test_inspect_docx_output_accepts_editable_template_word(tmp_path):
    output = tmp_path / "editable.docx"
    model = ThesisModel(
        metadata=Metadata(
            title_cn="PDF Pipeline Thesis",
            student_id="20370001",
            student_name="Zhang San",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026年7月",
            classification="TN953",
        ),
        sections=[
            ContentBlock(
                id="section-1",
                type="section",
                title="1 Introduction",
                text="This PDF contains extractable thesis text.",
            )
        ],
    )
    render_editable_buaa_docx(TEMPLATE, model, output)

    result = inspect_docx_output(
        output,
        model.metadata,
        source_kind="pdf",
        require_spine=True,
    )

    assert result.blocking_items == []
    assert result.editability["editable_characters"] > 0
    assert result.editability["paragraph_count"] > 0
    assert result.editability["page_screenshot_drawing_count"] == 0
    assert any("editable Word validation passed" in note for note in result.notes)


def test_inspect_docx_output_reports_editable_omml_equation_count(tmp_path):
    output = tmp_path / "editable-equation.docx"
    model = ThesisModel(
        metadata=Metadata(
            title_cn="Equation Thesis",
            student_id="20370001",
            student_name="Zhang San",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026-07",
            classification="TN953",
        ),
        sections=[ContentBlock(id="section-1", type="section", text="Editable body text.")],
        equations=[
            EquationItem(
                id="eq-1",
                kind="omml",
                text="x+y",
                omml=OMML_FRAGMENT,
                requires_review=False,
            )
        ],
    )
    render_editable_buaa_docx(TEMPLATE, model, output)

    result = inspect_docx_output(
        output,
        model.metadata,
        source_kind="docx",
        require_spine=True,
        expected_omml_equation_count=1,
    )

    assert result.blocking_items == []
    assert result.editability["omml_equation_count"] == 1
