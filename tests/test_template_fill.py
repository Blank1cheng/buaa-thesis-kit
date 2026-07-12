import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches

from buaa_thesis_kit.models import AssetItem, ContentBlock, EquationItem, Metadata, ThesisModel
from buaa_thesis_kit.template_fill import fill_word_template


TINY_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x04\x00\x00\x00\xb5\x1c\x0c\x02\x00\x00\x00\x0bIDATx\xdac\xfc\xff"
    b"\x1f\x00\x03\x03\x02\x00\xef\xbf\xa7\xdb\x00\x00\x00\x00IEND\xaeB`\x82"
)
OMML_FRAGMENT = (
    '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
    "<m:oMath><m:r><m:t>x+y</m:t></m:r></m:oMath>"
    "</m:oMathPara>"
)
OLE_OBJECT_XML = (
    '<w:object xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<v:shape id="_x0000_i1025" type="#_x0000_t75" style="width:120pt;height:24pt">'
    '<v:imagedata r:id="rIdEquationImage" o:title=""/>'
    "</v:shape>"
    '<o:OLEObject Type="Embed" ProgID="Equation.DSMT4" r:id="rIdEquation"/>'
    "</w:object>"
)


def _sample_model(tmp_path: Path) -> ThesisModel:
    missing_image = tmp_path / "missing-image.png"
    return ThesisModel(
        metadata=Metadata(
            title_cn="Data Driven Flight Control",
            title_en="Data Driven Flight Control",
            student_name="Zhang San",
            student_id="20370001",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026-06",
            classification="TP391",
            unit_code="10006",
        ),
        front_matter={
            "chinese_abstract": "Chinese abstract paragraph.",
            "english_abstract": "English abstract paragraph.",
        },
        sections=[
            ContentBlock(id="sec-1", type="chapter", title="1 Introduction", text="Opening paragraph.", level=1),
            ContentBlock(id="sec-2", type="section", title="1.1 Background", text="Background paragraph.", level=2),
        ],
        tables=[
            ContentBlock(
                id="tbl-1",
                type="table",
                title="Table 1 Metrics",
                text="Metric\tValue\nAccuracy\t98%",
            )
        ],
        figures=[
            AssetItem(
                id="fig-1",
                type="image",
                path=str(missing_image),
                caption="Figure 1 System overview",
                requires_review=True,
            )
        ],
        equations=[
            EquationItem(id="eq-1", kind="omml", text="x+y", number="(1)", requires_review=True),
        ],
        references=[
            ContentBlock(id="ref-1", type="reference", text="[1] Wang. Flight control study. 2026."),
            ContentBlock(id="ref-2", type="reference", text="[2] Smith. Data systems. 2025."),
        ],
        appendices=[
            ContentBlock(id="app-1", type="appendix", title="Appendix A", text="Supplemental material."),
        ],
    )


def _paragraph_texts(doc: Document) -> list[str]:
    return [paragraph.text for paragraph in doc.paragraphs]


def _all_text(doc: Document) -> str:
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            table_text.extend(cell.text for cell in row.cells)
    return "\n".join(_paragraph_texts(doc) + table_text)


def test_replaces_paragraph_and_table_cell_placeholders_without_leaving_tokens(tmp_path):
    template = tmp_path / "placeholder-template.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("Title: {{TITLE_CN}}")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "{{STUDENT_NAME}}"
    table.cell(1, 0).text = "Unknown"
    table.cell(1, 1).text = "{{NOT_A_FIELD}}"
    doc.save(template)

    fill_word_template(template, _sample_model(tmp_path), output)

    result = Document(output)
    text = _all_text(result)
    assert "Title: Data Driven Flight Control" in text
    assert "Zhang San" in text
    assert "{{" not in text
    assert "}}" not in text


def test_body_and_references_placeholders_expand_to_multiple_paragraphs(tmp_path):
    template = tmp_path / "block-template.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{BODY}}")
    doc.add_paragraph("{{REFERENCES}}")
    doc.save(template)

    fill_word_template(template, _sample_model(tmp_path), output)

    result = Document(output)
    paragraphs = [text for text in _paragraph_texts(result) if text]
    assert "1 Introduction" in paragraphs
    assert "Opening paragraph." in paragraphs
    assert "1.1 Background" in paragraphs
    assert "Background paragraph." in paragraphs
    assert "[1] Wang. Flight control study. 2026." in paragraphs
    assert "[2] Smith. Data systems. 2025." in paragraphs
    assert "\n" not in paragraphs[0]


def test_no_placeholder_template_fallback_builds_docx_from_model(tmp_path):
    template = tmp_path / "official-like-template.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("Official template cover text without fields.")
    doc.save(template)

    fill_word_template(template, _sample_model(tmp_path), output)

    result = Document(output)
    text = _all_text(result)
    assert "Data Driven Flight Control" in text
    assert "Zhang San" in text
    assert "20370001" in text
    assert "Chinese abstract paragraph." in text
    assert "English abstract paragraph." in text
    assert "1 Introduction" in text
    assert "Opening paragraph." in text
    assert "Metric" in text
    assert "Accuracy" in text
    assert "[图像缺失：Figure 1 System overview，需人工确认]" in text
    assert "missing-image.png" not in text
    assert "[公式缺失：" in text
    assert "x+y" not in text
    assert "[1] Wang. Flight control study. 2026." in text
    assert "Appendix A" in text
    assert "Supplemental material." in text
    assert "Official template cover text without fields." not in text


def test_fill_creates_only_requested_docx_in_empty_output_directory(tmp_path):
    template = tmp_path / "template.docx"
    output_dir = tmp_path / "empty-output"
    output = output_dir / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{TITLE_CN}}")
    doc.save(template)

    fill_word_template(template, _sample_model(tmp_path), output)

    assert sorted(path.name for path in output_dir.iterdir()) == ["thesis.docx"]


def test_generated_docx_opens_and_contains_document_xml(tmp_path):
    template = tmp_path / "template.docx"
    output = tmp_path / "nested" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{TITLE_CN}}")
    doc.save(template)

    fill_word_template(template, _sample_model(tmp_path), output)

    opened = Document(output)
    assert opened.paragraphs[0].text == "Data Driven Flight Control"
    with zipfile.ZipFile(output) as package:
        assert "word/document.xml" in package.namelist()


def test_scalar_only_placeholder_template_appends_unrepresented_major_blocks(tmp_path):
    template = tmp_path / "scalar-only-template.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{TITLE_CN}}")
    doc.save(template)

    fill_word_template(template, _sample_model(tmp_path), output)

    result = Document(output)
    text = _all_text(result)
    assert "Data Driven Flight Control" in text
    assert "Chinese abstract paragraph." in text
    assert "English abstract paragraph." in text
    assert "1 Introduction" in text
    assert "Metric" in text
    assert "Figure 1 System overview" in text
    assert "[公式缺失：" in text
    assert "[1] Wang. Flight control study. 2026." in text
    assert "Appendix A" in text


def test_placeholder_template_preserves_existing_section_margins(tmp_path):
    template = tmp_path / "custom-margin-template.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(1.0)
    doc.add_paragraph("{{TITLE_CN}}")
    doc.save(template)

    fill_word_template(template, _sample_model(tmp_path), output)

    result_section = Document(output).sections[0]
    assert result_section.top_margin == Inches(0.7)
    assert result_section.bottom_margin == Inches(0.8)
    assert result_section.left_margin == Inches(0.9)
    assert result_section.right_margin == Inches(1.0)


def test_tables_placeholder_inserts_real_word_table(tmp_path):
    template = tmp_path / "table-placeholder.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{TABLES}}")
    doc.save(template)

    fill_word_template(template, _sample_model(tmp_path), output)

    result = Document(output)
    assert result.tables
    assert result.tables[0].cell(0, 0).text == "Metric"
    assert result.tables[0].cell(0, 1).text == "Value"
    assert result.tables[0].cell(1, 0).text == "Accuracy"
    assert result.tables[0].cell(1, 1).text == "98%"


def test_figures_placeholder_inserts_supported_image_and_reviews_missing_image(tmp_path):
    image_path = tmp_path / "tiny.png"
    image_path.write_bytes(TINY_PNG)
    model = _sample_model(tmp_path)
    model.figures = [
        AssetItem(id="fig-ok", type="image", path=str(image_path), caption="Inserted figure"),
        AssetItem(id="fig-missing", type="image", path=str(tmp_path / "missing.png"), caption="Missing figure"),
    ]
    template = tmp_path / "figure-placeholder.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{FIGURES}}")
    doc.save(template)

    fill_word_template(template, model, output)

    result = Document(output)
    text = _all_text(result)
    assert "Inserted figure" in text
    assert "[Figure inserted]" not in text
    assert "[图像缺失：Missing figure，需人工确认]" in text
    assert "missing.png" not in text
    with zipfile.ZipFile(output) as package:
        assert any(name.startswith("word/media/") for name in package.namelist())


def test_equations_placeholder_inserts_omml_word_math(tmp_path):
    model = _sample_model(tmp_path)
    model.equations = [
        EquationItem(
            id="eq-omml",
            kind="omml",
            text="x+y",
            omml=OMML_FRAGMENT,
            requires_review=True,
        )
    ]
    template = tmp_path / "equation-placeholder.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{EQUATIONS}}")
    doc.save(template)

    fill_word_template(template, model, output)

    with zipfile.ZipFile(output) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "<m:oMathPara" in document_xml
    assert "<m:t>x+y</m:t>" in document_xml
    assert "[Equation requires review]" not in document_xml


def test_equations_placeholder_inserts_supported_embedded_equation_preview(tmp_path):
    preview = tmp_path / "equation-preview.png"
    preview.write_bytes(TINY_PNG)
    model = _sample_model(tmp_path)
    model.equations = [
        EquationItem(
            id="eq-preview",
            kind="embedded-object",
            text="equation.bin",
            preview_path=str(preview),
            requires_review=True,
        )
    ]
    template = tmp_path / "equation-preview-template.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{EQUATIONS}}")
    doc.save(template)

    fill_word_template(template, model, output)

    result = Document(output)
    text = _all_text(result)
    assert "[Equation preview inserted]" not in text
    assert "[Equation requires review]" not in text
    with zipfile.ZipFile(output) as package:
        assert any(name.startswith("word/media/") for name in package.namelist())


def test_equations_placeholder_preserves_editable_embedded_equation_object(tmp_path):
    preview = tmp_path / "equation-preview.png"
    preview.write_bytes(TINY_PNG)
    object_path = tmp_path / "equation.bin"
    object_path.write_bytes(b"equation ole payload")
    model = _sample_model(tmp_path)
    model.equations = [
        EquationItem(
            id="eq-object",
            kind="embedded-object",
            text="equation.bin",
            preview_path=str(preview),
            object_path=str(object_path),
            object_xml=OLE_OBJECT_XML,
            requires_review=True,
        )
    ]
    template = tmp_path / "equation-object-template.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{EQUATIONS}}")
    doc.save(template)

    fill_word_template(template, model, output)

    result = Document(output)
    text = _all_text(result)
    assert "[Equation preview inserted]" not in text
    assert "[Equation requires review]" not in text
    with zipfile.ZipFile(output) as package:
        names = package.namelist()
        document_xml = package.read("word/document.xml").decode("utf-8")
        rels_xml = package.read("word/_rels/document.xml.rels").decode("utf-8")
        content_types = package.read("[Content_Types].xml").decode("utf-8")
        embedding_names = [name for name in names if name.startswith("word/embeddings/")]
        media_names = [name for name in names if name.startswith("word/media/")]

        assert "<o:OLEObject" in document_xml
        assert 'ProgID="Equation.DSMT4"' in document_xml
        assert 'r:id="rIdEquation"' not in document_xml
        assert 'r:id="rIdEquationImage"' not in document_xml
        assert "oleObject" in rels_xml
        assert "buaa-equation-eq-object.bin" in "\n".join(embedding_names)
        assert "buaa-equation-eq-object.png" in "\n".join(media_names)
        assert package.read("word/embeddings/buaa-equation-eq-object.bin") == b"equation ole payload"
        assert "application/vnd.openxmlformats-officedocument.oleObject" in content_types


def test_inline_scalar_replacement_preserves_unrelated_bold_run(tmp_path):
    template = tmp_path / "runs-template.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    keep = paragraph.add_run("KeepBold ")
    keep.bold = True
    paragraph.add_run("{{TITLE_CN}}")
    doc.save(template)

    fill_word_template(template, _sample_model(tmp_path), output)

    result_paragraph = Document(output).paragraphs[0]
    assert result_paragraph.text == "KeepBold Data Driven Flight Control"
    assert result_paragraph.runs[0].text == "KeepBold "
    assert result_paragraph.runs[0].bold is True


def test_placeholder_scanning_does_not_create_header_or_footer_parts(tmp_path):
    template = tmp_path / "no-header-footer-template.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{TITLE_CN}}")
    doc.save(template)
    with zipfile.ZipFile(template) as package:
        assert not any(name.startswith("word/header") or name.startswith("word/footer") for name in package.namelist())

    fill_word_template(template, _sample_model(tmp_path), output)

    with zipfile.ZipFile(output) as package:
        assert not any(name.startswith("word/header") or name.startswith("word/footer") for name in package.namelist())


def test_corrupt_existing_png_is_reviewed_without_false_inserted_label(tmp_path):
    corrupt_image = tmp_path / "corrupt.png"
    corrupt_image.write_bytes(b"not a real png")
    model = _sample_model(tmp_path)
    model.figures = [AssetItem(id="fig-bad", type="image", path=str(corrupt_image), caption="Corrupt figure")]

    placeholder_template = tmp_path / "placeholder-figure.docx"
    placeholder_output = tmp_path / "placeholder-out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("{{FIGURES}}")
    doc.save(placeholder_template)
    fill_word_template(placeholder_template, model, placeholder_output)

    fallback_template = tmp_path / "fallback-figure.docx"
    fallback_output = tmp_path / "fallback-out" / "thesis.docx"
    doc = Document()
    doc.add_paragraph("official template without placeholders")
    doc.save(fallback_template)
    fill_word_template(fallback_template, model, fallback_output)

    for output in (placeholder_output, fallback_output):
        text = _all_text(Document(output))
        assert "[Figure inserted]" not in text
        assert "[Figure requires review]" not in text
        assert "[图像缺失：Corrupt figure，需人工确认]" in text
        assert "corrupt.png" not in text


def test_inline_scalar_replacement_preserves_trailing_space_before_bold_run(tmp_path):
    template = tmp_path / "placeholder-space-template.docx"
    output = tmp_path / "out" / "thesis.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("{{TITLE_CN}} ")
    tail = paragraph.add_run("Tail")
    tail.bold = True
    doc.save(template)

    fill_word_template(template, _sample_model(tmp_path), output)

    result_paragraph = Document(output).paragraphs[0]
    assert result_paragraph.text == "Data Driven Flight Control Tail"
    assert result_paragraph.runs[1].text == "Tail"
    assert result_paragraph.runs[1].bold is True
