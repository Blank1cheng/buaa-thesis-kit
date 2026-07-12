import base64
import json
import zipfile
from pathlib import Path

import pytest
import fitz
from docx import Document
from pypdf import PdfWriter

from buaa_thesis_kit.models import ContentBlock, EquationItem, ThesisModel
from buaa_thesis_kit.validate import validate_clean_output


pytestmark = pytest.mark.legacy_word


TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)
SUCCESS_OUTPUT_NAMES = [
    "harness",
    "image",
    "model.json",
    "render_smoke",
    "report.md",
    "template_diff",
    "template_inheritance_report.json",
    "thesis.docx",
    "thesis.pdf",
    "thesis.tex",
]
WORD_NO_PDF_OUTPUT_NAMES = [
    "harness",
    "image",
    "model.json",
    "report.md",
    "template_inheritance_report.json",
    "thesis.docx",
    "thesis.tex",
]


def _write_valid_pdf(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as handle:
        writer.write(handle)


def _stub_pdf_export(_source: Path, target: Path) -> tuple[bool, str]:
    _write_valid_pdf(target)
    return True, "stubbed Word PDF export"


def _failing_pdf_export(_source: Path, _target: Path) -> tuple[bool, str]:
    return False, "Word PDF export unavailable in this environment"


def _stub_doc_conversion(source: Path, target: Path) -> tuple[bool, str]:
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(Path(source).read_bytes())
    return True, "stubbed DOC to DOCX conversion"


def _failing_doc_conversion(_source: Path, _target: Path) -> tuple[bool, str]:
    return False, "Word DOC conversion unavailable in this environment"


def _write_source_docx(path: Path, *, with_image: bool = False) -> None:
    document = Document()
    for text in [
        "中文题目：流水线集成测试",
        "英文题目：Pipeline Integration Test",
        "学生姓名：张三",
        "学号：20370001",
        "学院：自动化科学与电气工程学院",
        "专业：自动化",
        "指导教师：李四",
        "日期：2026年6月",
        "摘要",
        "这是中文摘要。",
        "ABSTRACT",
        "This is the English abstract.",
        "1 绪论",
        "正文。",
        "参考文献",
        "[1] 王五. 流水线测试[J]. 2026.",
    ]:
        document.add_paragraph(text)
    if with_image:
        image_path = path.with_suffix(".png")
        image_path.write_bytes(TINY_PNG)
        document.add_picture(str(image_path))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _write_text_pdf(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text(
        (72, 72),
        "\n".join(
            [
                "Title: PDF Pipeline Thesis",
                "Student Name: Zhang San",
                "Student ID: 20370001",
                "College: Automation College",
                "Major: Automation",
                "Advisor: Li Si",
                "Date: 2026-07",
                "1 Introduction",
                "This PDF contains extractable thesis text.",
                "References",
                "[1] Wang Wu. Test reference. 2026.",
            ]
        ),
        fontsize=12,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    document.close()


def _write_text_pdf_with_table(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text(
        (72, 72),
        "\n".join(
            [
                "Title: PDF Table Thesis",
                "Student Name: Zhang San",
                "Student ID: 20370001",
                "College: Automation College",
                "Major: Automation",
                "Advisor: Li Si",
                "Date: 2026-07",
                "1 Introduction",
                "Before table paragraph.",
                "Table 1.1 Evaluation metrics",
                "Metric\tValue",
                "Accuracy\t98%",
                "Recall\t95%",
                "After table paragraph.",
            ]
        ),
        fontsize=12,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    document.close()


def _write_text_pdf_with_equation(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text(
        (72, 72),
        "\n".join(
            [
                "Title: PDF Equation Thesis",
                "Student Name: Zhang San",
                "Student ID: 20370001",
                "College: Automation College",
                "Major: Automation",
                "Advisor: Li Si",
                "Date: 2026-07",
                "1 Introduction",
                "The state transition model is defined below.",
                "x_k = F x_{k-1} + w_k (2.1)",
                "The equation above is used for prediction.",
            ]
        ),
        fontsize=12,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    document.close()


def _write_blank_pdf(path: Path) -> None:
    document = fitz.open()
    document.new_page(width=300, height=400)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    document.close()


def _public_names(output_dir: Path) -> list[str]:
    return sorted(path.name for path in output_dir.iterdir())


def _assert_success_output_contract(output: Path) -> None:
    assert _public_names(output) == SUCCESS_OUTPUT_NAMES


def _assert_word_no_pdf_output_contract(output: Path) -> None:
    assert _public_names(output) == WORD_NO_PDF_OUTPUT_NAMES


@pytest.fixture(autouse=True)
def _default_pipeline_output_text_harness_pass(monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    def pass_validate_output_text_file(candidate, output_report=None):
        report = {
            "status": "pass",
            "candidate": str(candidate),
            "failures": [],
            "counts": {},
        }
        if output_report is not None:
            Path(output_report).parent.mkdir(parents=True, exist_ok=True)
            Path(output_report).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report

    monkeypatch.setattr(pipeline, "validate_output_text_file", pass_validate_output_text_file)


def test_run_pipeline_success_writes_clean_contract_and_removes_work_dir(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.docx"
    output = tmp_path / "output"
    _write_source_docx(source, with_image=True)
    output.mkdir()
    (output / "old-debug.json").write_text("{}", encoding="utf-8")
    (output / "image").mkdir()
    (output / "image" / "manifest.json").write_text("old process file", encoding="utf-8")
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "needs_review"
    _assert_success_output_contract(output)
    ok, messages = validate_clean_output(output)
    assert ok is True
    assert messages == []
    assert not (output / "old-debug.json").exists()
    assert not (output / "image" / "manifest.json").exists()
    assert any((output / "template_diff").glob("*.png"))
    assert not list(tmp_path.glob("output_work*"))
    assert any("figure" in item.lower() for item in report["manual_review"])
    assert any("tex" in item.lower() for item in report["manual_review"])
    assert "image/" in (output / "thesis.tex").read_text(encoding="utf-8")


def test_run_pipeline_model_harness_failure_blocks_word_generation(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.docx"
    output = tmp_path / "output"
    _write_source_docx(source)
    poisoned_model = ThesisModel(
        sections=[
            ContentBlock(
                id="bad-template-leak",
                type="chapter",
                title="论文封面书脊",
                text="王小亮 MERGEFORMAT",
                level=1,
            )
        ]
    )

    monkeypatch.setattr(pipeline, "extract_thesis_model", lambda _path, _work_dir: poisoned_model)

    def unexpected_apply_word_fixes(_state):
        raise AssertionError("apply_word_fixes must not run after model harness failure")

    monkeypatch.setattr(pipeline, "apply_word_fixes", unexpected_apply_word_fixes)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "failed"
    assert not (output / "thesis.docx").exists()
    model_report = json.loads((output / "harness" / "model_validation_report.json").read_text(encoding="utf-8"))
    status = json.loads((output / "harness" / "status.json").read_text(encoding="utf-8"))
    assert model_report["status"] == "failed"
    assert status["failed_stage"] == "model"
    assert any("harness_model_failed" in item for item in report["blocking_items"])


def test_run_pipeline_output_text_harness_failure_blocks_pdf_export(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.docx"
    output = tmp_path / "output"
    _write_source_docx(source)

    def fake_validate_output_text_file(_candidate, output_report=None):
        report = {
            "status": "failed",
            "candidate": str(_candidate),
            "failures": [{"id": "body_contains_template_page_number_48"}],
            "counts": {},
        }
        if output_report is not None:
            Path(output_report).parent.mkdir(parents=True, exist_ok=True)
            Path(output_report).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report

    def unexpected_pdf_export(_source, _target):
        raise AssertionError("PDF export must not run after output_text harness failure")

    monkeypatch.setattr(pipeline, "validate_output_text_file", fake_validate_output_text_file, raising=False)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", unexpected_pdf_export)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "failed"
    assert (output / "thesis.docx").is_file()
    assert not (output / "thesis.pdf").exists()
    output_text_report = json.loads((output / "harness" / "output_text_report.json").read_text(encoding="utf-8"))
    status = json.loads((output / "harness" / "status.json").read_text(encoding="utf-8"))
    assert output_text_report["status"] == "failed"
    assert status["failed_stage"] == "output_text"
    assert any("harness_output_text_failed" in item for item in report["blocking_items"])


def test_run_pipeline_pdf_failure_reports_blocking_without_process_files(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.docx"
    output = tmp_path / "output"
    _write_source_docx(source)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _failing_pdf_export)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "failed"
    _assert_word_no_pdf_output_contract(output)
    assert (output / "thesis.docx").is_file()
    assert (output / "thesis.tex").is_file()
    assert not (output / "thesis.pdf").exists()
    assert any("pdf export" in item.lower() for item in report["blocking_items"])
    assert any("thesis.pdf" in item for item in report["blocking_items"])
    allowed_json = {
        "artifact_identity.json",
        "bad_fixture_regression_report.json",
        "evidence_packet.json",
        "failure_queue.json",
        "gate_board.json",
        "instrumented_template_report.json",
        "model.json",
        "template_inheritance_report.json",
        "model_validation_report.json",
        "output_text_report.json",
        "role_quiz_report.json",
        "status.json",
    }
    assert not any(
        path.name.endswith((".log", ".tmp")) or (path.suffix == ".json" and path.name not in allowed_json)
        for path in output.rglob("*")
        if path.is_file()
    )
    report_text = (output / "report.md").read_text(encoding="utf-8")
    assert "Word PDF export unavailable" in report_text
    assert "Missing required output: thesis.pdf" in report_text


def test_run_pipeline_keep_work_retains_adjacent_work_dir_and_keeps_public_output_clean(
    tmp_path, monkeypatch
):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.docx"
    output = tmp_path / "output"
    _write_source_docx(source, with_image=True)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    report = pipeline.run_pipeline(source, output, keep_work=True)

    work_dir = output.with_name(output.name + "_work")
    assert report["status"] == "needs_review"
    assert work_dir.is_dir()
    assert (work_dir / "source.docx").is_file()
    _assert_success_output_contract(output)
    ok, messages = validate_clean_output(output)
    assert ok is True
    assert messages == []


def test_run_pipeline_text_pdf_input_writes_clean_contract(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.pdf"
    output = tmp_path / "output"
    _write_text_pdf(source)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "needs_review"
    _assert_success_output_contract(output)
    ok, messages = validate_clean_output(output)
    assert ok is True
    assert messages == []
    assert any("pdf" in item.lower() and "layout" in item.lower() for item in report["manual_review"])
    assert "PDF Pipeline Thesis" in (output / "thesis.tex").read_text(encoding="utf-8")
    assert report["metadata"]["student_id"]["value"] == "20370001"
    assert report["metadata"]["student_id"]["evidence"]["method"] == "pdf-text-label"
    assert report["equation_ledger"] == []
    report_text = (output / "report.md").read_text(encoding="utf-8")
    assert "student_id: 20370001" in report_text
    assert "## Equation Ledger" in report_text


def test_run_pipeline_strict_mode_fails_when_review_items_remain(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.pdf"
    output = tmp_path / "output"
    _write_text_pdf(source)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    report = pipeline.run_pipeline(source, output, strict=True)

    assert report["status"] == "failed"
    assert any("strict_finalization_failed" in item for item in report["blocking_items"])
    _assert_success_output_contract(output)
    assert "strict_finalization_failed" in (output / "report.md").read_text(encoding="utf-8")


def test_run_pipeline_pdf_input_generates_editable_template_docx(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.pdf"
    output = tmp_path / "output"
    _write_text_pdf(source)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "needs_review"
    assert not any("visual-preserving" in note for note in report["notes"])
    with zipfile.ZipFile(output / "thesis.docx") as docx_zip:
        document_xml = docx_zip.read("word/document.xml").decode("utf-8")

    assert "PDF Pipeline Thesis" in document_xml
    assert "This PDF contains extractable thesis text." in document_xml
    assert "20370001" in document_xml
    assert "PDF Extracted Text" not in document_xml
    assert report["editability"]["editable_characters"] > 0
    assert report["editability"]["paragraph_count"] > 0
    assert report["editability"]["page_screenshot_drawing_count"] == 0
    assert report["editability"]["body_snippet_count"] > 0
    assert report["editability"]["body_snippet_hits"] > 0
    assert "## Editability Audit" in (output / "report.md").read_text(encoding="utf-8")


def test_run_pipeline_pdf_table_becomes_editable_word_table(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source-table.pdf"
    output = tmp_path / "output"
    _write_text_pdf_with_table(source)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "needs_review"
    assert report["summary"]["tables"] == 1
    document = Document(output / "thesis.docx")
    table_text = [
        [cell.text for cell in row.cells]
        for table in document.tables
        for row in table.rows
    ]
    assert ["Metric", "Value"] in table_text
    assert ["Accuracy", "98%"] in table_text
    assert ["Recall", "95%"] in table_text
    tex = (output / "thesis.tex").read_text(encoding="utf-8")
    assert "\\begin{table}" in tex
    assert "Accuracy & 98\\%" in tex


def test_run_pipeline_pdf_equation_writes_editable_omml_and_ledger(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source-equation.pdf"
    output = tmp_path / "output"
    _write_text_pdf_with_equation(source)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    report = pipeline.run_pipeline(source, output)

    assert report["summary"]["equations"] == 1
    assert report["equation_ledger"][0]["kind"] == "pdf-text-equation"
    assert report["equation_ledger"][0]["status"] == "editable_omml"
    assert report["equation_ledger"][0]["number"] == "(2.1)"
    assert report["equation_ledger"][0]["editable_in_word"] is True
    assert not any("Equation pdf-equation-1 requires review" in item for item in report["manual_review"])
    docx_text = Document(output / "thesis.docx").paragraphs
    assert not any("[Equation requires review]" in paragraph.text for paragraph in docx_text)
    with zipfile.ZipFile(output / "thesis.docx") as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "<m:oMathPara" in document_xml
    assert "<m:sSub>" in document_xml
    report_text = (output / "report.md").read_text(encoding="utf-8")
    assert "status=editable_omml" in report_text
    assert "pdf-text-equation" in report_text


def test_run_pipeline_scanned_pdf_keeps_page_image_as_ocr_evidence_not_word_screenshot(
    tmp_path, monkeypatch
):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "scanned.pdf"
    output = tmp_path / "output"
    _write_blank_pdf(source)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "needs_review"
    assert (output / "image" / "pdf-page-001.png").is_file()
    assert report["ocr_ledger"] == [
        {
            "page": 1,
            "status": "needs_ocr",
            "image_path": str((output / "image" / "pdf-page-001.png").resolve(strict=False)),
            "text_characters": 0,
            "confidence": 0.0,
            "requires_review": True,
        }
    ]
    assert any("ocr" in item.lower() for item in report["manual_review"])
    with zipfile.ZipFile(output / "thesis.docx") as docx_zip:
        document_xml = docx_zip.read("word/document.xml").decode("utf-8")
        package_names = docx_zip.namelist()
    assert report["editability"]["page_screenshot_drawing_count"] == 0
    assert not any("pdf-page" in name for name in package_names)
    assert "[Figure inserted]" not in document_xml
    assert "OCR/manual transcription required" not in document_xml
    tex = (output / "thesis.tex").read_text(encoding="utf-8")
    assert "pdf-page-001.png" in tex
    assert "\\includegraphics" not in tex
    report_text = (output / "report.md").read_text(encoding="utf-8")
    assert "## OCR Ledger" in report_text
    assert "page=1" in report_text
    assert "pdf-page-001.png" in report_text


def test_run_pipeline_scanned_pdf_with_ocr_writes_editable_word_text(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline
    from buaa_thesis_kit.pdf_extract import OcrResult, extract_pdf_model as real_extract_pdf_model

    source = tmp_path / "scanned-with-ocr.pdf"
    output = tmp_path / "output"
    _write_blank_pdf(source)

    def fake_extract_pdf_model(path: Path, work_dir: Path):
        return real_extract_pdf_model(
            path,
            work_dir,
            ocr_engine=lambda _image_path: OcrResult(
                text="1 Introduction\nOCR recovered editable body text.",
                confidence=0.86,
                engine="fake-ocr",
            ),
        )

    monkeypatch.setattr(pipeline, "extract_pdf_model", fake_extract_pdf_model)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "needs_review"
    assert report["ocr_ledger"][0]["status"] == "ocr_text_extracted"
    assert report["ocr_ledger"][0]["text_characters"] == len(
        "1 Introduction\nOCR recovered editable body text."
    )
    assert (output / "image" / "pdf-page-001.png").is_file()
    with zipfile.ZipFile(output / "thesis.docx") as docx_zip:
        document_xml = docx_zip.read("word/document.xml").decode("utf-8")
    assert "OCR recovered editable body text." in document_xml
    assert "OCR/manual transcription required" not in document_xml
    assert "OCR recovered editable body text." in (output / "thesis.tex").read_text(encoding="utf-8")


def test_run_pipeline_doc_input_converts_to_docx_before_extraction(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    docx_source = tmp_path / "source-as-docx.docx"
    source = tmp_path / "source.doc"
    output = tmp_path / "output"
    _write_source_docx(docx_source)
    source.write_bytes(docx_source.read_bytes())
    monkeypatch.setattr(pipeline, "convert_doc_to_docx", _stub_doc_conversion)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] in {"pass", "needs_review"}
    _assert_success_output_contract(output)
    ok, messages = validate_clean_output(output)
    assert ok is True
    assert messages == []
    assert any("doc to docx" in note.lower() for note in report["notes"])


def test_run_pipeline_doc_conversion_failure_writes_failed_report(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.doc"
    output = tmp_path / "output"
    source.write_bytes(b"legacy word bytes")
    monkeypatch.setattr(pipeline, "convert_doc_to_docx", _failing_doc_conversion)

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "failed"
    assert "image" in _public_names(output)
    assert "report.md" in _public_names(output)
    assert any("doc conversion" in item.lower() for item in report["blocking_items"])
    assert "Word DOC conversion unavailable" in (output / "report.md").read_text(encoding="utf-8")


def test_copy_final_equation_assets_moves_preview_images_to_public_image_dir(tmp_path):
    import buaa_thesis_kit.pipeline as pipeline

    preview = tmp_path / "work" / "equation-preview" / "formula.png"
    preview.parent.mkdir(parents=True)
    preview.write_bytes(TINY_PNG)
    image_dir = tmp_path / "output" / "image"
    model = ThesisModel(
        equations=[
            EquationItem(
                id="eq-preview",
                kind="embedded-object",
                text="equation.bin",
                preview_path=str(preview),
                requires_review=True,
            )
        ]
    )

    notes = pipeline._copy_final_equation_assets(model, image_dir)

    copied = image_dir / "formula.png"
    assert copied.read_bytes() == TINY_PNG
    assert model.equations[0].preview_path == str(copied.resolve(strict=False))
    assert notes == []


def test_equation_report_classifies_editable_and_review_equations(tmp_path):
    import buaa_thesis_kit.pipeline as pipeline

    preview = tmp_path / "work" / "formula.png"
    preview.parent.mkdir(parents=True)
    preview.write_bytes(TINY_PNG)
    object_path = tmp_path / "work" / "equation.bin"
    object_path.write_bytes(b"equation ole payload")
    image_dir = tmp_path / "output" / "image"
    model = ThesisModel(
        equations=[
            EquationItem(
                id="eq-omml",
                kind="omml",
                text="x+y",
                number="(1)",
                omml="<m:oMath/>",
                requires_review=True,
            ),
            EquationItem(
                id="eq-object",
                kind="embedded-object",
                text="equation.bin",
                number="(2)",
                preview_path=str(preview),
                object_path=str(object_path),
                object_xml="<w:object/>",
                requires_review=True,
            ),
            EquationItem(
                id="eq-preview",
                kind="embedded-object",
                text="preview only",
                preview_path=str(preview),
                requires_review=True,
            ),
            EquationItem(
                id="eq-latex-review",
                kind="pdf-text-equation",
                text="x+y",
                latex="x+y",
                requires_review=True,
            ),
            EquationItem(
                id="eq-manual",
                kind="unknown",
                text="manual",
                requires_review=True,
            ),
        ]
    )

    pipeline._copy_final_equation_assets(model, image_dir)
    report = pipeline._equation_report(model)

    by_id = {item["id"]: item for item in report}
    assert by_id["eq-omml"]["status"] == "editable_omml"
    assert by_id["eq-omml"]["editable_in_word"] is True
    assert by_id["eq-object"]["status"] == "editable_ole_object"
    assert by_id["eq-object"]["editable_in_word"] is True
    assert by_id["eq-object"]["preview_path"] == str((image_dir / "formula.png").resolve(strict=False))
    assert by_id["eq-preview"]["status"] == "preview_image_needs_review"
    assert by_id["eq-preview"]["editable_in_word"] is False
    assert by_id["eq-latex-review"]["status"] == "latex_needs_review"
    assert by_id["eq-latex-review"]["editable_in_word"] is False
    assert by_id["eq-manual"]["status"] == "manual_transcription_required"
    assert by_id["eq-manual"]["editable_in_word"] is False


def test_run_pipeline_invalid_pdf_input_writes_failed_report_without_exception(tmp_path):
    import buaa_thesis_kit.pipeline as pipeline

    source = tmp_path / "source.pdf"
    output = tmp_path / "output"
    source.write_bytes(b"%PDF-1.4\nnot a thesis source")

    report = pipeline.run_pipeline(source, output)

    assert report["status"] == "failed"
    assert "image" in _public_names(output)
    assert "report.md" in _public_names(output)
    assert not list(tmp_path.glob("output_work*"))
    assert any("pdf" in item.lower() for item in report["blocking_items"])
    assert "PDF extraction failed" in (output / "report.md").read_text(encoding="utf-8")


def test_run_pipeline_refuses_output_directory_that_contains_source_without_cleanup(tmp_path):
    import buaa_thesis_kit.pipeline as pipeline

    output = tmp_path / "output"
    output.mkdir()
    source = output / "source.docx"
    sentinel = output / "do-not-delete.txt"
    _write_source_docx(source)
    sentinel.write_text("keep", encoding="utf-8")

    with pytest.raises(ValueError, match="contains source file"):
        pipeline.run_pipeline(source, output)

    assert source.exists()
    assert sentinel.read_text(encoding="utf-8") == "keep"


def test_cli_main_prints_json_and_returns_status_based_exit_codes(tmp_path, monkeypatch, capsys):
    import buaa_thesis_kit.pipeline as pipeline
    import scripts.run_pipeline as cli

    source = tmp_path / "source.docx"
    output = tmp_path / "output"
    failed_output = tmp_path / "failed-output"
    failed_source = tmp_path / "source.pdf"
    _write_source_docx(source)
    failed_source.write_bytes(b"%PDF-1.4\nnot a thesis source")
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    pass_code = cli.main([str(source), "--out", str(output)])
    pass_payload = json.loads(capsys.readouterr().out)
    fail_code = cli.main([str(failed_source), "--out", str(failed_output)])
    fail_payload = json.loads(capsys.readouterr().out)

    assert pass_code == 0
    assert pass_payload["status"] in {"pass", "needs_review"}
    assert pass_payload["summary"]["sections"] >= 1
    assert fail_code == 1
    assert fail_payload["status"] == "failed"


def test_cli_help_describes_supported_source_types(capsys):
    import scripts.run_pipeline as cli

    with pytest.raises(SystemExit) as exc_info:
        cli.main(["--help"])

    output = capsys.readouterr().out
    assert exc_info.value.code == 0
    assert "DOC, DOCX, or PDF" in output
    assert "--strict" in output


def test_cli_strict_mode_returns_failure_for_needs_review(tmp_path, monkeypatch, capsys):
    import buaa_thesis_kit.pipeline as pipeline
    import scripts.run_pipeline as cli

    source = tmp_path / "source.pdf"
    output = tmp_path / "output"
    _write_text_pdf(source)
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", _stub_pdf_export)

    code = cli.main([str(source), "--out", str(output), "--strict"])
    payload = json.loads(capsys.readouterr().out)

    assert code == 1
    assert payload["status"] == "failed"
    assert any("strict_finalization_failed" in item for item in payload["blocking_items"])


def test_cli_sample_mode_truncated_is_passed_to_pipeline(tmp_path, monkeypatch, capsys):
    import scripts.run_pipeline as cli

    source = tmp_path / "source.docx"
    source.write_bytes(b"docx")
    output = tmp_path / "output"
    captured = {}

    def fake_run_pipeline(source_arg, output_arg, **kwargs):
        captured["source"] = source_arg
        captured["output"] = output_arg
        captured.update(kwargs)
        return {
            "status": "needs_review",
            "outputs": {},
            "summary": {},
            "blocking_items": [],
            "manual_review": [],
            "notes": [],
        }

    monkeypatch.setattr(cli, "run_pipeline", fake_run_pipeline)

    code = cli.main([str(source), "--out", str(output), "--sample-mode", "truncated"])
    payload = json.loads(capsys.readouterr().out)

    assert code == 0
    assert payload["status"] == "needs_review"
    assert captured["sample_mode"] == "truncated"
