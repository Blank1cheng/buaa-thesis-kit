import json
import os
from pathlib import Path

from docx import Document
from pypdf import PdfWriter


def _write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def test_progress_artifacts_show_gate_board_failure_queue_and_next_action(tmp_path):
    from buaa_thesis_kit.harness.progress import write_progress_artifacts

    harness = tmp_path / "harness"
    candidate = tmp_path / "thesis.docx"
    Document().save(candidate)
    _write_json(
        harness / "output_text_report.json",
        {
            "status": "failed",
            "candidate_path": str(candidate),
            "candidate_sha256": "abc",
            "candidate_size": candidate.stat().st_size,
            "failures": [
                {
                    "id": "template_instructions_or_sample_leak",
                    "tokens": ["论文封面书脊"],
                    "region": "spine",
                }
            ],
        },
    )
    _write_json(harness / "status.json", {"status": "failed", "failed_stage": "output_text", "stages": []})

    packet = write_progress_artifacts(
        harness,
        candidate_path=candidate,
        sample_mode="truncated",
        commands=["python scripts/run_harness.py --candidate thesis.docx"],
    )

    gate_board = json.loads((harness / "gate_board.json").read_text(encoding="utf-8"))
    failure_queue = json.loads((harness / "failure_queue.json").read_text(encoding="utf-8"))
    evidence_packet = json.loads((harness / "evidence_packet.json").read_text(encoding="utf-8"))
    progress = (harness / "progress.md").read_text(encoding="utf-8")
    assert packet["remaining_failures"] == ["H-001"]
    assert gate_board["overall_status"] == "failed"
    assert any(gate["name"] == "pre_finalize_output_text" and gate["status"] == "failed" for gate in gate_board["gates"])
    assert failure_queue["failures"][0]["id"] == "H-001"
    assert failure_queue["failures"][0]["reason"] == "template_instructions_or_sample_leak"
    assert failure_queue["failures"][0]["region"] == "spine"
    assert evidence_packet["commands"] == ["python scripts/run_harness.py --candidate thesis.docx"]
    assert "Next Failure To Fix" in progress
    assert "/fix H-001" in progress


def test_progress_artifacts_include_rejection_identity_metadata(tmp_path):
    from buaa_thesis_kit.harness.artifact_identity import artifact_identity
    from buaa_thesis_kit.harness.progress import write_progress_artifacts

    harness = tmp_path / "harness"
    source_candidate = tmp_path / "output" / "pipeline_gate" / "thesis.docx"
    bad_fixture = tmp_path / "tests" / "fixtures" / "bad_outputs" / "abc123_thesis.docx"
    source_candidate.parent.mkdir(parents=True, exist_ok=True)
    bad_fixture.parent.mkdir(parents=True, exist_ok=True)
    Document().save(source_candidate)
    bad_fixture.write_bytes(source_candidate.read_bytes())
    identity = artifact_identity(candidate_path=bad_fixture)
    identity.update(
        {
            "source_candidate_path": str(source_candidate),
            "bad_fixture_path": str(bad_fixture),
            "commit": "test-commit",
        }
    )
    _write_json(harness / "artifact_identity.json", identity)

    packet = write_progress_artifacts(harness, candidate_path=bad_fixture, sample_mode="truncated")

    gate_board = json.loads((harness / "gate_board.json").read_text(encoding="utf-8"))
    failure_queue = json.loads((harness / "failure_queue.json").read_text(encoding="utf-8"))
    assert gate_board["source_candidate_path"] == str(source_candidate)
    assert gate_board["bad_fixture_path"] == str(bad_fixture)
    assert gate_board["candidate_size"] == bad_fixture.stat().st_size
    assert packet["source_candidate_path"] == str(source_candidate)
    assert packet["bad_fixture_path"] == str(bad_fixture)
    assert packet["candidate_size"] == bad_fixture.stat().st_size
    assert failure_queue["source_candidate_path"] == str(source_candidate)
    assert failure_queue["bad_fixture_path"] == str(bad_fixture)


def test_progress_artifacts_enrich_existing_status_with_identity(tmp_path):
    from buaa_thesis_kit.harness.progress import write_progress_artifacts

    harness = tmp_path / "harness"
    candidate = tmp_path / "thesis.docx"
    Document().save(candidate)
    _write_json(harness / "status.json", {"status": "failed", "failed_stage": "render_smoke", "stages": []})

    write_progress_artifacts(harness, candidate_path=candidate, sample_mode="truncated")

    status = json.loads((harness / "status.json").read_text(encoding="utf-8"))
    assert status["candidate_path"] == str(candidate)
    assert status["candidate_sha256"]
    assert status["candidate_size"] == candidate.stat().st_size
    assert status["commit"]


def test_progress_artifacts_fail_when_report_sha_differs_from_candidate(tmp_path):
    from buaa_thesis_kit.harness.progress import write_progress_artifacts

    harness = tmp_path / "harness"
    candidate = tmp_path / "bad.docx"
    Document().save(candidate)
    _write_json(
        harness / "artifact_identity.json",
        {
            "candidate_path": str(candidate),
            "candidate_sha256": "0" * 64,
            "candidate_size": candidate.stat().st_size,
            "source_candidate_path": "output/pipeline_gate/thesis.docx",
            "bad_fixture_path": str(candidate),
            "commit": "test-commit",
        },
    )

    write_progress_artifacts(harness, candidate_path=candidate, sample_mode="truncated")

    gate_board = json.loads((harness / "gate_board.json").read_text(encoding="utf-8"))
    failure_queue = json.loads((harness / "failure_queue.json").read_text(encoding="utf-8"))
    by_name = {gate["name"]: gate for gate in gate_board["gates"]}
    assert by_name["artifact_identity"]["status"] == "failed"
    assert failure_queue["failures"][0]["reason"] == "artifact_identity_mismatch"


def test_progress_schema_includes_paths_and_interaction_evidence(tmp_path, capsys):
    from buaa_thesis_kit.harness.progress import write_progress_artifacts
    from scripts import harness_status, harness_triage

    harness = tmp_path / "harness"
    candidate = tmp_path / "thesis.docx"
    model = tmp_path / "model.json"
    template = tmp_path / "template.docx"
    Document().save(candidate)
    model.write_text("{}", encoding="utf-8")
    Document().save(template)
    _write_json(
        harness / "output_text_report.json",
        {
            "status": "failed",
            "candidate_path": str(candidate),
            "failures": [{"id": "cn_abstract_contains_english_title", "actual": "Research on", "region": "abstract_cn"}],
        },
    )

    write_progress_artifacts(
        harness,
        candidate_path=candidate,
        model_path=model,
        template_path=template,
        sample_mode="truncated",
    )

    gate_board = json.loads((harness / "gate_board.json").read_text(encoding="utf-8"))
    assert gate_board["template_path"] == str(template)
    assert gate_board["template_sha256"]
    assert gate_board["model_path"] == str(model)
    assert gate_board["model_sha256"]

    assert harness_status.main(["--out", str(harness)]) == 0
    status_payload = json.loads(capsys.readouterr().out)
    assert status_payload["candidate_sha256"] == gate_board["candidate_sha256"]

    assert harness_triage.main(["--out", str(harness), "--evidence", "H-001"]) == 0
    evidence_payload = json.loads(capsys.readouterr().out)
    assert evidence_payload["failure"]["id"] == "H-001"
    assert evidence_payload["candidate_sha256"] == gate_board["candidate_sha256"]
    assert evidence_payload["reports"]["failure_queue"].endswith("failure_queue.json")
    assert evidence_payload["artifacts"]["docx"] == str(candidate)


def test_harness_interaction_document_exists():
    doc = Path("docs/HARNESS_INTERACTION.md")

    assert doc.is_file()
    text = doc.read_text(encoding="utf-8")
    for command in ("/status", "/triage", "/fix H-xxx", "/evidence H-xxx", "/reject"):
        assert command in text
    assert "Do not opportunistically fix unrelated issues" in text


def test_pipeline_writes_progress_artifacts_and_render_smoke_report(tmp_path, monkeypatch):
    import buaa_thesis_kit.pipeline as pipeline

    def stub_pdf_export(_source: Path, target: Path) -> tuple[bool, str]:
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with target.open("wb") as handle:
            writer.write(handle)
        return True, "stubbed pdf export"

    output = tmp_path / "pipeline_gate"
    monkeypatch.setattr(pipeline, "export_pdf_from_docx", stub_pdf_export)

    report = pipeline.run_pipeline(
        Path("tests/fixtures/truncated_input.docx"),
        output,
        sample_mode="truncated",
    )

    assert report["status"] in {"needs_review", "failed", "pass"}
    assert (output / "harness" / "gate_board.json").is_file()
    assert (output / "harness" / "failure_queue.json").is_file()
    assert (output / "harness" / "evidence_packet.json").is_file()
    assert (output / "harness" / "progress.md").is_file()
    assert (output / "render_smoke" / "report.json").is_file()
    gate_board = json.loads((output / "harness" / "gate_board.json").read_text(encoding="utf-8"))
    by_name = {gate["name"]: gate for gate in gate_board["gates"]}
    assert by_name["role_quiz"]["status"] == "pass"
    assert by_name["instrumented_template"]["status"] == "pass"
    assert by_name["bad_fixture_regression"]["status"] == "skipped"


def test_harness_doctor_prefers_newer_pipeline_gate_over_stale_root_harness(tmp_path):
    from scripts.harness_doctor import run_doctor

    root = tmp_path / "output"
    stale_harness = root / "harness"
    pipeline_harness = root / "pipeline_gate" / "harness"
    stale_docx = tmp_path / "stale.docx"
    current_docx = root / "pipeline_gate" / "thesis.docx"
    Document().save(stale_docx)
    current_docx.parent.mkdir(parents=True, exist_ok=True)
    Document().save(current_docx)
    _write_json(
        stale_harness / "status.json",
        {"status": "failed", "candidate_path": str(stale_docx), "stages": []},
    )
    _write_json(
        stale_harness / "output_text_report.json",
        {"status": "failed", "candidate": str(stale_docx), "failures": [{"id": "stale_failure"}]},
    )
    _write_json(
        pipeline_harness / "status.json",
        {"status": "pass", "candidate_path": str(current_docx), "stages": []},
    )
    _write_json(
        pipeline_harness / "output_text_report.json",
        {"status": "pass", "candidate": str(current_docx), "failures": []},
    )
    os.utime(stale_harness / "status.json", (1, 1))
    os.utime(pipeline_harness / "status.json", (2, 2))

    packet = run_doctor(out_dir=stale_harness, sample_mode="truncated")

    queue = json.loads((stale_harness / "failure_queue.json").read_text(encoding="utf-8"))
    assert packet["candidate_path"] == str(current_docx)
    assert not any(item["reason"] == "stale_failure" for item in queue["failures"])


def test_harness_doctor_custom_out_uses_its_own_reports_even_when_pipeline_gate_exists(tmp_path):
    from scripts.harness_doctor import run_doctor

    root = tmp_path / "output"
    pipeline_harness = root / "pipeline_gate" / "harness"
    custom_harness = root / "harness_thesis10"
    pipeline_docx = root / "pipeline_gate" / "thesis.docx"
    custom_docx = tmp_path / "thesis10.docx"
    pipeline_docx.parent.mkdir(parents=True, exist_ok=True)
    Document().save(pipeline_docx)
    Document().save(custom_docx)
    _write_json(
        pipeline_harness / "status.json",
        {"status": "pass", "candidate_path": str(pipeline_docx), "stages": []},
    )
    _write_json(root / "template_inheritance_report.json", {"status": "pass", "failures": []})
    _write_json(
        custom_harness / "status.json",
        {"status": "failed", "candidate_path": str(custom_docx), "stages": []},
    )
    _write_json(
        custom_harness / "output_text_report.json",
        {"status": "failed", "candidate": str(custom_docx), "failures": [{"id": "custom_failure"}]},
    )
    render_dir = custom_harness / "render_smoke"
    render_dir.mkdir(parents=True, exist_ok=True)
    (render_dir / "page_001_cover.png").write_bytes(b"png")
    _write_json(render_dir / "report.json", {"status": "pass", "failures": []})

    packet = run_doctor(out_dir=custom_harness, sample_mode="truncated")

    queue = json.loads((custom_harness / "failure_queue.json").read_text(encoding="utf-8"))
    gate_board = json.loads((custom_harness / "gate_board.json").read_text(encoding="utf-8"))
    by_name = {gate["name"]: gate for gate in gate_board["gates"]}
    assert packet["candidate_path"] == str(custom_docx)
    assert packet["artifacts"]["page_images"] == [str(render_dir / "page_001_cover.png")]
    assert by_name["template_inheritance"]["report"] is None
    assert any(item["reason"] == "custom_failure" for item in queue["failures"])


def test_harness_doctor_uses_candidate_adjacent_render_smoke_when_root_harness_exists(tmp_path):
    from scripts.harness_doctor import run_doctor

    root = tmp_path / "output"
    harness = root / "harness"
    candidate = root / "pipeline_gate" / "thesis.docx"
    render_dir = candidate.parent / "render_smoke"
    candidate.parent.mkdir(parents=True, exist_ok=True)
    Document().save(candidate)
    _write_json(
        harness / "status.json",
        {"status": "pass", "candidate_path": str(candidate), "stages": []},
    )
    _write_json(
        harness / "output_text_report.json",
        {"status": "pass", "candidate": str(candidate), "failures": []},
    )
    _write_json(
        render_dir / "report.json",
        {
            "status": "failed",
            "candidate": str(candidate),
            "sample_mode": "truncated",
            "failures": [
                {
                    "id": "cover_metadata_anchor_misaligned",
                    "region": "cover",
                    "evidence_text": "学    号     17375303",
                }
            ],
            "artifacts": {"page_images": []},
        },
    )

    run_doctor(out_dir=harness, sample_mode="truncated")

    gate_board = json.loads((harness / "gate_board.json").read_text(encoding="utf-8"))
    queue = json.loads((harness / "failure_queue.json").read_text(encoding="utf-8"))
    by_name = {gate["name"]: gate for gate in gate_board["gates"]}
    assert by_name["render_smoke"]["status"] == "failed"
    assert Path(by_name["render_smoke"]["report"]).parts[-2:] == ("render_smoke", "report.json")
    assert queue["failures"][0]["reason"] == "cover_metadata_anchor_misaligned"
