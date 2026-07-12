from __future__ import annotations

import hashlib
import json
from pathlib import Path

import fitz
import pytest
from PIL import Image
from docx import Document
from docx.shared import Inches

import buaa_thesis_kit.delivery_packaging as delivery_packaging
from buaa_thesis_kit.delivery_packaging import (
    flatten_latex_document,
    package_agent_delivery,
)
from buaa_thesis_kit.harness.delivery import validate_delivery
from scripts import package_agent_delivery as package_cli


REQUIRED_GATES = ("G20", "G21", "G22", "G23", "G24", "G25", "G27", "G28")
REQUIRED_REGIONS = (
    "cover",
    "spine",
    "taskbook",
    "declaration",
    "abstract_cn",
    "abstract_en",
    "toc",
    "chapter_openers",
    "figures_tables_equations",
    "references",
    "acknowledgement_appendix",
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_pdf(path: Path, page_count: int) -> None:
    document = fitz.open()
    for _ in range(page_count):
        document.new_page(width=595, height=842)
    document.save(path)
    document.close()


def _write_png(path: Path, color: str = "white") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (32, 32), color=color).save(path, format="PNG")


def _write_pdf_screenshot(
    pdf_path: Path,
    *,
    page_number: int,
    bbox: list[float],
    destination: Path,
) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with fitz.open(pdf_path) as document:
        pixmap = document[page_number - 1].get_pixmap(
            matrix=fitz.Matrix(2, 2),
            clip=fitz.Rect(bbox),
            alpha=False,
        )
        pixmap.save(destination)


def _source_record(path: Path) -> dict[str, object]:
    return {
        "candidate_path": str(path.resolve()),
        "source_sha256": _sha256(path),
        "source_size": path.stat().st_size,
        "source_type": path.suffix.lower().lstrip("."),
    }


def test_pipeline_source_identity_rejects_relabelled_candidate(tmp_path):
    source_a = tmp_path / "candidate-a.docx"
    source_b = tmp_path / "candidate-b.docx"
    source_a.write_bytes(b"candidate A")
    source_b.write_bytes(b"candidate B")
    recorded = _source_record(source_a)

    with pytest.raises(ValueError, match="artifact_identity_mismatch"):
        delivery_packaging.validate_pipeline_source_identity(
            source_b,
            model={"source": recorded},
            pipeline_report={"source_identity": recorded},
            extraction_identity={
                "source_path": recorded["candidate_path"],
                "sha256": recorded["source_sha256"],
                "size": recorded["source_size"],
            },
        )


def test_replace_target_rejects_arbitrary_or_protected_directories(tmp_path):
    arbitrary = tmp_path / "valuable-project"
    arbitrary.mkdir()
    sentinel = arbitrary / "do-not-delete.txt"
    sentinel.write_text("keep", encoding="utf-8")

    with pytest.raises(ValueError, match="unsafe_output_replace_target"):
        delivery_packaging.validate_replace_target(arbitrary, protected_paths=[])
    assert sentinel.read_text(encoding="utf-8") == "keep"

    output = tmp_path / "output"
    protected_run = output / "pipeline-run"
    protected_run.mkdir(parents=True)
    with pytest.raises(ValueError, match="unsafe_output_replace_target"):
        delivery_packaging.validate_replace_target(
            output,
            protected_paths=[protected_run],
        )
    assert protected_run.is_dir()


def test_semantic_docx_validation_rejects_fake_or_page_screenshot_word(tmp_path):
    fake = tmp_path / "fake.docx"
    fake.write_bytes(b"not a Word package")
    with pytest.raises(ValueError, match="invalid_editable_docx"):
        delivery_packaging.validate_semantic_docx(fake, source_type="pdf")

    screenshot = tmp_path / "page.png"
    _write_png(screenshot)
    image_word = tmp_path / "image-only.docx"
    document = Document()
    document.add_paragraph("Editable text is present but the source page is still rasterized.")
    document.add_picture(str(screenshot), width=Inches(6), height=Inches(8))
    document.save(image_word)

    with pytest.raises(ValueError, match="page_screenshot"):
        delivery_packaging.validate_semantic_docx(image_word, source_type="pdf")


def test_flattened_latex_verification_rejects_pdf_pixel_mismatch(tmp_path):
    staging = tmp_path / "staging"
    (staging / "image").mkdir(parents=True)
    (staging / "thesis.tex").write_text(
        "\\documentclass{book}\\begin{document}x\\end{document}",
        encoding="utf-8",
    )
    template = tmp_path / "template"
    template.mkdir()
    pipeline_pdf = tmp_path / "pipeline.pdf"
    _write_pdf(pipeline_pdf, 1)

    def compile_different_pdf(workdir: Path, entry: str) -> dict:
        rebuilt = fitz.open()
        page = rebuilt.new_page(width=595, height=842)
        page.insert_text((72, 72), "different output")
        rebuilt.save(workdir / "thesis.pdf")
        rebuilt.close()
        return {
            "status": "success",
            "pdf_exists": True,
            "pdf_path": str(workdir / "thesis.pdf"),
            "engine": "fake-xelatex",
        }

    with pytest.raises(ValueError, match="flattened_latex_pdf_mismatch"):
        delivery_packaging.verify_flattened_latex_equivalence(
            staging,
            pipeline_pdf=pipeline_pdf,
            template_dir=template,
            compiler=compile_different_pdf,
        )


def test_flatten_latex_document_inlines_generated_tree_and_rewrites_assets(tmp_path):
    workdir = tmp_path / "workdir"
    (workdir / "data" / "bachelor").mkdir(parents=True)
    (workdir / "assets" / "figures").mkdir(parents=True)
    _write_png(workdir / "assets" / "figures" / "figure.png")
    (workdir / "thesis.tex").write_text(
        "\\documentclass{book}\n"
        "\\begin{document}\n"
        "\\include{data/front}\n"
        "\\input{data/body}\n"
        "\\end{document}\n",
        encoding="utf-8",
    )
    (workdir / "data" / "front.tex").write_text(
        "front\n\\input{bachelor/meta}\n", encoding="utf-8"
    )
    (workdir / "data" / "bachelor" / "meta.tex").write_text(
        "metadata\n", encoding="utf-8"
    )
    (workdir / "data" / "body.tex").write_text(
        "body\n\\includegraphics[width=1cm]{assets/figures/figure.png}\n",
        encoding="utf-8",
    )

    result = flatten_latex_document(workdir / "thesis.tex", workdir)

    assert "\\include{data/front}" not in result.text
    assert "\\input{data/body}" not in result.text
    assert "front" in result.text
    assert "metadata" in result.text
    assert "body" in result.text
    assert result.text.count("\\clearpage") == 2
    assert "\\includegraphics[width=1cm]{image/figure.png}" in result.text
    assert [(item.source, item.destination) for item in result.assets] == [
        (workdir / "assets" / "figures" / "figure.png", "figure.png")
    ]


def test_flatten_latex_document_rejects_include_outside_workdir(tmp_path):
    workdir = tmp_path / "workdir"
    workdir.mkdir()
    (tmp_path / "outside.tex").write_text("outside", encoding="utf-8")
    (workdir / "thesis.tex").write_text(
        "\\input{../outside}\n", encoding="utf-8"
    )

    try:
        flatten_latex_document(workdir / "thesis.tex", workdir)
    except ValueError as exc:
        assert "outside source root" in str(exc)
    else:
        raise AssertionError("path traversal must be rejected")


def test_package_agent_delivery_produces_flat_bundle_accepted_by_harness(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "candidate.docx"
    source_document = Document()
    source_document.add_paragraph("Editable semantic thesis source document.")
    source_document.save(source)

    run_dir = tmp_path / "run"
    workdir = run_dir / "workdir"
    (workdir / "data").mkdir(parents=True)
    (workdir / "assets" / "figures").mkdir(parents=True)
    _write_png(workdir / "assets" / "figures" / "figure.png", color="blue")
    (workdir / "thesis.tex").write_text(
        "\\documentclass{buaathesis}\n"
        "\\begin{document}\n"
        "\\include{data/body}\n"
        "\\end{document}\n",
        encoding="utf-8",
    )
    (workdir / "data" / "body.tex").write_text(
        "\\includegraphics{assets/figures/figure.png}\n", encoding="utf-8"
    )
    (workdir / "compile.log").write_text(
        "Latexmk: This is Latexmk, John Collins, 9 March 2026. Version 4.88.\n"
        "This is XeTeX, Version 3.141592653-2.6-0.999998 (TeX Live 2026)\n"
        "Document Class: buaathesis 2020/03/05 v0.9 The LaTeX template for thesis of BUAA\n",
        encoding="utf-8",
    )
    _write_pdf(run_dir / "thesis.pdf", len(REQUIRED_REGIONS))
    source_identity = _source_record(source)
    _write_json(
        run_dir / "model.json",
        {"source": source_identity, "metadata": {"title_cn": "测试论文"}},
    )
    _write_json(
        run_dir / "template_inspection.json",
        {"clone_url": "https://github.com/BHOSC/BUAAthesis"},
    )
    _write_json(
        run_dir / "report.json",
        {
            "status": "pass",
            "source_identity": source_identity,
            "compile_status": {"status": "success", "engine": "latexmk-xelatex"},
            "equations": {"total": 2, "native_latex": 2, "needs_review": 0},
            "figures": {"images_extracted": 1},
        },
    )
    _write_json(run_dir / "harness" / "failure_queue.json", [])
    _write_json(
        run_dir / "harness" / "source_identity_report.json",
        {
            "source_path": source_identity["candidate_path"],
            "sha256": source_identity["source_sha256"],
            "size": source_identity["source_size"],
        },
    )
    gate_board = run_dir / "harness" / "gate_board.json"
    _write_json(gate_board, {gate: {"status": "pass"} for gate in REQUIRED_GATES})

    review_dir = tmp_path / "review"
    reviews = []
    for page, region in enumerate(REQUIRED_REGIONS, start=1):
        screenshot = f"page_{page:03d}_{region}.png"
        bbox = [0, 0, 100, 100]
        _write_pdf_screenshot(
            run_dir / "thesis.pdf",
            page_number=page,
            bbox=bbox,
            destination=review_dir / screenshot,
        )
        reviews.append(
            {
                "region": region,
                "pages": [page],
                "screenshot": screenshot,
                "status": "pass",
                "checks": ["content", "layout"],
                "bbox": bbox,
                "failure_ids": [],
            }
        )
    visual_manifest = review_dir / "visual_review.json"
    _write_json(
        visual_manifest,
        {
            "pdf_sha256": _sha256(run_dir / "thesis.pdf"),
            "pdf_page_count": len(REQUIRED_REGIONS),
            "reviews": reviews,
        },
    )

    template_dir = tmp_path / "template"
    template_dir.mkdir()
    (template_dir / "buaathesis.cls").write_text(
        "\\setmainfont{Times New Roman}\n"
        "\\setCJKfamilyfont{songti}{SimSun}\n"
        "\\setCJKfamilyfont{heiti}{SimHei}\n",
        encoding="utf-8",
    )
    output = tmp_path / "output"

    verification_calls: list[tuple[Path, Path, Path]] = []

    def verify_flattened(staging_dir: Path, *, pipeline_pdf: Path, template_dir: Path):
        verification_calls.append((staging_dir, pipeline_pdf, template_dir))
        return {
            "status": "pass",
            "engine": "fake-xelatex",
            "page_count": len(REQUIRED_REGIONS),
            "pixel_mismatch_pages": [],
        }

    monkeypatch.setattr(
        delivery_packaging,
        "verify_flattened_latex_equivalence",
        verify_flattened,
    )

    package_report = package_agent_delivery(
        source_path=source,
        pipeline_run_dir=run_dir,
        output_dir=output,
        visual_manifest_path=visual_manifest,
        template_dir=template_dir,
    )

    assert package_report["status"] == "packaged"
    assert package_report["latex_verification"]["status"] == "pass"
    assert len(verification_calls) == 1
    assert {item.name for item in output.iterdir()} == {
        "thesis.docx",
        "thesis.pdf",
        "thesis.tex",
        "model.json",
        "report.md",
        "failure_queue.json",
        "image",
    }
    assert not (output / "data").exists()
    assert (output / "image" / "figure.png").is_file()
    assert "\\include{data/body}" not in (output / "thesis.tex").read_text(
        encoding="utf-8"
    )
    assert "image/figure.png" in (output / "thesis.tex").read_text(encoding="utf-8")
    model = json.loads((output / "model.json").read_text(encoding="utf-8"))
    assert model["source"] == {
        "candidate_path": str(source.resolve()),
        "source_sha256": _sha256(source),
        "source_size": source.stat().st_size,
        "source_type": "docx",
    }
    assert json.loads((output / "failure_queue.json").read_text(encoding="utf-8")) == {
        "failures": []
    }
    delivery_report = (output / "report.md").read_text(encoding="utf-8")
    for token in (
        "https://github.com/BHOSC/BUAAthesis",
        "2020/03/05 v0.9",
        "TeX Live 2026",
        "Latexmk 4.88",
        "Times New Roman",
        "SimSun",
        "SimHei",
        "PDF font inventory",
        "G20=pass",
        "G28=pass",
    ):
        assert token in delivery_report
    assert "- Latexmk runtime: `Latexmk 4.88`" in delivery_report

    result = validate_delivery(
        output,
        profile="latex_pdf",
        gate_board_path=gate_board,
        candidate_path=source,
    )
    assert result["status"] == "pass"
    assert result["failures"] == []


def test_package_delivery_cli_writes_report_outside_public_output(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    run_dir = tmp_path / "run"
    output = tmp_path / "output"
    visual_manifest = tmp_path / "visual_review.json"
    template = tmp_path / "template"
    report = tmp_path / "package_report.json"
    captured = {}

    def fake_package(**kwargs):
        captured.update(kwargs)
        return {"status": "packaged", "output": str(output)}

    monkeypatch.setattr(package_cli, "package_agent_delivery", fake_package)

    code = package_cli.main(
        [
            str(source),
            "--run",
            str(run_dir),
            "--output",
            str(output),
            "--visual-manifest",
            str(visual_manifest),
            "--template",
            str(template),
            "--replace",
            "--out",
            str(report),
        ]
    )

    assert code == 0
    assert captured == {
        "source_path": source,
        "pipeline_run_dir": run_dir,
        "output_dir": output,
        "visual_manifest_path": visual_manifest,
        "template_dir": template,
        "editable_docx_path": None,
        "replace_existing": True,
    }
    assert json.loads(report.read_text(encoding="utf-8"))["status"] == "packaged"
