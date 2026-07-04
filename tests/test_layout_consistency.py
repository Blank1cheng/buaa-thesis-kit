import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION


def _write_reference(path: Path) -> None:
    document = Document()
    document.add_paragraph("reference")
    document.save(path)


def _write_candidate_with_non_cover_table(path: Path) -> None:
    document = Document()
    cover = document.add_table(rows=1, cols=1)
    cover.cell(0, 0).text = "cover table"
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("任务书")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "forbidden task table"
    document.save(path)


def test_validate_layout_consistency_cli_writes_report_and_fails_non_cover_tables(tmp_path, capsys):
    import scripts.validate_layout_consistency as cli

    reference = tmp_path / "reference.docx"
    candidate = tmp_path / "candidate.docx"
    report_path = tmp_path / "layout_consistency_report.json"
    _write_reference(reference)
    _write_candidate_with_non_cover_table(candidate)

    code = cli.main(
        [
            "--reference",
            str(reference),
            "--candidate",
            str(candidate),
            "--sample-mode",
            "truncated",
            "--out",
            str(report_path),
        ]
    )
    payload = json.loads(capsys.readouterr().out)
    report = json.loads(report_path.read_text(encoding="utf-8"))

    assert code == 1
    assert payload["status"] == "failed"
    assert report["checks"]["non_cover_tables"]["status"] == "failed"
    assert report["sample_mode"] == "truncated"


def test_validate_layout_consistency_does_not_flag_cover_table_labels(tmp_path):
    from scripts.validate_layout_consistency import validate_layout_consistency

    reference = tmp_path / "reference.docx"
    candidate = tmp_path / "candidate.docx"
    _write_reference(reference)
    document = Document()
    cover = document.add_table(rows=4, cols=2)
    rows = [
        ("院（系）名称", "Automation College"),
        ("专业名称", "Automation"),
        ("学生姓名", "Zhang San"),
        ("指导教师", "Li Si"),
    ]
    for row, (label, value) in zip(cover.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("1 Introduction")
    document.add_paragraph("This is a natural body paragraph with enough length.")
    document.save(candidate)

    report = validate_layout_consistency(reference, candidate, sample_mode="truncated")

    assert report["checks"]["forbidden_body_text"]["status"] == "pass"
