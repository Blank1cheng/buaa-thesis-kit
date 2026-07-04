import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

from buaa_thesis_kit.editable_template_render import _body_sections, render_editable_buaa_docx
from buaa_thesis_kit.models import AssetItem, ContentBlock, EquationItem, Metadata, ThesisModel


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "buaa_undergraduate_thesis_template.docx"
OMML_FRAGMENT = (
    '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
    "<m:oMath><m:r><m:t>x+y</m:t></m:r></m:oMath>"
    "</m:oMathPara>"
)
TINY_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x04\x00\x00\x00\xb5\x1c\x0c\x02\x00\x00\x00\x0bIDATx\xdac\xfc\xff"
    b"\x1f\x00\x03\x03\x02\x00\xef\xbf\xa7\xdb\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _body_paragraph_items(docx_path: Path) -> list[tuple[str, bool]]:
    with zipfile.ZipFile(docx_path) as docx_zip:
        root = ET.fromstring(docx_zip.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    body = root.find("w:body", ns)
    assert body is not None
    items: list[tuple[str, bool]] = []
    for paragraph in body.findall("w:p", ns):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", ns))
        has_drawing = paragraph.find(".//w:drawing", ns) is not None
        if text or has_drawing:
            items.append((text, has_drawing))
    return items


def _document_xml(docx_path: Path) -> str:
    with zipfile.ZipFile(docx_path) as docx_zip:
        return docx_zip.read("word/document.xml").decode("utf-8")


def _all_word_texts(docx_path: Path) -> list[str]:
    root = ET.fromstring(_document_xml(docx_path))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    return [node.text or "" for node in root.findall(".//w:t", ns)]


def test_render_editable_buaa_docx_reuses_template_without_page_screenshots(tmp_path):
    model = ThesisModel(
        metadata=Metadata(
            title_cn="Editable Template Thesis",
            student_name="Zhang San",
            student_id="20370001",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026年7月",
            classification="TN953",
        ),
        sections=[
            ContentBlock(
                id="body-1",
                type="section",
                title="1 Introduction",
                text="Editable body paragraph from the extracted source.",
                level=1,
            )
        ],
        references=[
            ContentBlock(
                id="ref-1",
                type="reference",
                text="[1] Wang Wu. Editable template test. 2026.",
            )
        ],
        status="needs_review",
    )
    output = tmp_path / "thesis.docx"

    render_editable_buaa_docx(TEMPLATE, model, output)

    visible_text = "\n".join(_all_word_texts(output))
    assert "Editable Template Thesis" in visible_text
    assert "20370001" in visible_text
    assert "Editable body paragraph from the extracted source." in visible_text
    assert "书脊" not in visible_text
    assert "PDF Extracted Text" not in visible_text

    with zipfile.ZipFile(output) as docx_zip:
        document_xml = docx_zip.read("word/document.xml").decode("utf-8")
        media_names = [name for name in docx_zip.namelist() if name.startswith("word/media/")]

    assert "BUAA_VERTICAL_SPINE" in document_xml
    assert 'w:textDirection w:val="tbRl"' in document_xml
    assert "Book Spine" not in document_xml
    assert "Micro-process of Oil/Water" not in document_xml
    assert not any("pdf-page" in name for name in media_names)


def test_render_editable_buaa_docx_preserves_cover_spacing_for_long_titles(tmp_path):
    model = ThesisModel(
        metadata=Metadata(
            title_cn="基于数据驱动的捷联惯性导航组件寿命预测方法研究",
            student_name="宋郭睿",
            student_id="20375284",
            college="沈元学院",
            major="自动化",
            advisor="彭朝琴",
            date="2024年6月",
            classification="TN953",
        ),
        sections=[ContentBlock(id="body-1", type="section", text="正文内容。")],
    )
    output = tmp_path / "long-title.docx"

    render_editable_buaa_docx(TEMPLATE, model, output)

    document = Document(str(output))
    before_spine = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "本科毕业设计（论文）任务书":
            break
        before_spine.append(paragraph.text)

    blank_cover_paragraphs = [text for text in before_spine if not text.strip()]
    assert "2024年6月" in before_spine
    assert 8 <= len(blank_cover_paragraphs) <= 12


def test_render_editable_buaa_docx_shrinks_long_cover_table_values(tmp_path):
    model = ThesisModel(
        metadata=Metadata(
            title_cn="基于实拍图像的光电系统性能评估关键技术研究",
            student_name="崔润昊",
            student_id="17375303",
            college="自动化科学与电气工程学院",
            major="自动化",
            advisor="唐荻音",
            date="2021年5月",
            classification="TP273",
        ),
        sections=[ContentBlock(id="body-1", type="section", text="正文内容。")],
    )
    output = tmp_path / "long-cover-table-value.docx"

    render_editable_buaa_docx(TEMPLATE, model, output)

    document = Document(str(output))
    college_cell = next(
        cell
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text == "自动化科学与电气工程学院"
    )
    sizes = [
        run.font.size.pt
        for paragraph in college_cell.paragraphs
        for run in paragraph.runs
        if run.font.size is not None
    ]
    assert college_cell.text == "自动化科学与电气工程学院"
    assert sizes
    assert max(sizes) <= 12


def test_render_editable_buaa_docx_preserves_omml_equations_as_word_math(tmp_path):
    model = ThesisModel(
        metadata=Metadata(
            title_cn="Equation Template Thesis",
            student_name="Zhang San",
            student_id="20370001",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026-06",
            classification="TP273",
        ),
        sections=[ContentBlock(id="body-1", type="section", text="Body text.")],
        equations=[
            EquationItem(
                id="eq-omml",
                kind="omml",
                text="x+y",
                omml=OMML_FRAGMENT,
                requires_review=True,
            )
        ],
    )
    output = tmp_path / "omml-equation.docx"

    render_editable_buaa_docx(TEMPLATE, model, output)

    with zipfile.ZipFile(output) as docx_zip:
        document_xml = docx_zip.read("word/document.xml").decode("utf-8")
    assert "<m:oMathPara" in document_xml
    assert "<m:t>x+y</m:t>" in document_xml
    assert "[Equation requires review]" not in document_xml


def test_render_editable_buaa_docx_generates_front_matter_and_toc_before_body(tmp_path):
    model = ThesisModel(
        metadata=Metadata(
            title_cn="Paged Front Matter Thesis",
            student_name="Zhang San",
            student_id="20370001",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026-06",
            classification="TP273",
        ),
        front_matter={
            "chinese_abstract": "Chinese abstract body.",
            "english_abstract": "English abstract body.",
        },
        sections=[
            ContentBlock(id="body", type="chapter", title="1 Introduction", text="Body text.", level=1),
        ],
    )
    output = tmp_path / "front-matter-pages.docx"

    render_editable_buaa_docx(TEMPLATE, model, output)

    document = Document(str(output))
    paragraphs = list(document.paragraphs)
    paragraph_texts = [paragraph.text for paragraph in paragraphs]
    toc_text = next(text for text in paragraph_texts if text.replace(" ", "") == "目录")
    for marker in ("本科毕业设计（论文）任务书", "本人声明", "摘    要", "Abstract", toc_text, "1 Introduction"):
        assert marker in paragraph_texts
    assert paragraph_texts.index("本科毕业设计（论文）任务书") < paragraph_texts.index("本人声明")
    assert paragraph_texts.index("本人声明") < paragraph_texts.index("摘    要")
    assert paragraph_texts.index("摘    要") < paragraph_texts.index("Abstract")
    assert paragraph_texts.index("Abstract") < paragraph_texts.index(toc_text)
    assert paragraph_texts.index(toc_text) < paragraph_texts.index("1 Introduction")
    document_xml = _document_xml(output)
    for marker in (
        "Ⅰ、毕业设计（论文）题目：",
        "Ⅱ、毕业设计（论文）使用的原始资料（数据）及设计技术要求：",
        "Ⅲ、毕业设计（论文）工作内容：",
        "Ⅳ、主要参考资料：",
    ):
        assert marker in document_xml
    assert "我声明，本论文及其研究工作是由本人在导师指导下独立完成的" in document_xml
    assert "本人郑重声明" not in document_xml
    assert "指导教师签名" not in document_xml
    assert document_xml.count("<w:sectPr") >= 6
    assert '<w:pgNumType w:fmt="upperRoman" w:start="1"' in document_xml
    assert '<w:pgNumType w:fmt="decimal" w:start="1"' in document_xml


def test_render_editable_buaa_docx_uses_buaa_abstract_labels_and_keywords(tmp_path):
    model = ThesisModel(
        metadata=Metadata(
            title_cn="Abstract Label Thesis",
            student_name="Zhang San",
            student_id="20370001",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026-06",
            classification="TP273",
        ),
        front_matter={
            "chinese_abstract": "这是中文摘要正文。",
            "keywords_cn": "光电系统，调制传递函数",
            "english_abstract": "This is the English abstract.",
            "keywords_en": "electro-optical system, MTF",
        },
        sections=[ContentBlock(id="body", type="chapter", title="1 Introduction", text="Body text.", level=1)],
    )
    output = tmp_path / "abstract-labels.docx"

    render_editable_buaa_docx(TEMPLATE, model, output)

    document = Document(str(output))
    visible_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "摘    要" in visible_text
    assert "关键词：光电系统，调制传递函数" in visible_text
    assert "Abstract" in visible_text
    assert "Key Words: electro-optical system, MTF" in visible_text
    assert "Chinese Abstract" not in visible_text
    assert "English Abstract" not in visible_text
    document_xml = _document_xml(output)
    assert '<w:pgNumType w:fmt="upperRoman" w:start="1"' in document_xml
    assert '<w:pgNumType w:fmt="decimal" w:start="1"' in document_xml


def test_render_editable_buaa_docx_separates_cn_and_en_abstract_pages(tmp_path):
    model = ThesisModel(
        metadata=Metadata(
            title_cn="基于实拍图像的光电系统性能评估关键技术研究",
            student_name="崔润昊",
            student_id="17375303",
            college="自动化科学与电气工程学院",
            major="自动化",
            advisor="唐荻音",
            date="2021年5月",
            classification="TP273",
        ),
        front_matter={
            "title_en": "Research on Key Technologies of Performance Evaluation of Electro-Optical System",
            "author_en": "CUI Run-hao",
            "tutor_en": "TANG Di-yin",
            "chinese_abstract": "这是中文摘要正文。",
            "keywords_cn": "光电系统，调制传递函数",
            "english_abstract": "This is the English abstract.",
            "keywords_en": "electro-optical system, MTF",
        },
        sections=[ContentBlock(id="body", type="chapter", title="1 绪论", text="正文。", level=1)],
    )
    output = tmp_path / "split-abstracts.docx"

    render_editable_buaa_docx(TEMPLATE, model, output)

    texts = _all_word_texts(output)
    cn_index = texts.index("摘    要")
    keywords_index = texts.index("关键词：光电系统，调制传递函数")
    english_title_index = next(index for index, text in enumerate(texts) if text.startswith("Research on"))
    author_index = texts.index("Author: CUI Run-hao")
    tutor_index = texts.index("Tutor: TANG Di-yin")
    en_index = texts.index("Abstract")
    en_page_text = "\n".join(texts[english_title_index:])
    assert "这是中文摘要正文。" in texts[cn_index:keywords_index]
    assert cn_index < keywords_index < english_title_index < author_index < tutor_index < en_index
    assert "Research on" in en_page_text
    assert "Te\nchnologies" not in en_page_text
    assert "Author: CUI Run-hao" in en_page_text
    assert "Tutor: TANG Di-yin" in en_page_text


def test_render_editable_buaa_docx_inlines_captioned_figures_near_body_caption(tmp_path):
    image_path = tmp_path / "figure.png"
    image_path.write_bytes(TINY_PNG)
    caption = "Fig. 1.1 System architecture"
    model = ThesisModel(
        metadata=Metadata(
            title_cn="Inline Figure Thesis",
            student_name="Zhang San",
            student_id="20370001",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026-06",
            classification="TP273",
        ),
        sections=[
            ContentBlock(
                id="body",
                type="chapter",
                title="1 Introduction",
                text=f"Before paragraph.\n{caption}\nAfter paragraph.",
                level=1,
            )
        ],
        figures=[
            AssetItem(
                id="fig-1",
                type="pdf-figure-image",
                path=str(image_path),
                caption=caption,
                requires_review=True,
            )
        ],
    )
    output = tmp_path / "inline-figure.docx"

    render_editable_buaa_docx(TEMPLATE, model, output)

    items = _body_paragraph_items(output)
    before_index = next(index for index, item in enumerate(items) if item[0] == "Before paragraph.")
    caption_index = next(index for index, item in enumerate(items) if item[0] == caption)
    after_index = next(index for index, item in enumerate(items) if item[0] == "After paragraph.")
    drawing_index = next(index for index, item in enumerate(items) if item[1])

    assert before_index < drawing_index < caption_index < after_index
    visible_text = "\n".join(text for text, _has_drawing in items)
    assert "[Figure inserted]" not in visible_text
    assert str(image_path) not in visible_text


def test_body_sections_filter_pdf_front_matter_residue():
    sections = [
        ContentBlock(
            id="cover-residue",
            type="chapter",
            title="1 分类号",
            text="TN953\n毕业设计(论文)\n学院名称\n沈元学院\n学生姓名\n宋郭睿",
            level=1,
        ),
        ContentBlock(
            id="spine-residue",
            type="chapter",
            title="2024 年6 月",
            text="论文封面书脊\n四号黑体字",
            level=1,
        ),
        ContentBlock(
            id="task-residue",
            type="chapter",
            title="Ⅱ、毕业设计（论文）使用的原始资料（数据）及设计技术要求：",
            text="实验室具有计算机硬件以及python软件开发平台。",
            level=1,
        ),
        ContentBlock(id="body", type="section", title="1.1 课题来源", text="正文内容。", level=2),
    ]

    rendered = _body_sections(sections)

    assert [section.title for section in rendered] == ["1.1 课题来源"]


def test_render_editable_buaa_docx_inlines_figures_when_pdf_caption_is_split(tmp_path):
    image_path = tmp_path / "figure.png"
    image_path.write_bytes(TINY_PNG)
    model = ThesisModel(
        metadata=Metadata(
            title_cn="Split Caption Figure Thesis",
            student_name="Zhang San",
            student_id="20370001",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026-06",
            classification="TP273",
        ),
        sections=[
            ContentBlock(
                id="body",
                type="chapter",
                title="1 Introduction",
                text="Before paragraph.\nFig. 1.1\nSystem architecture\nAfter paragraph.",
                level=1,
            )
        ],
        figures=[
            AssetItem(
                id="fig-1",
                type="pdf-figure-image",
                path=str(image_path),
                caption="Fig. 1.1 System architecture",
                requires_review=True,
            )
        ],
    )
    output = tmp_path / "split-caption-figure.docx"

    render_editable_buaa_docx(TEMPLATE, model, output)

    items = _body_paragraph_items(output)
    marker_index = next(index for index, item in enumerate(items) if item[0] == "Fig. 1.1")
    drawing_index = next(index for index, item in enumerate(items) if item[1])
    assert drawing_index + 1 == marker_index
