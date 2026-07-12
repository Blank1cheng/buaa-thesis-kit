import zipfile
from pathlib import Path

import yaml
from docx import Document

from buaa_thesis_kit.frontmatter_render.capture_reference import capture_reference_frontmatter
from buaa_thesis_kit.frontmatter_render.replace_placeholders import replace_docx_placeholders


def _xml_text(docx_path: Path) -> str:
    with zipfile.ZipFile(docx_path) as package:
        return package.read("word/document.xml").decode("utf-8")


def test_replace_docx_placeholders_handles_split_runs_and_headers(tmp_path):
    source = tmp_path / "template.docx"
    output = tmp_path / "output.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("{{TITLE_")
    paragraph.add_run("CN}}")
    document.sections[0].header.add_paragraph("{{STUDENT_NAME}}")
    document.save(source)

    replace_docx_placeholders(
        source,
        output,
        {
            "TITLE_CN": "基于实拍图像的光电系统性能评估关键技术研究",
            "STUDENT_NAME": "崔润昊",
        },
    )

    result = Document(str(output))
    assert "基于实拍图像的光电系统性能评估关键技术研究" in "\n".join(
        paragraph.text for paragraph in result.paragraphs
    )
    assert "崔润昊" in "\n".join(paragraph.text for paragraph in result.sections[0].header.paragraphs)
    assert "{{TITLE_" not in _xml_text(output)


def test_capture_reference_frontmatter_reports_missing_reference(tmp_path):
    result = capture_reference_frontmatter(tmp_path / "missing.docx", tmp_path / "captured")

    assert result.status == "failed"
    assert any("reference_docx_missing" in note for note in result.notes)


def test_frontmatter_layout_spec_and_style_registry_exist():
    root = Path(__file__).resolve().parents[1]
    spec_path = root / "templates" / "frontmatter" / "layout_spec.yaml"

    assert spec_path.is_file()
    spec = yaml.safe_load(spec_path.read_text(encoding="utf-8"))
    for page_name in ("cover", "spine", "taskbook", "declaration", "abstract_cn", "abstract_en", "toc"):
        assert page_name in spec

    assert spec["cover"]["metadata"]["unit_code"]["value"]["font"] == "Times New Roman"
    assert spec["cover"]["title"]["line1"]["font"] == "SimHei"
    assert spec["spine"]["textbox"]["text_direction"] == "tbRl"
    assert spec["abstract_cn"]["abstract_heading"]["text"] == "摘    要"

    from buaa_thesis_kit.styles.buaa_styles import BUAA_STYLES

    for style_name in (
        "CoverTitle",
        "CoverThesisTitle",
        "CoverFieldLabel",
        "CoverFieldValue",
        "TaskBookTitle",
        "TaskBookHeading",
        "TaskBookBody",
        "TaskBookReference",
        "DeclarationTitle",
        "DeclarationBody",
        "DeclarationSignature",
        "AbstractTitleCN",
        "AbstractBodyCN",
        "AbstractTitleEN",
        "AbstractBodyEN",
        "TOCTitle",
        "TOC1",
        "TOC2",
        "TOC3",
        "BodyHeading1",
        "BodyHeading2",
        "BodyNormal",
    ):
        assert style_name in BUAA_STYLES


def test_validate_frontmatter_render_cli_accepts_sample_mode_pages(tmp_path, monkeypatch, capsys):
    import scripts.validate_frontmatter_render as cli

    captured = {}

    def fake_validate_render_frontmatter(reference, candidate, output_dir, config=None, pages=None):
        captured["reference"] = reference
        captured["candidate"] = candidate
        captured["output_dir"] = output_dir
        captured["pages"] = pages
        captured["sample_mode"] = config.sample_mode
        return {"status": "needs_review", "blocking_items": [], "notes": ["fake"]}

    monkeypatch.setattr(cli, "validate_render_frontmatter", fake_validate_render_frontmatter)

    code = cli.main(
        [
            "--reference",
            str(tmp_path / "reference.docx"),
            "--candidate",
            str(tmp_path / "candidate.docx"),
            "--pages",
            "cover,spine,toc",
            "--sample-mode",
            "truncated",
            "--out",
            str(tmp_path / "frontmatter_diff"),
        ]
    )
    payload = capsys.readouterr().out

    assert code == 0
    assert '"status": "needs_review"' in payload
    assert captured["pages"] == ["cover", "spine", "toc"]
    assert captured["sample_mode"] == "truncated"


def test_frontmatter_validators_accept_spaced_toc_title():
    from buaa_thesis_kit.frontmatter_render.validate_render import _has_toc_render
    from scripts.validate_front_matter import _require_order

    assert _has_toc_render("目    录\n1 绪论 ................................................................ 1")
    blocking_items = []
    _require_order(blocking_items, ["本人声明", "目    录", "1 绪论"], ("本人声明", "目录"))
    assert blocking_items == []
