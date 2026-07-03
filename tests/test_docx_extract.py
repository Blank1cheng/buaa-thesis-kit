import base64
import json
import zipfile
from pathlib import Path

from docx import Document
from jsonschema import validate

from buaa_thesis_kit import docx_extract as docx_extract_module
from buaa_thesis_kit.docx_extract import extract_thesis_model


TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


def _save_docx(path: Path, paragraphs: list[str], metadata_table: bool = False) -> None:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if metadata_table:
        table = doc.add_table(rows=4, cols=4)
        rows = [
            ("学生姓名", "张三", "学号", "20370001"),
            ("学院", "自动化科学与电气工程学院", "专业", "自动化"),
            ("指导教师", "李四", "日期", "2026年6月"),
            ("分类号", "TP391", "单位代码", "10006"),
        ]
        for row, values in zip(table.rows, rows):
            for cell, value in zip(row.cells, values):
                cell.text = value
    doc.save(path)


def _patch_docx_zip(
    docx_path: Path,
    *,
    add_omml: bool = False,
    add_embedding: bool = False,
    media_entries: dict[str, bytes] | None = None,
) -> None:
    patched_path = docx_path.with_suffix(".patched.docx")
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(patched_path, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if add_omml and info.filename == "word/document.xml":
                omml = (
                    '<w:p><m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                    "<m:oMath><m:r><m:t>x+y</m:t></m:r></m:oMath>"
                    "</m:oMathPara></w:p>"
                ).encode("utf-8")
                data = data.replace(b"</w:body>", omml + b"</w:body>")
            target.writestr(info, data)
        if add_embedding:
            target.writestr("word/embeddings/oleObject1.bin", b"fake ole payload")
        for name, payload in (media_entries or {}).items():
            target.writestr(name, payload)
    patched_path.replace(docx_path)


def _patch_docx_with_ole_objects(docx_path: Path) -> None:
    patched_path = docx_path.with_suffix(".ole.docx")
    equation_object = (
        '<w:p><w:r><w:object>'
        '<v:shape xmlns:v="urn:schemas-microsoft-com:vml">'
        '<v:imagedata xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'r:id="rIdEquationImage"/>'
        "</v:shape>"
        '<o:OLEObject xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'Type="Embed" ProgID="Equation.DSMT4" r:id="rIdEquation"/>'
        "</w:object></w:r></w:p>"
    ).encode("utf-8")
    visio_object = (
        '<w:p><w:r><w:object>'
        '<o:OLEObject xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'Type="Embed" ProgID="Visio.Drawing.15" r:id="rIdVisio"/>'
        "</w:object></w:r></w:p>"
    ).encode("utf-8")
    relationships = (
        '<Relationship Id="rIdEquation" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" '
        'Target="embeddings/equation.bin"/>'
        '<Relationship Id="rIdEquationImage" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
        'Target="media/equation-preview.png"/>'
        '<Relationship Id="rIdVisio" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" '
        'Target="embeddings/visio.vsdx"/>'
    ).encode("utf-8")
    content_types = (
        '<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.oleObject"/>'
        '<Default Extension="png" ContentType="image/png"/>'
        '<Default Extension="vsdx" ContentType="application/vnd.ms-visio.drawing.main+xml"/>'
    ).encode("utf-8")
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(patched_path, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                data = data.replace(b"</w:body>", equation_object + visio_object + b"</w:body>")
            elif info.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", relationships + b"</Relationships>")
            elif info.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", content_types + b"</Types>")
            target.writestr(info, data)
        target.writestr("word/embeddings/equation.bin", b"equation ole payload")
        target.writestr("word/embeddings/visio.vsdx", b"visio payload")
        target.writestr("word/media/equation-preview.png", TINY_PNG)
    patched_path.replace(docx_path)


def test_extracts_metadata_from_paragraphs_and_tables_with_evidence(tmp_path):
    source = tmp_path / "metadata.docx"
    work_dir = tmp_path / "work"
    _save_docx(
        source,
        [
            "中文题目：基于数据驱动的系统研究",
            "英文题目：Research on Data-driven Systems",
            "摘要",
            "本文研究数据驱动系统。",
            "第一章 绪论",
            "正文内容。",
            "参考文献",
            "[1] 王五. 数据驱动系统研究[J]. 自动化学报, 2026.",
        ],
        metadata_table=True,
    )
    original_bytes = source.read_bytes()

    model = extract_thesis_model(source, work_dir)

    assert (work_dir / "source.docx").read_bytes() == original_bytes
    assert source.read_bytes() == original_bytes
    assert model.metadata.title_cn == "基于数据驱动的系统研究"
    assert model.metadata.title_en == "Research on Data-driven Systems"
    assert model.metadata.student_name == "张三"
    assert model.metadata.student_id == "20370001"
    assert model.metadata.college == "自动化科学与电气工程学院"
    assert model.metadata.major == "自动化"
    assert model.metadata.advisor == "李四"
    assert model.metadata.date == "2026年6月"
    assert model.metadata.classification == "TP391"
    assert model.metadata.unit_code == "10006"
    for field in ("title_cn", "student_name", "student_id", "unit_code"):
        evidence = model.metadata.evidence[field]
        assert evidence.file == "source.docx"
        assert evidence.method
        assert evidence.confidence > 0
        assert evidence.requires_review is False


def test_metadata_prefers_longer_overlapping_labels(tmp_path):
    source = tmp_path / "overlapping-labels.docx"
    work_dir = tmp_path / "work"
    _save_docx(
        source,
        [
            "中文题目：重叠标签测试",
            "学生姓名：张三",
            "学号：20370001",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师姓名：李四",
            "日期：2026年6月",
            "1 绪论",
            "正文内容。",
            "参考文献",
            "[1] 王五. 重叠标签测试[J]. 2026.",
        ],
    )

    model = extract_thesis_model(source, work_dir)

    assert model.metadata.advisor == "李四"


def test_splits_front_matter_body_sections_and_references_without_toc_entries(tmp_path):
    source = tmp_path / "sections.docx"
    work_dir = tmp_path / "work"
    _save_docx(
        source,
        [
            "中文题目：无人机控制系统研究",
            "学生姓名：张三",
            "学号：20370001",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师：李四",
            "日期：2026年6月",
            "目录",
            "1 绪论..........................1",
            "1.1 研究背景.....................2",
            "摘要",
            "这是中文摘要。",
            "关键词：无人机；控制",
            "ABSTRACT",
            "This is the English abstract.",
            "Keywords: UAV; control",
            "1 绪论",
            "绪论正文。",
            "1.1 研究背景",
            "背景正文。",
            "致谢",
            "感谢老师和同学。",
            "参考文献",
            "[1] 王五. 无人机控制研究[J]. 航空学报, 2026.",
            "[2] Smith J. Control Systems[M]. 2025.",
        ],
    )

    model = extract_thesis_model(source, work_dir)

    assert model.front_matter["chinese_abstract"] == "这是中文摘要。"
    assert model.front_matter["english_abstract"] == "This is the English abstract."
    body_titles = [section.title for section in model.sections if section.title.startswith("1")]
    assert body_titles == ["1 绪论", "1.1 研究背景"]
    assert all("..." not in section.title for section in model.sections)
    assert [ref.text for ref in model.references] == [
        "[1] 王五. 无人机控制研究[J]. 航空学报, 2026.",
        "[2] Smith J. Control Systems[M]. 2025.",
    ]


def test_skips_normalized_toc_entries_with_plain_page_numbers(tmp_path):
    source = tmp_path / "normalized-toc.docx"
    work_dir = tmp_path / "work"
    _save_docx(
        source,
        [
            "中文题目：目录测试",
            "学生姓名：张三",
            "学号：20370001",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师：李四",
            "日期：2026年6月",
            "目录",
            "1 绪论 1",
            "1.1 Background 2",
            "第一章 绪论 3",
            "摘要",
            "这是中文摘要。",
            "ABSTRACT",
            "This is the English abstract.",
            "1 绪论",
            "绪论正文。",
            "1.1 Background",
            "Background body.",
            "参考文献",
            "[1] 王五. 目录测试[J]. 2026.",
        ],
    )

    model = extract_thesis_model(source, work_dir)

    assert [section.title for section in model.sections] == ["1 绪论", "1.1 Background"]


def test_body_table_cells_are_not_duplicated_into_section_text(tmp_path):
    source = tmp_path / "body-table.docx"
    work_dir = tmp_path / "work"
    doc = Document()
    for text in [
        "中文题目：表格测试",
        "学生姓名：张三",
        "学号：20370001",
        "学院：自动化科学与电气工程学院",
        "专业：自动化",
        "指导教师：李四",
        "日期：2026年6月",
        "摘要",
        "这是中文摘要。",
        "ABSTRACT",
        "This is the English abstract.",
        "1 绪论",
        "表前正文。",
    ]:
        doc.add_paragraph(text)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "不要进入章节"
    table.cell(1, 1).text = "42"
    doc.add_paragraph("表后正文。")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 王五. 表格测试[J]. 2026.")
    doc.save(source)

    model = extract_thesis_model(source, work_dir)

    assert len(model.tables) == 1
    assert "不要进入章节" in model.tables[0].text
    assert model.sections[0].text == "表前正文。\n表后正文。"


def test_extracts_images_to_work_dir_and_marks_review(tmp_path):
    source = tmp_path / "image.docx"
    image_path = tmp_path / "tiny.png"
    image_path.write_bytes(TINY_PNG)
    doc = Document()
    doc.add_paragraph("中文题目：图像测试")
    doc.add_paragraph("学生姓名：张三")
    doc.add_paragraph("学号：20370001")
    doc.add_paragraph("学院：自动化科学与电气工程学院")
    doc.add_paragraph("专业：自动化")
    doc.add_paragraph("指导教师：李四")
    doc.add_paragraph("日期：2026年6月")
    doc.add_picture(str(image_path))
    doc.add_paragraph("1 绪论")
    doc.add_paragraph("正文。")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 王五. 图像测试[J]. 2026.")
    doc.save(source)

    model = extract_thesis_model(source, tmp_path / "work")

    assert len(model.figures) == 1
    copied_image = Path(model.figures[0].path)
    assert copied_image.exists()
    assert copied_image.parent.name == "image"
    assert copied_image.read_bytes() == TINY_PNG
    assert model.figures[0].requires_review is True
    assert model.status == "needs_review"


def test_detects_omml_and_embedded_equations(tmp_path):
    source = tmp_path / "equations.docx"
    _save_docx(
        source,
        [
            "中文题目：公式测试",
            "学生姓名：张三",
            "学号：20370001",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师：李四",
            "日期：2026年6月",
            "1 绪论",
            "正文。",
            "参考文献",
            "[1] 王五. 公式测试[J]. 2026.",
        ],
    )
    _patch_docx_zip(source, add_omml=True)
    _patch_docx_with_ole_objects(source)

    model = extract_thesis_model(source, tmp_path / "work")

    by_kind = {equation.kind: equation for equation in model.equations}
    assert by_kind["omml"].requires_review is True
    assert by_kind["omml"].omml.startswith("<m:oMathPara")
    assert "<m:t>x+y</m:t>" in by_kind["omml"].omml
    assert by_kind["embedded-object"].requires_review is True
    assert Path(by_kind["embedded-object"].preview_path).read_bytes() == TINY_PNG
    assert "OMML equations require TeX review" in model.extraction_warnings
    assert model.status == "needs_review"


def test_embedded_visio_objects_are_not_counted_as_equations(tmp_path):
    source = tmp_path / "ole-classification.docx"
    _save_docx(
        source,
        [
            "中文题目：嵌入对象分类测试",
            "学生姓名：张三",
            "学号：20370001",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师：李四",
            "日期：2026年6月",
            "1 绪论",
            "正文。",
            "参考文献",
            "[1] 王五. 嵌入对象测试[J]. 2026.",
        ],
    )
    _patch_docx_with_ole_objects(source)

    model = extract_thesis_model(source, tmp_path / "work")

    assert [(equation.kind, equation.text) for equation in model.equations] == [
        ("embedded-object", "equation.bin")
    ]
    assert Path(model.equations[0].preview_path).read_bytes() == TINY_PNG


def test_duplicate_media_basenames_are_extracted_to_unique_paths(tmp_path):
    source = tmp_path / "duplicate-media.docx"
    _save_docx(
        source,
        [
            "中文题目：媒体测试",
            "学生姓名：张三",
            "学号：20370001",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师：李四",
            "日期：2026年6月",
            "1 绪论",
            "正文。",
            "参考文献",
            "[1] 王五. 媒体测试[J]. 2026.",
        ],
    )
    other_payload = b"not really a png but still a media resource"
    _patch_docx_zip(
        source,
        media_entries={
            "word/media/a/image.png": TINY_PNG,
            "word/media/b/image.png": other_payload,
        },
    )

    model = extract_thesis_model(source, tmp_path / "work")

    paths = [Path(figure.path) for figure in model.figures]
    assert len(paths) == 2
    assert len({path.name for path in paths}) == 2
    assert paths[0].read_bytes() == TINY_PNG
    assert paths[1].read_bytes() == other_payload


def test_oversized_media_is_skipped_with_warning(tmp_path, monkeypatch):
    source = tmp_path / "oversized-media.docx"
    _save_docx(
        source,
        [
            "中文题目：大媒体测试",
            "学生姓名：张三",
            "学号：20370001",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师：李四",
            "日期：2026年6月",
            "1 绪论",
            "正文。",
            "参考文献",
            "[1] 王五. 大媒体测试[J]. 2026.",
        ],
    )
    monkeypatch.setattr(docx_extract_module, "MAX_MEDIA_BYTES", 4)
    _patch_docx_zip(source, media_entries={"word/media/large.png": b"12345"})

    model = extract_thesis_model(source, tmp_path / "work")

    assert model.figures == []
    assert "skipped oversized media: large.png" in model.extraction_warnings
    assert model.status == "needs_review"


def test_total_media_size_limit_skips_later_media_with_warning(tmp_path, monkeypatch):
    source = tmp_path / "total-media-limit.docx"
    _save_docx(
        source,
        [
            "中文题目：总媒体测试",
            "学生姓名：张三",
            "学号：20370001",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师：李四",
            "日期：2026年6月",
            "1 绪论",
            "正文。",
            "参考文献",
            "[1] 王五. 总媒体测试[J]. 2026.",
        ],
    )
    monkeypatch.setattr(docx_extract_module, "MAX_MEDIA_BYTES", 10)
    monkeypatch.setattr(docx_extract_module, "MAX_TOTAL_MEDIA_BYTES", 6)
    _patch_docx_zip(
        source,
        media_entries={
            "word/media/first.png": b"1234",
            "word/media/second.png": b"5678",
        },
    )

    model = extract_thesis_model(source, tmp_path / "work")

    assert [Path(figure.path).name for figure in model.figures] == ["first.png"]
    assert "skipped media after total size limit: second.png" in model.extraction_warnings
    assert model.status == "needs_review"


def test_conflicting_metadata_values_emit_warning_and_needs_review(tmp_path):
    source = tmp_path / "metadata-conflict.docx"
    _save_docx(
        source,
        [
            "中文题目：冲突测试",
            "学生姓名：张三",
            "学生姓名：李四",
            "学号：20370001",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师：王五",
            "日期：2026年6月",
            "1 绪论",
            "正文。",
            "参考文献",
            "[1] 赵六. 冲突测试[J]. 2026.",
        ],
    )

    model = extract_thesis_model(source, tmp_path / "work")

    assert model.metadata.student_name == "张三"
    assert "conflicting metadata field: student_name" in model.extraction_warnings
    assert model.status == "needs_review"


def test_conflicting_student_ids_emit_warning_and_needs_review(tmp_path):
    source = tmp_path / "student-id-conflict.docx"
    _save_docx(
        source,
        [
            "中文题目：学号冲突测试",
            "学生姓名：张三",
            "学号：20370001",
            "学号：20370002",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师：王五",
            "日期：2026年6月",
            "1 绪论",
            "正文。",
            "参考文献",
            "[1] 赵六. 学号冲突测试[J]. 2026.",
        ],
    )

    model = extract_thesis_model(source, tmp_path / "work")

    assert model.metadata.student_id == "20370001"
    assert "conflicting metadata field: student_id" in model.extraction_warnings
    assert model.status == "needs_review"


def test_extracts_spaced_cover_student_id_and_unlabeled_cover_date(tmp_path):
    source = tmp_path / "buaa-cover.docx"
    work_dir = tmp_path / "work"
    _save_docx(
        source,
        [
            "单位代码       10006",
            "学    号      17375303",
            "分类号    TP273",
            "毕业设计(论文)",
            "基于实拍图像的光电系统性能评估",
            "关键技术研究",
            "院（系）名称",
            "自动化科学与电气工程学院",
            "专业名称：自动化",
            "学生姓名：崔润昊",
            "指导教师：唐荻音",
            "2021年5月",
            "摘    要",
            "这是摘要。",
            "ABSTRACT",
            "This is the abstract.",
            "1 绪论",
            "正文。",
            "参考文献",
            "[1] 王五. 测试[J]. 2021.",
        ],
    )

    model = extract_thesis_model(source, work_dir)

    assert model.metadata.title_cn == "基于实拍图像的光电系统性能评估关键技术研究"
    assert model.metadata.student_id == "17375303"
    assert model.metadata.date == "2021年5月"
    assert model.metadata.classification == "TP273"


def test_cover_table_metadata_beats_later_task_book_paragraphs(tmp_path):
    source = tmp_path / "cover-table-priority.docx"
    work_dir = tmp_path / "work"
    doc = Document()
    for text in [
        "单位代码       10006",
        "学    号      17375303",
        "分类号    TP273",
        "毕业设计(论文)",
        "基于实拍图像的光电系统性能评估",
        "关键技术研究",
        "2021年5月",
    ]:
        doc.add_paragraph(text)
    table = doc.add_table(rows=4, cols=2)
    rows = [
        ("院（系）名称", "自动化科学与电气工程学院"),
        ("专业名称", "自动化"),
        ("学生姓名", "崔润昊"),
        ("指导教师", "唐荻音"),
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    for text in [
        "北京航空航天大学",
        "本科毕业设计（论文）任务书",
        "申请人所在院系：自动化 专业类 170325 班",
        "申请人专业：类 170325 班",
        "摘    要",
        "这是摘要。",
        "ABSTRACT",
        "This is the abstract.",
        "1 绪论",
        "正文。",
        "参考文献",
        "[1] 王五. 测试[J]. 2021.",
    ]:
        doc.add_paragraph(text)
    doc.save(source)

    model = extract_thesis_model(source, work_dir)

    assert model.metadata.college == "自动化科学与电气工程学院"
    assert model.metadata.major == "自动化"
    assert model.metadata.student_name == "崔润昊"
    assert model.metadata.advisor == "唐荻音"


def test_extracted_model_payload_validates_against_schema(tmp_path):
    source = tmp_path / "schema.docx"
    _save_docx(
        source,
        [
            "中文题目：模式验证测试",
            "学生姓名：张三",
            "学号：20370001",
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "指导教师：李四",
            "日期：2026年6月",
            "摘要",
            "这是中文摘要。",
            "ABSTRACT",
            "This is the English abstract.",
            "1 绪论",
            "正文。",
            "参考文献",
            "[1] 王五. 模式验证测试[J]. 2026.",
        ],
    )
    schema_path = Path(__file__).resolve().parents[1] / "schemas" / "thesis-model.schema.json"
    schema = json.loads(schema_path.read_text(encoding="utf-8"))

    model = extract_thesis_model(source, tmp_path / "work")

    validate(instance=model.to_dict(), schema=schema)
