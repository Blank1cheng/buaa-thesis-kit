from __future__ import annotations

import html
import zipfile
from pathlib import Path

from docx import Document

from buaa_thesis_kit.docx_extract import extract_thesis_model


def _save_minimal_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("中文题目：元数据抽取测试")
    doc.add_paragraph("学号：17375303")
    doc.add_paragraph("分类号：TP273")
    doc.add_paragraph("2021 年 5 月")
    doc.add_paragraph("摘要")
    doc.add_paragraph("这是中文摘要。")
    doc.add_paragraph("ABSTRACT")
    doc.add_paragraph("This is the English abstract.")
    doc.add_paragraph("1 绪论")
    doc.add_paragraph("正文内容。")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 王五. 测试[J]. 2021.")
    doc.save(path)


def _patch_docx_with_textbox(docx_path: Path, lines: list[str]) -> None:
    patched_path = docx_path.with_suffix(".textbox.docx")
    paragraph_xml = "".join(
        f"<w:p><w:r><w:t>{html.escape(line)}</w:t></w:r></w:p>" for line in lines
    )
    textbox_xml = (
        '<w:p><w:r><w:pict>'
        '<v:shape xmlns:v="urn:schemas-microsoft-com:vml" id="TextBox1" type="#_x0000_t202">'
        '<v:textbox><w:txbxContent>'
        f"{paragraph_xml}"
        "</w:txbxContent></v:textbox>"
        "</v:shape>"
        "</w:pict></w:r></w:p>"
    ).encode("utf-8")
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(patched_path, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                data = data.replace(b"</w:body>", textbox_xml + b"</w:body>")
            target.writestr(info, data)
    patched_path.replace(docx_path)


def test_docx_metadata_reads_textboxes_from_raw_word_xml_with_evidence(tmp_path: Path):
    source = tmp_path / "textbox-metadata.docx"
    work_dir = tmp_path / "work"
    _save_minimal_docx(source)
    _patch_docx_with_textbox(
        source,
        [
            "学院：自动化科学与电气工程学院",
            "专业：自动化",
            "学生姓名：文本框生",
            "指导教师：框内导师",
        ],
    )

    model = extract_thesis_model(source, work_dir)

    assert model.metadata.college == "自动化科学与电气工程学院"
    assert model.metadata.major == "自动化"
    assert model.metadata.student_name == "文本框生"
    assert model.metadata.advisor == "框内导师"
    for field in ("college", "major", "student_name", "advisor"):
        evidence = model.metadata.evidence[field]
        assert evidence.file == "source.docx"
        assert "raw-docx-xml" in evidence.method
        assert "textbox" in evidence.method
        assert getattr(evidence, "source_type") == "docx"
        assert getattr(evidence, "evidence_text")
    assert not list(work_dir.rglob("*.pdf"))


def test_docx_table_metadata_evidence_is_raw_ooxml_table_source(tmp_path: Path):
    source = tmp_path / "table-metadata.docx"
    work_dir = tmp_path / "work"
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    rows = [
        ("学院", "自动化科学与电气工程学院"),
        ("专业", "自动化"),
        ("学生姓名", "表格学生"),
        ("指导教师", "表格导师"),
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.add_paragraph("中文题目：表格元数据测试")
    doc.add_paragraph("2021 年 5 月")
    doc.add_paragraph("1 绪论")
    doc.add_paragraph("正文内容。")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 王五. 测试[J]. 2021.")
    doc.save(source)

    model = extract_thesis_model(source, work_dir)

    assert model.metadata.student_name == "表格学生"
    assert model.metadata.advisor == "表格导师"
    for field in ("college", "major", "student_name", "advisor"):
        evidence = model.metadata.evidence[field]
        assert "raw-docx-xml" in evidence.method
        assert "table" in evidence.method
        assert getattr(evidence, "source_type") == "docx"
        assert getattr(evidence, "evidence_text")
