import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Cm

from buaa_thesis_kit.editable_template_render import render_editable_buaa_docx
from buaa_thesis_kit.models import ContentBlock, Metadata, ThesisModel
from buaa_thesis_kit.validate.no_table_validator import validate_no_tables_after_cover


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "buaa_undergraduate_thesis_template.docx"


def _document_xml(docx_path: Path) -> str:
    with zipfile.ZipFile(docx_path) as package:
        return package.read("word/document.xml").decode("utf-8")


def _write_cover_table_only_docx(path: Path) -> None:
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "cover table allowed"
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("task book as paragraphs")
    document.save(path)


def _write_non_cover_table_docx(path: Path) -> None:
    document = Document()
    cover = document.add_table(rows=1, cols=1)
    cover.cell(0, 0).text = "cover table allowed"
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("task book")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "table after cover forbidden"
    document.save(path)


def _write_header_table_docx(path: Path) -> None:
    document = Document()
    cover = document.add_table(rows=1, cols=1)
    cover.cell(0, 0).text = "cover table allowed"
    section = document.add_section(WD_SECTION.NEW_PAGE)
    header_table = section.header.add_table(rows=1, cols=1, width=Cm(10))
    header_table.cell(0, 0).text = "header table forbidden"
    document.add_paragraph("abstract")
    document.save(path)


def test_no_table_validator_allows_cover_table_before_first_section_break(tmp_path):
    docx_path = tmp_path / "cover-table-only.docx"
    _write_cover_table_only_docx(docx_path)

    result = validate_no_tables_after_cover(docx_path)

    assert result.status == "pass"
    assert result.blocking_items == []


def test_no_table_validator_blocks_tables_after_cover_section(tmp_path):
    docx_path = tmp_path / "non-cover-table.docx"
    _write_non_cover_table_docx(docx_path)

    result = validate_no_tables_after_cover(docx_path)

    assert result.status == "failed"
    assert any("document.xml" in item for item in result.blocking_items)
    assert result.table_locations


def test_no_table_validator_blocks_header_footer_tables(tmp_path):
    docx_path = tmp_path / "header-table.docx"
    _write_header_table_docx(docx_path)

    result = validate_no_tables_after_cover(docx_path)

    assert result.status == "failed"
    assert any("header" in item for item in result.blocking_items)


def test_render_editable_buaa_docx_has_no_tables_after_cover(tmp_path):
    output = tmp_path / "thesis.docx"
    model = ThesisModel(
        metadata=Metadata(
            title_cn="Editable Flow Thesis",
            student_name="Zhang San",
            student_id="20370001",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026-06",
            classification="TP273",
        ),
        front_matter={
            "chinese_abstract": "第一行中文摘要\n第二行中文摘要",
            "english_abstract": "The first English abstract line.\nThe second English abstract line.",
        },
        sections=[ContentBlock(id="body", type="chapter", title="1 Introduction", text="Body text.", level=1)],
    )

    render_editable_buaa_docx(TEMPLATE, model, output)

    result = validate_no_tables_after_cover(output)
    assert result.status == "pass"
    assert '<w:tbl' in _document_xml(output)


def test_render_editable_buaa_docx_filters_cover_metadata_tables_from_body(tmp_path):
    output = tmp_path / "cover-table-filtered.docx"
    model = ThesisModel(
        metadata=Metadata(
            title_cn="Cover Table Filter Thesis",
            student_name="Zhang San",
            student_id="20370001",
            college="Automation College",
            major="Automation",
            advisor="Li Si",
            date="2026-06",
            classification="TP273",
        ),
        sections=[ContentBlock(id="body", type="chapter", title="1 Introduction", text="Body text.", level=1)],
        tables=[
            ContentBlock(
                id="cover-table",
                type="table",
                text="院（系）名称\tAutomation College\n专业名称\tAutomation\n学生姓名\tZhang San\n指导教师\tLi Si",
            )
        ],
    )

    render_editable_buaa_docx(TEMPLATE, model, output)

    result = validate_no_tables_after_cover(output)
    assert result.status == "pass"
    assert "院（系）名称\tAutomation College" not in _document_xml(output)
