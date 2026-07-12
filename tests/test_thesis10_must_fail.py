import json
from pathlib import Path


def test_render_smoke_does_not_mutate_candidate(tmp_path, monkeypatch):
    from docx import Document
    from pypdf import PdfWriter

    import scripts.validate_render_smoke as smoke
    from buaa_thesis_kit.harness.artifact_identity import sha256_file

    candidate = tmp_path / "candidate.docx"
    document = Document()
    document.add_paragraph("单位代码 10006")
    document.save(candidate)
    before = sha256_file(candidate)

    def fake_export(source: Path, target: Path) -> tuple[bool, str]:
        assert source != candidate
        copied = Document(source)
        copied.add_paragraph("export side effect")
        copied.save(source)
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with target.open("wb") as handle:
            writer.write(handle)
        return True, "fake export"

    monkeypatch.setattr(smoke, "export_pdf_from_docx", fake_export)

    smoke.validate_render_smoke(candidate=candidate, out_dir=tmp_path / "smoke", sample_mode="truncated")

    assert sha256_file(candidate) == before


def test_thesis10_bad_fixture_is_captured_by_harness(tmp_path):
    from scripts.validate_render_smoke import validate_render_smoke

    candidate = Path("tests/fixtures/bad_outputs/thesis10.docx")
    if not candidate.exists():
        raise AssertionError("thesis10.docx fixture must be copied from the rejected artifact before this test runs")

    report = validate_render_smoke(
        candidate=candidate,
        out_dir=tmp_path / "test_thesis10_render_smoke",
        sample_mode="truncated",
    )

    reason_ids = {item["id"] for item in report.get("failures", [])}
    assert report["status"] == "failed"
    assert {
        "cover_classification_split",
        "cover_metadata_misaligned",
        "cover_title_orphan_or_bad_wrap",
        "declaration_contamination",
        "frontmatter_blank_or_placeholder_fields",
        "taskbook_blank_or_placeholder_fields",
        "abstract_region_issue",
        "abstract_render_smoke_failed",
        "toc_template_sample_leak",
        "body_template_page_residue",
    }.issubset(reason_ids), json.dumps(report, ensure_ascii=False, indent=2)
