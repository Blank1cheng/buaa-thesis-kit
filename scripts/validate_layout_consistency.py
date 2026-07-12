from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from buaa_thesis_kit.validate.no_table_validator import validate_no_tables_after_cover  # noqa: E402


def validate_layout_consistency(
    reference: Path,
    candidate: Path,
    *,
    sample_mode: str = "full",
    out: Path | None = None,
) -> dict[str, object]:
    report: dict[str, object] = {
        "status": "pass",
        "sample_mode": sample_mode,
        "reference": str(reference),
        "candidate": str(candidate),
        "checks": {},
        "blocking_items": [],
        "notes": [],
    }
    checks: dict[str, dict[str, object]] = report["checks"]  # type: ignore[assignment]
    blocking_items: list[str] = report["blocking_items"]  # type: ignore[assignment]
    notes: list[str] = report["notes"]  # type: ignore[assignment]

    if not Path(reference).exists():
        _set_check(checks, "reference_exists", "needs_review", f"reference missing: {reference}")
        notes.append("reference document missing; structural candidate checks still ran.")
    else:
        _set_check(checks, "reference_exists", "pass", "")

    if not Path(candidate).exists():
        _set_check(checks, "candidate_exists", "failed", f"candidate missing: {candidate}")
        blocking_items.append(f"candidate_missing: {candidate}")
        report["status"] = "failed"
        _write_report_if_requested(report, out)
        return report
    _set_check(checks, "candidate_exists", "pass", "")

    document_xml = _read_part(candidate, "word/document.xml")
    all_text = _visible_text(candidate)
    no_table = validate_no_tables_after_cover(candidate)
    _set_check(
        checks,
        "non_cover_tables",
        no_table.status,
        "; ".join(no_table.blocking_items),
        {"locations": no_table.table_locations},
    )
    blocking_items.extend(no_table.blocking_items)

    _set_check(
        checks,
        "toc_field",
        "pass" if 'TOC \\o "1-3"' in document_xml else "failed",
        "Word TOC field missing." if 'TOC \\o "1-3"' not in document_xml else "",
    )
    _set_check(
        checks,
        "roman_front_matter_page_numbers",
        "pass" if 'w:fmt="upperRoman"' in document_xml else "failed",
        "Roman front matter page numbering missing." if 'w:fmt="upperRoman"' not in document_xml else "",
    )
    _set_check(
        checks,
        "body_decimal_page_restart",
        "pass"
        if re.search(r"<w:pgNumType\b[^>]*w:start=\"1\"[^>]*(?:w:fmt=\"decimal\")?", document_xml)
        else "failed",
        "Body decimal page restart missing."
        if not re.search(r"<w:pgNumType\b[^>]*w:start=\"1\"[^>]*(?:w:fmt=\"decimal\")?", document_xml)
        else "",
    )

    forbidden = _forbidden_visible_markers(all_text, candidate)
    _set_check(
        checks,
        "forbidden_body_text",
        "failed" if forbidden else "pass",
        ", ".join(forbidden),
        {"markers": forbidden},
    )
    if forbidden:
        blocking_items.append("forbidden_body_text: " + ", ".join(forbidden))

    paragraph_sample = _natural_paragraph_sample(candidate)
    _set_check(
        checks,
        "natural_paragraph_flow_sample",
        "pass" if paragraph_sample["short_line_ratio"] <= 0.75 else "needs_review",
        "",
        paragraph_sample,
    )

    failed_checks = [
        name
        for name, payload in checks.items()
        if payload.get("status") == "failed"
    ]
    if failed_checks:
        blocking_items.extend(
            f"layout_check_failed: {name}"
            for name in failed_checks
            if not any(name in item for item in blocking_items)
        )
    if sample_mode == "truncated":
        notes.append("truncated sample mode: body completeness is not checked.")
    report["status"] = "failed" if blocking_items else _aggregate_status(checks)
    _write_report_if_requested(report, out)
    return report


def _set_check(
    checks: dict[str, dict[str, object]],
    name: str,
    status: str,
    message: str,
    extra: dict[str, object] | None = None,
) -> None:
    payload: dict[str, object] = {"status": status}
    if message:
        payload["message"] = message
    if extra:
        payload.update(extra)
    checks[name] = payload


def _read_part(docx_path: Path, part_name: str) -> str:
    try:
        with zipfile.ZipFile(docx_path) as package:
            return package.read(part_name).decode("utf-8", errors="replace")
    except Exception:
        return ""


def _visible_text(docx_path: Path) -> str:
    parts: list[str] = []
    try:
        document = Document(str(docx_path))
    except Exception:
        return ""
    parts.extend(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    parts.append(_xml_text(_read_part(docx_path, "word/document.xml")))
    return "\n".join(part for part in parts if part)


def _xml_text(document_xml: str) -> str:
    try:
        root = ET.fromstring(document_xml)
    except ET.ParseError:
        return ""
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    return "\n".join(node.text or "" for node in root.findall(".//w:t", ns))


def _forbidden_visible_markers(text: str, docx_path: Path) -> list[str]:
    literal_markers = (
        "[Figure inserted]",
        "[Figure requires review]",
        "D:\\",
        ".worktrees",
        "output_work_",
        "MERGEFORMAT",
        "公式章",
        "下一章",
    )
    found = [marker for marker in literal_markers if marker in text]
    found.extend(sorted(set(re.findall(r"image\d+\.(?:wmf|emf)", text, flags=re.IGNORECASE))))
    found.extend(_cover_field_paragraph_leaks(docx_path))
    if re.search(r"(?m)^References\s*$", text):
        found.append("References")
    return sorted(set(found))


def _cover_field_paragraph_leaks(docx_path: Path) -> list[str]:
    labels = ("院（系）名称", "专业名称", "学生姓名", "指导教师")
    try:
        document = Document(str(docx_path))
    except Exception:
        return []
    leaks: list[str] = []
    for paragraph in document.paragraphs:
        compact = re.sub(r"\s+", "", paragraph.text or "")
        if not compact:
            continue
        hits = [label for label in labels if re.sub(r"\s+", "", label) in compact]
        exact_hits = [label for label in labels if compact == re.sub(r"\s+", "", label)]
        if exact_hits:
            leaks.extend(exact_hits)
        elif len(hits) >= 2:
            leaks.extend(hits)
    return leaks


def _natural_paragraph_sample(docx_path: Path) -> dict[str, object]:
    try:
        document = Document(str(docx_path))
    except Exception:
        return {"sample_count": 0, "short_line_ratio": 1.0, "paragraphs": []}
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    start = next(
        (
            index
            for index, text in enumerate(paragraphs)
            if "\t" not in text and re.match(r"^1\s+\S+", text)
        ),
        0,
    )
    candidates = [
        text
        for text in paragraphs[start + 1 :]
        if len(text) >= 8
        and not text.startswith(("第", "图", "表"))
        and not re.match(r"^\d+(?:\.\d+)*\s+\S+", text)
    ][:5]
    if not candidates:
        return {"sample_count": 0, "short_line_ratio": 1.0, "paragraphs": []}
    short_count = sum(1 for item in candidates if len(item) < 18)
    return {
        "sample_count": len(candidates),
        "short_line_ratio": short_count / len(candidates),
        "paragraphs": candidates,
    }


def _aggregate_status(checks: dict[str, dict[str, object]]) -> str:
    statuses = {str(payload.get("status", "")) for payload in checks.values()}
    if "failed" in statuses:
        return "failed"
    if "needs_review" in statuses:
        return "needs_review"
    return "pass"


def _write_report_if_requested(report: dict[str, object], out: Path | None) -> None:
    if out is None:
        return
    destination = Path(out)
    if destination.suffix.lower() != ".json":
        destination = destination / "layout_consistency_report.json"
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Validate render-level BUAA layout consistency.")
    parser.add_argument("--reference", type=Path, required=True)
    parser.add_argument("--candidate", type=Path, required=True)
    parser.add_argument("--sample-mode", choices=("full", "truncated"), default="full")
    parser.add_argument("--out", type=Path, default=Path("layout_consistency_report.json"))
    args = parser.parse_args(argv)

    report = validate_layout_consistency(
        args.reference,
        args.candidate,
        sample_mode=args.sample_mode,
        out=args.out,
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if report["status"] == "failed" else 0


if __name__ == "__main__":
    raise SystemExit(main())
